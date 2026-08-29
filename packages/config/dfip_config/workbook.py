"""Parse and validate publisher Logic / Labels workbooks.

Accepts ``.xlsx`` only. Invalid files raise ``CatalogWorkbookError`` and do
not produce a catalog snapshot. Duplicate Campaign Name keys are allowed
(first ``row_order`` wins at resolve time). Duplicate Filter Logic 1 values
in a Labels file are rejected.
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from dfip_config.catalog import LABELS_KIND, LOGIC_KIND, CatalogKind
from dfip_config.resolve import match_key
from dfip_config.store import CAMPAIGN_OUTPUT_FIELDS

LOGIC_REQUIRED = ("campaign_name",)
LABEL_REQUIRED = ("group_name", "filter_logic_1_value")

LOGIC_ALIASES: dict[str, tuple[str, ...]] = {
    "campaign_name": ("campaign_name", "Campaign Name"),
    "filter_logic_1": ("filter_logic_1", "Filter Logic 1"),
    "filter_logic_2": ("filter_logic_2", "Filter Logic 2"),
    "amc_status_filter_logic_3": (
        "amc_status_filter_logic_3",
        "AMC Status - Filter Logic 3",
    ),
    "amc_device_category_filter_logic_4": (
        "amc_device_category_filter_logic_4",
        "AMC Device Category -  Filter Logic 4",
        "AMC Device Category - Filter Logic 4",
    ),
    "amc_product_cat_filter_logic_5": (
        "amc_product_cat_filter_logic_5",
        "AMC Product Cat -  Filter Logic 5",
        "AMC Product Cat - Filter Logic 5",
    ),
    "manual_or_automated": ("manual_or_automated", "Manual Or Automated"),
}

LABEL_ALIASES: dict[str, tuple[str, ...]] = {
    "group_name": ("group_name", "Group Name", "Filter Logic 1_2"),
    "filter_logic_1_value": (
        "filter_logic_1_value",
        "Filter Logic 1",
        "Filter Logic 1 Value",
    ),
}


@dataclass(frozen=True)
class CatalogIssue:
    row: int | None
    code: str
    detail: str


class CatalogWorkbookError(Exception):
    def __init__(self, message: str, issues: tuple[CatalogIssue, ...] = ()) -> None:
        super().__init__(message)
        self.message = message
        self.issues = issues


@dataclass(frozen=True)
class ParsedCatalog:
    kind: CatalogKind
    rows: tuple[dict[str, Any], ...]
    distinct_key_count: int
    duplicate_key_count: int


def parse_catalog_workbook(kind: CatalogKind, payload: bytes) -> ParsedCatalog:
    """Return a validated snapshot. Does not persist."""
    if kind not in {LOGIC_KIND, LABELS_KIND}:
        raise CatalogWorkbookError(
            "Unknown catalog kind.",
            (CatalogIssue(None, "INVALID_KIND", "kind must be logic or labels"),),
        )
    try:
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:
        raise CatalogWorkbookError(
            "Workbook could not be read.",
            (CatalogIssue(None, "INVALID_FILE_TYPE", "Only .xlsx workbooks are accepted."),),
        ) from exc
    try:
        sheet = workbook.active
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if header_row is None:
            raise CatalogWorkbookError(
                "Workbook is missing a header row.",
                (CatalogIssue(1, "INVALID_SCHEMA", "Header row is required."),),
            )
        headers = [_cell_text(value) for value in header_row]
        aliases = LOGIC_ALIASES if kind == LOGIC_KIND else LABEL_ALIASES
        required = LOGIC_REQUIRED if kind == LOGIC_KIND else LABEL_REQUIRED
        mapping = _map_headers(headers, aliases)
        missing = [name for name in required if name not in mapping]
        if missing:
            raise CatalogWorkbookError(
                "Workbook is missing required columns.",
                tuple(
                    CatalogIssue(1, "INVALID_SCHEMA", f"Missing required column: {name}")
                    for name in missing
                ),
            )
        parsed: list[dict[str, Any]] = []
        issues: list[CatalogIssue] = []
        for excel_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            cells = [_cell_text(value) for value in values]
            record = {
                field: cells[index] if index < len(cells) else None
                for field, index in mapping.items()
            }
            if _row_empty(record, required):
                continue
            if kind == LOGIC_KIND:
                if record.get("campaign_name") is None:
                    issues.append(
                        CatalogIssue(
                            excel_row,
                            "INVALID_CONTENT",
                            "Campaign Name is required.",
                        )
                    )
                    continue
            else:
                for field in required:
                    if record.get(field) is None:
                        issues.append(
                            CatalogIssue(
                                excel_row,
                                "INVALID_CONTENT",
                                f"{field} is required.",
                            )
                        )
                if any(item.row == excel_row for item in issues):
                    continue
            record["_source_row"] = excel_row
            parsed.append(record)
    finally:
        workbook.close()

    if issues:
        raise CatalogWorkbookError("Workbook content is invalid.", tuple(issues))
    if not parsed:
        raise CatalogWorkbookError(
            "Workbook has no data rows.",
            (CatalogIssue(None, "INVALID_CONTENT", "At least one data row is required."),),
        )

    rows, distinct, duplicates, extra_issues = _finalize_rows(kind, parsed)
    if extra_issues:
        raise CatalogWorkbookError("Workbook content is invalid.", extra_issues)
    return ParsedCatalog(
        kind=kind,
        rows=rows,
        distinct_key_count=distinct,
        duplicate_key_count=duplicates,
    )


def _cell_text(value: object) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value)


def _map_headers(headers: list[str | None], aliases: dict[str, tuple[str, ...]]) -> dict[str, int]:
    present = {header: index for index, header in enumerate(headers) if header is not None}
    mapping: dict[str, int] = {}
    for field, names in aliases.items():
        for name in names:
            if name in present:
                mapping[field] = present[name]
                break
    return mapping


def _row_empty(record: dict[str, Any], required: tuple[str, ...]) -> bool:
    return all(record.get(field) is None for field in required)


def _finalize_rows(
    kind: CatalogKind, parsed: list[dict[str, Any]]
) -> tuple[tuple[dict[str, Any], ...], int, int, tuple[CatalogIssue, ...]]:
    issues: list[CatalogIssue] = []
    rows: list[dict[str, Any]] = []
    seen: dict[str, int] = {}
    for index, record in enumerate(parsed, start=1):
        if kind == LOGIC_KIND:
            row = {
                "row_order": index,
                "campaign_name": record["campaign_name"],
            }
            for field in CAMPAIGN_OUTPUT_FIELDS:
                row[field] = record.get(field)
            key = match_key(record["campaign_name"])
            if key is not None:
                seen[key] = seen.get(key, 0) + 1
            rows.append(row)
        else:
            value = record["filter_logic_1_value"]
            if value in seen:
                issues.append(
                    CatalogIssue(
                        int(record.get("_source_row") or index + 1),
                        "INVALID_CONTENT",
                        "Duplicate Filter Logic 1 value.",
                    )
                )
                continue
            seen[str(value)] = 1
            rows.append(
                {
                    "row_order": index,
                    "group_name": record["group_name"],
                    "filter_logic_1_value": value,
                }
            )
    if issues:
        return (), 0, 0, tuple(issues)
    distinct = len(seen)
    if kind == LOGIC_KIND:
        duplicate = sum(1 for count in seen.values() if count > 1)
    else:
        duplicate = 0
    return tuple(rows), distinct, duplicate, ()
