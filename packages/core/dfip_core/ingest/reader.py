"""Read-only Web Engage worksheet access. Does not transform or label rows."""

from __future__ import annotations

from collections.abc import Iterator
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import date, datetime, time
from hashlib import sha256 as sha256_hash
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from dfip_core.ingest.headers import (
    HEADER_SCAN_ROWS,
    PREFERRED_SHEET,
    HeaderContractError,
    HeaderMatch,
    expected_source_headers,
    match_header_rows,
)


@dataclass(frozen=True)
class WorkbookLayout:
    path: str
    sheet_names: tuple[str, ...]
    worksheet_name: str
    match: HeaderMatch


@dataclass(frozen=True)
class SourceDataRow:
    source_row_number: int
    values: tuple[Any, ...]
    raw: dict[str, Any]
    campaign_id: str | None
    variation_id: str | None
    day: date | None
    is_empty: bool


def sha256_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    digest = sha256_hash()
    with path.open("rb") as handle:
        while True:
            chunk = handle.read(chunk_size)
            if not chunk:
                break
            digest.update(chunk)
    return digest.hexdigest()


def serialize_cell(value: Any) -> Any:
    """JSON-safe form of an Excel cell. Strings are not trimmed or cased."""
    if value is None:
        return None
    if isinstance(value, datetime):
        if value.tzinfo is None and value.time() == time.min:
            return value.date().isoformat()
        return value.replace(microsecond=0).isoformat()
    if isinstance(value, date) and not isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, time):
        return value.replace(microsecond=0).isoformat()
    if isinstance(value, bool | int | float | str):
        return value
    return str(value)


def typed_day(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and len(value) == 10 and value[4] == "-" and value[7] == "-":
        try:
            return date.fromisoformat(value)
        except ValueError:
            return None
    return None


def typed_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        return value
    serialized = serialize_cell(value)
    if serialized is None:
        return None
    return str(serialized)


def _pad(values: tuple[Any, ...] | list[Any], length: int) -> list[Any]:
    items = list(values)
    if len(items) < length:
        items.extend([None] * (length - len(items)))
    return items


def payload_from_values(
    values: list[Any],
    start_index: int,
    extra_headers: tuple[str, ...] = (),
) -> dict[str, Any]:
    headers = expected_source_headers()
    needed = start_index + 57 + len(extra_headers)
    padded = _pad(values, needed)
    region = padded[start_index : start_index + 57]
    if len(region) != 57:
        raise HeaderContractError(
            "INSUFFICIENT_COLUMNS",
            f"row has {len(values)} cells; expected at least {needed}",
        )
    payload = {header: serialize_cell(cell) for header, cell in zip(headers, region, strict=True)}
    extra_region = padded[start_index + 57 : start_index + 57 + len(extra_headers)]
    for header, cell in zip(extra_headers, extra_region, strict=True):
        payload[header] = serialize_cell(cell)
    return payload


def row_is_empty(raw: dict[str, Any]) -> bool:
    headers = expected_source_headers()
    return all(value is None or value == "" for key, value in raw.items() if key in headers)


def build_source_row(
    excel_row: int,
    values: list[Any],
    start_index: int,
    extra_headers: tuple[str, ...] = (),
) -> SourceDataRow:
    raw = payload_from_values(values, start_index, extra_headers)
    canonical = tuple(_pad(values, start_index + 57)[start_index : start_index + 57])
    return SourceDataRow(
        source_row_number=excel_row,
        values=canonical,
        raw=raw,
        campaign_id=typed_text(raw.get("Campaign ID")),
        variation_id=typed_text(raw.get("Variation ID")),
        day=typed_day(raw.get("Day")),
        is_empty=row_is_empty(raw),
    )


@contextmanager
def open_source_workbook(path: Path) -> Iterator[tuple[WorkbookLayout, Any]]:
    """Open a workbook once for header inspect + row iteration.

    Callers must finish iterating rows before the context exits. ``read_only``
    worksheets cannot be rewound.
    """
    path = Path(path)
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:  # noqa: BLE001 — surface Excel parse failures
        raise HeaderContractError("UNREADABLE_WORKBOOK", f"{path.name}: {exc}") from exc
    try:
        sheet_names = tuple(workbook.sheetnames)
        worksheet_name, match = _resolve_layout(workbook, sheet_names)
        layout = WorkbookLayout(
            path=str(path),
            sheet_names=sheet_names,
            worksheet_name=worksheet_name,
            match=match,
        )
        yield layout, workbook
    finally:
        workbook.close()


def inspect_workbook(path: Path) -> WorkbookLayout:
    """Validate sheet + 57-header region. Read-only. Does not ingest rows.

    Trailing approved extras after the 57-window are recorded on
    ``match.extra_headers``. Unapproved trailing headers fail closed.
    """
    with open_source_workbook(path) as (layout, _workbook):
        return layout


def worksheet_data_row_hint(workbook: Any, layout: WorkbookLayout) -> int | None:
    """Excel dimension hint for data rows below the header. May be absent.

    Read-only worksheets without a stored dimension must not call ``max_row``:
    openpyxl may stream the whole sheet just to count rows, which doubles the
    later staging pass. Missing hints keep progress indeterminate with truthful
    current counts.
    """
    worksheet = workbook[layout.worksheet_name]
    header_row = layout.match.header_row
    if getattr(workbook, "read_only", False):
        cached = getattr(worksheet, "_max_row", None)
        if not isinstance(cached, int):
            return None
        max_row = cached
    else:
        max_row = getattr(worksheet, "max_row", None)
    if isinstance(max_row, int) and isinstance(header_row, int) and max_row > header_row:
        return max_row - header_row
    return None


def iter_source_rows_from_workbook(
    workbook: Any, layout: WorkbookLayout
) -> Iterator[SourceDataRow]:
    worksheet = workbook[layout.worksheet_name]
    start = layout.match.start_index
    extra_headers = layout.match.extra_headers
    max_col = start + 57 + len(extra_headers)
    first_data = layout.match.header_row + 1
    for excel_row, row in enumerate(
        worksheet.iter_rows(
            min_row=first_data,
            max_col=max_col,
            values_only=True,
        ),
        start=first_data,
    ):
        yield build_source_row(excel_row, list(row), start, extra_headers)


def iter_source_rows(path: Path, layout: WorkbookLayout) -> Iterator[SourceDataRow]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        yield from iter_source_rows_from_workbook(workbook, layout)
    finally:
        workbook.close()


def _scan_headers(worksheet: Any) -> HeaderMatch:
    scanned: list[tuple[int, list[object]]] = []
    for excel_row, row in enumerate(
        worksheet.iter_rows(
            min_row=1,
            max_row=HEADER_SCAN_ROWS,
            max_col=120,
            values_only=True,
        ),
        start=1,
    ):
        scanned.append((excel_row, list(row)))
    return match_header_rows(scanned)


def _resolve_layout(workbook: Any, sheet_names: tuple[str, ...]) -> tuple[str, HeaderMatch]:
    if PREFERRED_SHEET in workbook.sheetnames:
        return PREFERRED_SHEET, _scan_headers(workbook[PREFERRED_SHEET])
    matches: list[tuple[str, HeaderMatch]] = []
    for name in sheet_names:
        try:
            matches.append((name, _scan_headers(workbook[name])))
        except HeaderContractError:
            continue
    if len(matches) == 1:
        return matches[0]
    if len(matches) > 1:
        found = ", ".join(item[0] for item in matches)
        raise HeaderContractError(
            "AMBIGUOUS_REGION",
            f"multiple worksheets contain the 57 source headers: {found}",
        )
    raise HeaderContractError(
        "MISSING_SHEET",
        "worksheet "
        f"{PREFERRED_SHEET!r} was not found and no other sheet has the 57 "
        f"source headers. sheets={list(sheet_names)}.",
    )
