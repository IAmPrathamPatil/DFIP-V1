"""Native Excel PivotTable + slicer package for the nine Daily Report sheets.

Authoring is ZIP/OOXML (no COM generation path). PivotTable parts are
re-serialized through openpyxl so desktop Excel accepts them. One shared
pivot cache binds to the PublishedFacts worksheet in the tracked template, and
P8 rebinds that cache to the publication's static PublishedFacts cells.

Python UNIQUE/FILTER/SUMIFS in daily_report.report_formula remains the P10
reconstruction contract. It is not stored on the delivered report sheets.
"""

from __future__ import annotations

import io
import re
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path
from uuid import NAMESPACE_URL, UUID, uuid5
from zipfile import ZipFile

from openpyxl.utils import get_column_letter

from dfip_web.client_workbook import FACT_HEADERS, FAKE_MASHUP_ZIP_PARTS
from dfip_web.daily_report import (
    AMC_DAYWISE,
    AMC_GROUP,
    AMC_SPLIT,
    AMC_VERTICAL,
    CAPTION_KEY_COL,
    CAPTION_NAME_COL,
    CHANNEL,
    D2C,
    D2C_GROUP,
    EMPTY_STATE,
    MEASURE_HEADERS,
    MEASURE_OPS,
    NS_MAIN,
    NS_REL,
    OVERALL,
    PUBLISHED_HEADER,
    REPORT_CONTRACTS,
    REPORT_SHEET_NAMES,
    SERVICE,
    SERVICE_FILTER_LOGIC_1,
    SUB_SPLIT,
    VERTICAL,
    XF_GROUP_BANNER,
    XF_HINT,
    XF_PURPOSE,
    XF_TITLE,
    ZERO_VS_BLANK_HINT,
    ZERO_VS_BLANK_HINT_VERTICAL,
    ReportSheetContract,
    _caption_pairs,
    _clone_zipinfo,
    _col,
    _ensure_report_styles,
    _inline,
    _new_sheet_zipinfo,
    _next_relationship_id,
    _rels_with_metadata,
    _sheet_part_map,
    _types_with_metadata,
    _xml_attr,
    _xml_text,
    mutate_xlsx,
)
from dfip_web.published_facts_mashup import extract_published_facts_section_from_package
from dfip_web.report_format import apply_reference_formatting, datafield_numfmt

PIVOT_CACHE_ID = 1
# x14 slicer tabular id. Distinct from workbook cacheId; Excel stores both.
SLICER_PIVOT_CACHE_ID = 110000011
PIVOT_CACHE_PART = "xl/pivotCache/pivotCacheDefinition1.xml"
PIVOT_RECORDS_PART = "xl/pivotCache/pivotCacheRecords1.xml"
# Excel worksheet max row. Refreshable cache must grow past the download stamp.
PIVOT_CACHE_REFRESHABLE_ROWS = 1_048_575
# Query-table defined name Power Query resizes on Refresh All. A frozen
# worksheet ``ref`` (even A1:AS1048576) is rewritten by Excel after the first
# cache refresh to the then-current row count, so later months never enter
# the native PivotCache.
PIVOT_CACHE_QUERY_NAME = "ExternalData_1"

NS_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
NS_XR = "http://schemas.microsoft.com/office/spreadsheetml/2014/revision"
NS_SLICER_MAIN = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"
NS_XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
NS_PIVOT_CACHE_REL = f"{NS_REL}/pivotCacheDefinition"
NS_PIVOT_TABLE_REL = f"{NS_REL}/pivotTable"
NS_DRAWING_REL = f"{NS_REL}/drawing"
NS_SLICER_REL = "http://schemas.microsoft.com/office/2007/relationships/slicer"
NS_SLICER_CACHE_REL = "http://schemas.microsoft.com/office/2007/relationships/slicerCache"
NS_PIVOT_RECORDS_REL = f"{NS_REL}/pivotCacheRecords"

CT_CACHE = "application/vnd.openxmlformats-officedocument.spreadsheetml.pivotCacheDefinition+xml"
CT_RECORDS = "application/vnd.openxmlformats-officedocument.spreadsheetml.pivotCacheRecords+xml"
CT_PIVOT = "application/vnd.openxmlformats-officedocument.spreadsheetml.pivotTable+xml"
CT_SLICER = "application/vnd.ms-excel.slicer+xml"
CT_SLICER_CACHE = "application/vnd.ms-excel.slicerCache+xml"
CT_DRAWING = "application/vnd.openxmlformats-officedocument.drawing+xml"

SLICER_CACHES_URI = "{BBE1A952-AA13-448e-AADC-164F8A28A991}"
SHEET_SLICER_URI = "{A8765BA9-456A-4dab-B4F3-ACF838C121DE}"
CACHE_X14_URI = "{725AE2AE-9491-48be-B2B4-4EB974FC3084}"

PIVOT_TABLE_NAMES: dict[str, str] = {
    REPORT_CONTRACTS[0].name: "PivotOverall",
    REPORT_CONTRACTS[1].name: "PivotVertical",
    REPORT_CONTRACTS[2].name: "PivotChannel",
    REPORT_CONTRACTS[3].name: "PivotSubSplit",
    REPORT_CONTRACTS[4].name: "PivotAmcDaywise",
    REPORT_CONTRACTS[5].name: "PivotAmcSplit",
    REPORT_CONTRACTS[6].name: "PivotAmcVertical",
    REPORT_CONTRACTS[7].name: "PivotD2c",
    REPORT_CONTRACTS[8].name: "PivotService",
}

CACHE_FIELD_CAPTIONS: dict[str, str] = {
    "filter_logic_1_group": "Filter Logic 1_2",
}

# Stable template uniques for page fields. Snapshot generation replaces these
# with first-seen PublishedFacts values so item indexes match the worksheet.
PAGE_GROUP_ITEMS: tuple[str, ...] = (
    "Group1",
    "Group2",
    "Group3",
    "Group4",
    "Group5",
    "Group6",
    "Group7",
)

_SKIP_ZIP_PREFIXES = (
    "xl/pivotCache/",
    "xl/pivotTables/",
    "xl/slicers/",
    "xl/slicerCaches/",
)

_SLICER_FALLBACK = "This shape represents a slicer. Slicers are supported in Excel 2010 or later."


@dataclass(frozen=True)
class ReportPivotBinding:
    contract: ReportSheetContract
    sheet_part: str
    sheet_id: int
    index: int  # 1-based pivot/slicer/drawing number

    @property
    def pivot_name(self) -> str:
        return PIVOT_TABLE_NAMES[self.contract.name]

    @property
    def pivot_part(self) -> str:
        return f"xl/pivotTables/pivotTable{self.index}.xml"

    @property
    def slicer_part(self) -> str:
        return f"xl/slicers/slicer{self.index}.xml"

    @property
    def drawing_part(self) -> str:
        return f"xl/drawings/drawing{self.index}.xml"


def cache_field_names() -> tuple[str, ...]:
    names = list(FACT_HEADERS)
    for header, kind, *_rest in MEASURE_OPS:
        if kind != "sum":
            names.append(header)
    return tuple(names)


def cache_field_index(header: str) -> int:
    published = PUBLISHED_HEADER.get(header, header)
    names = cache_field_names()
    return names.index(published)


def calculated_cache_fields() -> tuple[tuple[str, str], ...]:
    fields: list[tuple[str, str]] = []
    for op in MEASURE_OPS:
        header, kind = op[0], op[1]
        if kind == "sum":
            continue
        left = _formula_atom(op[2])
        right = _formula_atom(op[3])
        formula = f"{left}-{right}" if kind == "diff" else f"{left}/{right}"
        fields.append((header, formula))
    return tuple(fields)


def _formula_atom(header: str) -> str:
    if re.search(r"[^A-Za-z0-9]", header):
        return "'" + header.replace("'", "''") + "'"
    return header


def _uid(key: str) -> str:
    raw = uuid5(NAMESPACE_URL, f"dfip-p11:{key}")
    blob = bytearray(raw.bytes)
    blob[6] = (blob[6] & 0x0F) | 0x40
    blob[8] = (blob[8] & 0x3F) | 0x80
    return str(UUID(bytes=bytes(blob))).upper()


def _string_shared_items_xml(values: Sequence[str]) -> str:
    """String shared items with no numeric type flags (those flags crashed refresh)."""
    if not values:
        return "<sharedItems/>"
    body = "".join(f'<s v="{_xml_attr(value)}"/>' for value in values)
    return f'<sharedItems count="{len(values)}">{body}</sharedItems>'


def _shared_items_xml(field_name: str) -> str:
    """Empty shared items except page-filter fields, which need named defaults.

    Numeric min/max=0 on blank columns, and ``rowItems t=data x=0`` against empty
    Month/Day, crashed desktop Excel (RPC 0x8001010A). Page-filter fields only
    store the default string values Excel matches by caption after populate.
    """
    if field_name == "filter_logic_1_group":
        return _string_shared_items_xml(PAGE_GROUP_ITEMS)
    if field_name == "Filter Logic 1":
        return _string_shared_items_xml((SERVICE_FILTER_LOGIC_1,))
    return "<sharedItems/>"


def _published_facts_ref(row_count: int) -> str:
    last_row = max(1 + max(row_count, 1), 2)
    last_col = get_column_letter(len(FACT_HEADERS))
    return f"A1:{last_col}{last_row}"


def _cache_source_xml(*, row_count: int = 1) -> str:
    ref = _published_facts_ref(row_count)
    return (
        '<cacheSource type="worksheet">'
        f'<worksheetSource ref="{_xml_attr(ref)}" sheet="PublishedFacts"/>'
        "</cacheSource>"
    )


def _cache_source_query_xml() -> str:
    return (
        '<cacheSource type="worksheet">'
        f'<worksheetSource name="{_xml_attr(PIVOT_CACHE_QUERY_NAME)}" '
        'sheet="PublishedFacts"/>'
        "</cacheSource>"
    )


def pivot_cache_definition_xml() -> str:
    fields: list[str] = []
    for name in FACT_HEADERS:
        caption = CACHE_FIELD_CAPTIONS.get(name)
        cap = f' caption="{_xml_attr(caption)}"' if caption else ""
        fields.append(
            f'<cacheField name="{_xml_attr(name)}"{cap} numFmtId="0">'
            f"{_shared_items_xml(name)}</cacheField>"
        )
    for name, formula in calculated_cache_fields():
        fields.append(
            f'<cacheField name="{_xml_attr(name)}" numFmtId="0" '
            f'formula="{_xml_attr(formula)}" databaseField="0" sqlType="0"/>'
        )
    uid = _uid("pivot-cache-1")
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
        f'<pivotCacheDefinition xmlns="{NS_MAIN}" xmlns:r="{NS_REL}" '
        f'xmlns:mc="{NS_MC}" mc:Ignorable="xr" xmlns:xr="{NS_XR}" '
        'r:id="rId1" refreshedBy="DFIP" backgroundQuery="0" createdVersion="8" '
        'refreshedVersion="8" minRefreshableVersion="3" recordCount="0" '
        f'refreshOnLoad="0" xr:uid="{{{uid}}}">'
        f"{_cache_source_xml(row_count=1)}"
        f'<cacheFields count="{len(fields)}">{"".join(fields)}</cacheFields>'
        f'<extLst><ext uri="{CACHE_X14_URI}" '
        'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
        f'<x14:pivotCacheDefinition pivotCacheId="{SLICER_PIVOT_CACHE_ID}"/>'
        "</ext></extLst></pivotCacheDefinition>"
    )


def pivot_cache_records_xml() -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
        f'<pivotCacheRecords xmlns="{NS_MAIN}" count="0"/>'
    )


def bind_pivot_cache_for_snapshot(cache_xml: str, row_count: int) -> str:
    """Point the shared cache at static PublishedFacts cells (P8 snapshot).

    Historical snapshots keep ``refreshOnLoad="0"`` so Excel does not auto-
    refresh from the live API. ``recordCount`` stays 0 until Excel rebuilds
    records; inflating it without ``pivotCacheRecords`` crashes refresh.
    """
    source = _cache_source_xml(row_count=row_count)
    xml = re.sub(r"<cacheSource\b.*?</cacheSource>", source, cache_xml, count=1, flags=re.DOTALL)
    xml = re.sub(r"<cacheSource\b[^>]*/>", source, xml, count=1)
    if "refreshOnLoad=" in xml:
        xml = re.sub(r'refreshOnLoad="\d+"', 'refreshOnLoad="0"', xml, count=1)
    else:
        xml = xml.replace("<pivotCacheDefinition ", '<pivotCacheDefinition refreshOnLoad="0" ', 1)
    xml = re.sub(r'recordCount="\d+"', 'recordCount="0"', xml, count=1)
    return xml


def bind_refreshable_query_defined_name(workbook_xml: str, row_count: int) -> str:
    """Point ExternalData_1 at the stamped PublishedFacts rows.

    The tracked template name is ``PublishedFacts!$A$1`` (one cell). A cache
    bound to that name cannot refresh until Power Query resizes it. Seed the
    name to the snapshot so PivotCache.Refresh and Refresh All both see a
    real table; Power Query still rewrites the name when history grows.
    """
    last_row = max(1 + max(row_count, 1), 2)
    last_col = get_column_letter(len(FACT_HEADERS))
    formula = f"PublishedFacts!$A$1:${last_col}${last_row}"
    replacement = (
        f'<definedName name="{PIVOT_CACHE_QUERY_NAME}" localSheetId="0">{formula}</definedName>'
    )
    pattern = rf'<definedName name="{re.escape(PIVOT_CACHE_QUERY_NAME)}"[^>]*>[^<]*</definedName>'
    if re.search(pattern, workbook_xml):
        return re.sub(pattern, replacement, workbook_xml, count=1)
    return workbook_xml


def bind_pivot_cache_for_refreshable(cache_xml: str) -> str:
    """Point the shared cache at the Power Query table name.

    Refreshable workbooks accumulate months after Refresh All. Binding the
    cache to a worksheet ``ref`` — including a full-column range — lets Excel
    rewrite the source to the first refresh's row count. Later history then
    lands in ``ExternalData_1`` / PublishedFacts UsedRange while all nine
    PivotTables keep reading the frozen slice. The query-table name grows
    with the connection refresh and is the cache source Refresh All uses.
    ``recordCount`` stays 0 until Excel rebuilds records.
    """
    source = _cache_source_query_xml()
    xml = re.sub(r"<cacheSource\b.*?</cacheSource>", source, cache_xml, count=1, flags=re.DOTALL)
    xml = re.sub(r"<cacheSource\b[^>]*/>", source, xml, count=1)
    if "refreshOnLoad=" in xml:
        xml = re.sub(r'refreshOnLoad="\d+"', 'refreshOnLoad="0"', xml, count=1)
    else:
        xml = xml.replace("<pivotCacheDefinition ", '<pivotCacheDefinition refreshOnLoad="0" ', 1)
    xml = re.sub(r'recordCount="\d+"', 'recordCount="0"', xml, count=1)
    return xml


def _page_fields(contract: ReportSheetContract) -> list[tuple[str, str | None]]:
    """(display header, selected shared-item value or None for all)."""
    pages: list[tuple[str, str | None]] = []
    # Service Filter Logic 1 is pinned by captionEqual + slicer, not a pageField
    # item index. Excel remaps that index after Refresh All and hides every row.
    if contract.default_filter_logic_1 and contract.name != SERVICE:
        pages.append(("Filter Logic 1", contract.default_filter_logic_1))
    if contract.name == SERVICE:
        pages.append(("Filter Logic 2", None))
    elif contract.default_group:
        pages.append(("Filter Logic 1_2", contract.default_group))
    return pages


def _axis_for_field(contract: ReportSheetContract, field_name: str) -> str | None:
    published_rows = [PUBLISHED_HEADER.get(name, name) for name in contract.display_fields]
    if field_name in published_rows:
        return "axisRow"
    for header, _selected in _page_fields(contract):
        if PUBLISHED_HEADER.get(header, header) == field_name:
            return "axisPage"
    return None


def _page_field_attrs(field_name: str) -> str:
    if field_name in {"Filter Logic 1", "Filter Logic 2"}:
        return ' multipleItemSelectionAllowed="1"'
    return ""


def _page_item_values(field_name: str) -> tuple[str, ...]:
    if field_name in {"Filter Logic 1_2", "filter_logic_1_group"}:
        return PAGE_GROUP_ITEMS
    if field_name == "Filter Logic 1":
        return (SERVICE_FILTER_LOGIC_1,)
    return ()


def _page_items_xml(field_name: str, selected: str | None) -> str:
    # Filter Logic 1 uniques grow as months accumulate. A singleton shared
    # item list with x="0" remaps after Refresh All and empties Service.
    if field_name == "Filter Logic 1":
        return '<items count="1"><item t="default"/></items>'
    values = _page_item_values(field_name)
    if not values or selected is None:
        return '<items count="1"><item t="default"/></items>'
    body = "".join(f'<item x="{index}"/>' for index in range(len(values)))
    return f'<items count="{len(values) + 1}">{body}<item t="default"/></items>'


def _pivot_field_xml(contract: ReportSheetContract, field_name: str, is_data: bool) -> str:
    if is_data:
        extra = ""
        if field_name not in FACT_HEADERS:
            extra = ' dragToRow="0" dragToCol="0" dragToPage="0" defaultSubtotal="0"'
        return f'<pivotField dataField="1" showAll="0"{extra}/>'
    axis = _axis_for_field(contract, field_name)
    selected = None
    for header, value in _page_fields(contract):
        if PUBLISHED_HEADER.get(header, header) == field_name:
            selected = value
            break
    axis_attr = f' axis="{axis}"' if axis else ""
    extra = _page_field_attrs(field_name) if axis == "axisPage" else ""
    if axis == "axisRow":
        items = '<items count="1"><item t="default" sd="1"/></items>'
        return f'<pivotField{axis_attr} showAll="0">{items}</pivotField>'
    if axis == "axisPage":
        items = _page_items_xml(field_name, selected)
        return f'<pivotField{axis_attr} showAll="0"{extra}>{items}</pivotField>'
    return f'<pivotField{axis_attr} showAll="0"/>'


def measure_datafield_display_name(header: str) -> str:
    return header if header not in FACT_HEADERS else f"  {header}"


def measure_data_fields_xml() -> str:
    data_xml = []
    for header in MEASURE_HEADERS:
        fld = cache_field_index(header)
        fmt = datafield_numfmt(header)
        num = 0 if fmt is None else fmt
        display = measure_datafield_display_name(header)
        data_xml.append(
            f'<dataField name="{_xml_attr(display)}" fld="{fld}" '
            f'baseField="0" baseItem="0" numFmtId="{num}"/>'
        )
    return f'<dataFields count="{len(data_xml)}">{"".join(data_xml)}</dataFields>'


def measure_col_items_xml() -> str:
    col_items = "".join(
        f'<i t="data"><x v="{index}"/></i>' for index in range(len(MEASURE_HEADERS))
    )
    return f'<colItems count="{len(MEASURE_HEADERS)}">{col_items}</colItems>'


def sum_measure_headers() -> tuple[str, ...]:
    return tuple(header for header, kind, *_rest in MEASURE_OPS if kind == "sum")


def sum_measure_data_fields_xml() -> str:
    """Additive sums only. Ratio/diff fields cannot survive a Power Query cache rebuild."""
    data_xml = []
    for header in sum_measure_headers():
        fld = cache_field_index(header)
        fmt = datafield_numfmt(header)
        num = 0 if fmt is None else fmt
        display = measure_datafield_display_name(header)
        data_xml.append(
            f'<dataField name="{_xml_attr(display)}" fld="{fld}" '
            f'baseField="0" baseItem="0" numFmtId="{num}"/>'
        )
    return f'<dataFields count="{len(data_xml)}">{"".join(data_xml)}</dataFields>'


def sum_measure_col_items_xml() -> str:
    headers = sum_measure_headers()
    col_items = "".join(f'<i t="data"><x v="{index}"/></i>' for index in range(len(headers)))
    return f'<colItems count="{len(headers)}">{col_items}</colItems>'


def ensure_measure_data_fields_xml(xml: str) -> str:
    """Keep all 22 reference measures, including ratio/diff calculated fields."""
    xml, replaced = re.subn(
        r"<dataFields\b.*?</dataFields>",
        measure_data_fields_xml(),
        xml,
        count=1,
        flags=re.DOTALL,
    )
    if replaced != 1:
        return xml
    xml, _n = re.subn(
        r"<colItems\b.*?</colItems>",
        measure_col_items_xml(),
        xml,
        count=1,
        flags=re.DOTALL,
    )
    return xml


def restore_pivot_measures_excel(pivot) -> None:
    """Re-add MEASURE_HEADERS in reference order after a cache rebuild.

    Excel PivotCache.Refresh from a Power Query table drops calculated
    dataFields (22 → 10 additive sums). Unique CTC / Unique Conversions
    survive but shift left, so the Click-Through / Overall banners sit over
    blank cells. Re-adding in MEASURE_HEADERS order restores the layout.
    """
    xl_hidden = 0
    while int(pivot.DataFields.Count) > 0:
        pivot.DataFields.Item(1).Orientation = xl_hidden
    for header, kind, *_rest in MEASURE_OPS:
        if kind != "sum":
            name, formula = header, None
            for calc_name, calc_formula in calculated_cache_fields():
                if calc_name == header:
                    formula = calc_formula
                    break
            if formula:
                try:
                    pivot.CalculatedFields().Add(name, formula)
                except Exception:
                    pass
        display = measure_datafield_display_name(header)
        field = None
        for candidate in (display, header):
            try:
                field = pivot.PivotFields(candidate)
                break
            except Exception:
                continue
        if field is None:
            continue
        try:
            pivot.AddDataField(field)
        except Exception:
            continue
        fmt = datafield_numfmt(header)
        if fmt is None:
            continue
        try:
            pivot.DataFields.Item(pivot.DataFields.Count).NumberFormat = _excel_number_format(fmt)
        except Exception:
            pass


def _excel_number_format(num_fmt_id: int) -> str:
    return {
        2: "0.00",
        3: "#,##0",
        4: "#,##0.00",
        9: "0%",
        10: "0.00%",
        164: "0.0%",
    }.get(num_fmt_id, "General")


def restore_service_page_excel(pivot) -> None:
    """Pin Service Filter Logic 1 by caption after Excel rebuilds page items."""
    try:
        pivot.PageFields("Filter Logic 1").CurrentPage = SERVICE_FILTER_LOGIC_1
    except Exception:
        pass


def restore_report_pivots_excel(workbook) -> None:
    for contract in REPORT_CONTRACTS:
        pivot = workbook.Worksheets(contract.name).PivotTables(1)
        restore_pivot_measures_excel(pivot)
        if contract.name == SERVICE and contract.default_filter_logic_1:
            restore_service_page_excel(pivot)


def pivot_table_xml(binding: ReportPivotBinding) -> str:
    contract = binding.contract
    names = cache_field_names()
    measure_set = set(MEASURE_HEADERS)
    pivot_fields = []
    for name in names:
        is_measure = name in measure_set
        used_as_axis = _axis_for_field(contract, name) is not None
        pivot_fields.append(_pivot_field_xml(contract, name, is_measure and not used_as_axis))

    row_refs = "".join(
        f'<field x="{cache_field_index(header)}"/>' for header in contract.display_fields
    )
    pages = _page_fields(contract)
    if pages:
        page_xml = "".join(_page_field_tag(header, selected) for header, selected in pages)
        page_block = f'<pageFields count="{len(pages)}">{page_xml}</pageFields>'
    else:
        page_block = ""
    data_xml = []
    for header in MEASURE_HEADERS:
        fld = cache_field_index(header)
        fmt = datafield_numfmt(header)
        num = 0 if fmt is None else fmt
        display = measure_datafield_display_name(header)
        data_xml.append(
            f'<dataField name="{_xml_attr(display)}" fld="{fld}" '
            f'baseField="0" baseItem="0" numFmtId="{num}"/>'
        )
    uid = _uid(f"pivot-{binding.pivot_name}")
    caption = contract.display_fields[0]
    col_items = "".join(
        f'<i t="data"><x v="{index}"/></i>' for index in range(len(MEASURE_HEADERS))
    )
    # Excel requires rowItems/colItems immediately after their field lists.
    # Compact `rowFields, colFields, pageFields, dataFields, rowItems` is rejected.
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
        f'<pivotTableDefinition xmlns="{NS_MAIN}" '
        f'xmlns:r="{NS_REL}" xmlns:mc="{NS_MC}" mc:Ignorable="xr" xmlns:xr="{NS_XR}" '
        f'xr:uid="{{{uid}}}" name="{binding.pivot_name}" cacheId="{PIVOT_CACHE_ID}" '
        'r:id="rId1" dataCaption="Values" showError="1" updatedVersion="8" '
        'minRefreshableVersion="3" itemPrintTitles="1" createdVersion="8" '
        'indent="0" compact="1" outline="1" outlineData="1" compactData="1" '
        f'multipleFieldFilters="0" applyWidthHeightFormats="1" '
        f'rowHeaderCaption="{_xml_attr(caption)}">'
        '<location ref="B9:X10" firstHeaderRow="0" firstDataRow="1" firstDataCol="1"/>'
        f'<pivotFields count="{len(names)}">{"".join(pivot_fields)}</pivotFields>'
        f'<rowFields count="{len(contract.display_fields)}">{row_refs}</rowFields>'
        '<rowItems count="1"><i t="grand"><x/></i></rowItems>'
        '<colFields count="1"><field x="-2"/></colFields>'
        f'<colItems count="{len(MEASURE_HEADERS)}">{col_items}</colItems>'
        f"{page_block}"
        f'<dataFields count="{len(data_xml)}">{"".join(data_xml)}</dataFields>'
        f"{_contract_caption_filters_xml(contract)}"
        '<pivotTableStyleInfo name="PivotStyleLight16" showRowHeaders="1" '
        'showColHeaders="1" showRowStripes="0" showColStripes="0" showLastColumn="1"/>'
        "</pivotTableDefinition>"
    )


def _excel_safe_pivot_xml(xml: str) -> str:
    """Re-emit PivotTable XML in the attribute/child order desktop Excel writes.

    Hand-authored compact parts are valid OOXML but Excel 16 refuses to open
    them (0x800A03EC). openpyxl's TableDefinition serializer matches the
    verbose form Excel accepts, including r:id and rowItems-before-colFields.
    """
    from openpyxl.pivot.table import TableDefinition
    from openpyxl.xml.functions import fromstring, tostring

    table = TableDefinition.from_tree(fromstring(xml.encode("utf-8")))
    table.id = "rId1"
    payload = tostring(table.to_tree())
    if isinstance(payload, bytes):
        text = payload.decode("utf-8")
    else:
        text = payload
    return re.sub(
        r"<rowItems count=\"1\"><i t=\"data\"[^>]*>.*?</i></rowItems>",
        '<rowItems count="1"><i t="grand"><x/></i></rowItems>',
        text,
        count=1,
        flags=re.DOTALL,
    )


def _page_field_tag(header: str, selected: str | None) -> str:
    fld = cache_field_index(header)
    # Filter Logic 1 uniques grow as months accumulate. A numeric pageField
    # item index is rewritten by Excel after Refresh All and empties Service
    # Campaigns (captionEqual + the wrong item → zero rows). Pin by caption
    # / slicer instead. Closed Group1–Group7 lists stay index-stable.
    if header == "Filter Logic 1":
        return f'<pageField fld="{fld}" hier="-1"/>'
    values = _page_item_values(header)
    if selected and selected in values:
        return f'<pageField fld="{fld}" item="{values.index(selected)}" hier="-1"/>'
    return f'<pageField fld="{fld}" hier="-1"/>'


def _items_for_values(values: Sequence[str], selected: str | None) -> str:
    if not values or selected is None:
        return '<items count="1"><item t="default"/></items>'
    body = "".join(f'<item x="{index}"/>' for index in range(len(values)))
    return f'<items count="{len(values) + 1}">{body}<item t="default"/></items>'


def _nth_pivot_field(xml: str, fld: int) -> re.Match[str]:
    matches = list(
        re.finditer(r"<pivotField\b[^>]*/>|<pivotField\b[^>]*>.*?</pivotField>", xml, re.DOTALL)
    )
    if fld >= len(matches):
        raise ValueError(f"pivotField {fld} missing")
    return matches[fld]


def _pin_one_page_field(xml: str, header: str, selected: str | None, values: Sequence[str]) -> str:
    fld = cache_field_index(header)
    if header == "Filter Logic 1":
        selected = None
        values = ()
    if selected and selected in values:
        tag = f'<pageField fld="{fld}" item="{values.index(selected)}" hier="-1"/>'
    else:
        tag = f'<pageField fld="{fld}" hier="-1"/>'
    xml, n = re.subn(rf'<pageField fld="{fld}"[^/]*/>', tag, xml, count=1)
    if n != 1:
        return xml
    match = _nth_pivot_field(xml, fld)
    old = match.group(0)
    if "<items" not in old:
        return xml
    new, replaced = re.subn(
        r"<items[^>]*>.*?</items>",
        _items_for_values(values, selected),
        old,
        count=1,
        flags=re.DOTALL,
    )
    if replaced != 1:
        return xml
    return xml[: match.start()] + new + xml[match.end() :]


def _drop_unmatched_caption_filters(
    xml: str,
    *,
    fl1_values: Sequence[str],
    group_values: Sequence[str],
) -> str:
    """Remove captionEqual pins whose value is not in the snapshot.

    A phantom Service Filter Logic 1 pin empties that sheet when the tenant
    has no matching campaign, even after cumulative history is in the cache.
    """
    fl1_fld = str(cache_field_index("Filter Logic 1"))
    group_fld = str(cache_field_index("Filter Logic 1_2"))

    def _keep(match: re.Match[str]) -> str:
        block = match.group(0)
        fld = re.search(r'fld="(\d+)"', block)
        val = re.search(r'<filter val="([^"]*)"/>', block)
        if fld is None or val is None:
            return block
        value = (
            val.group(1)
            .replace("&amp;", "&")
            .replace("&lt;", "<")
            .replace("&gt;", ">")
            .replace("&quot;", '"')
        )
        if fld.group(1) == fl1_fld and value not in fl1_values:
            return ""
        if fld.group(1) == group_fld and value not in group_values:
            return ""
        return block

    xml = re.sub(
        r'<filter fld="\d+" type="captionEqual"[^>]*>.*?</filter>',
        _keep,
        xml,
        flags=re.DOTALL,
    )
    remaining = len(re.findall(r'type="captionEqual"', xml))
    xml = re.sub(r'<filters count="\d+">', f'<filters count="{remaining}">', xml, count=1)
    if remaining == 0:
        xml = re.sub(r"<filters count=\"0\"></filters>", "", xml, count=1)
    return xml


def _pin_page_field_state(
    xml: str,
    contract: ReportSheetContract,
    group_values: Sequence[str],
    fl1_values: Sequence[str],
) -> str:
    """Re-apply pageField item indexes after openpyxl re-serialization."""
    for header, selected in _page_fields(contract):
        if header == "Filter Logic 2":
            xml = _pin_one_page_field(xml, header, None, ())
        elif header == "Filter Logic 1":
            xml = _pin_one_page_field(xml, header, None, ())
        else:
            xml = _pin_one_page_field(xml, header, selected, group_values)
    if contract.name == SERVICE:
        xml = _strip_service_fl1_page_field(xml)
    return _drop_unmatched_caption_filters(xml, fl1_values=fl1_values, group_values=group_values)


def _strip_service_fl1_page_field(xml: str) -> str:
    """Remove the Service Filter Logic 1 page field so Excel cannot remap item=0."""
    fld = cache_field_index("Filter Logic 1")
    xml, _n = re.subn(rf'<pageField fld="{fld}"[^/]*/>', "", xml)
    remaining = len(re.findall(r"<pageField ", xml))
    xml = re.sub(r'<pageFields count="\d+">', f'<pageFields count="{remaining}">', xml, count=1)
    if remaining == 0:
        xml = re.sub(r"<pageFields count=\"0\"></pageFields>", "", xml, count=1)
    match = _nth_pivot_field(xml, fld)
    old = match.group(0)
    new = old.replace(' axis="axisPage"', "")
    new, _r = re.subn(
        r"<items[^>]*>.*?</items>",
        '<items count="1"><item t="default"/></items>',
        new,
        count=1,
        flags=re.DOTALL,
    )
    return xml[: match.start()] + new + xml[match.end() :]


def _replace_cache_shared_items(cache_xml: str, field_name: str, values: Sequence[str]) -> str:
    shared = _string_shared_items_xml(values)
    pattern = (
        rf'(<cacheField name="{re.escape(field_name)}"[^>]*>)\s*'
        r"(?:<sharedItems/>|<sharedItems\b.*?</sharedItems>)\s*(</cacheField>)"
    )
    xml, n = re.subn(pattern, rf"\1{shared}\2", cache_xml, count=1, flags=re.DOTALL)
    if n != 1:
        raise ValueError(f"cache field {field_name} sharedItems not replaced")
    return xml


def _first_seen(values: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    out: list[str] = []
    for value in values:
        if not value or value in seen:
            continue
        seen.add(value)
        out.append(value)
    return out


def _row_text(row: Mapping[str, object], *keys: str) -> str:
    for key in keys:
        value = row.get(key)
        if value is not None and str(value) != "":
            return str(value)
    return ""


def page_filter_uniques_from_rows(
    rows: Sequence[Mapping[str, object]],
) -> tuple[list[str], list[str]]:
    """First-seen page-filter values, with report defaults present for indexing."""
    groups = _first_seen(_row_text(row, "filter_logic_1_group", "Filter Logic 1_2") for row in rows)
    fl1 = _first_seen(_row_text(row, "filter_logic_1", "Filter Logic 1") for row in rows)
    for default in (AMC_GROUP, D2C_GROUP, "Group2"):
        if default not in groups:
            groups.append(default)
    return groups, fl1


def apply_page_filter_defaults_into(
    parts: dict[str, bytes], rows: Sequence[Mapping[str, object]]
) -> None:
    """Pin pageField item indexes onto an in-memory package."""
    if not rows:
        return
    groups, fl1 = page_filter_uniques_from_rows(rows)
    workbook_xml = parts["xl/workbook.xml"].decode("utf-8")
    rels_xml = parts["xl/_rels/workbook.xml.rels"].decode("utf-8")
    bindings = bindings_from_workbook(workbook_xml, rels_xml)
    cache = _replace_cache_shared_items(
        parts[PIVOT_CACHE_PART].decode("utf-8"), "filter_logic_1_group", groups
    )
    cache = _replace_cache_shared_items(cache, "Filter Logic 1", fl1)
    parts[PIVOT_CACHE_PART] = cache.encode("utf-8")
    for binding in bindings:
        xml = _pin_page_field_state(
            parts[binding.pivot_part].decode("utf-8"),
            binding.contract,
            groups,
            fl1,
        )
        parts[binding.pivot_part] = ensure_measure_data_fields_xml(xml).encode("utf-8")


def apply_page_filter_defaults_xml(body: bytes, rows: Sequence[Mapping[str, object]]) -> bytes:
    """Pin pageField item indexes to first-seen snapshot values (no stale blanks)."""
    if not rows:
        return body
    return mutate_xlsx(body, lambda parts: apply_page_filter_defaults_into(parts, rows))


# Compact layout: column B holds every row-field caption; C:L are the 10 sums.
# Click-Through O:S and Overall T:X sit outside the PivotTable so Refresh All
# cannot drop them. Identity is the same compact row as C:L (Month/Day context
# of that outline row). SEQUENCE spill does not survive PivotTable row rewrite;
# stored fill-down formulas on each row do.
CONVERSION_STYLE_MARK = "<!--DFIP-CONV-XF:"
CONVERSION_PIVOT_LOCATION = "B9:L10"
CONVERSION_VALUE_ROW = 10
CONVERSION_FILL_ROWS = 12000
CONVERSION_LAST_FORMULA_ROW = CONVERSION_VALUE_ROW + CONVERSION_FILL_ROWS - 1
# (column, header, style key, formula builder name)
CONVERSION_SECTION_FIELDS: tuple[tuple[str, str, str, str], ...] = (
    ("O", "Unique Click-Through Conversions", "int", "uct"),
    ("P", "Click-Through Revenue (INR)", "int", "ctc_rev"),
    ("Q", "Cost/ UCT conversion", "int", "cost_uct"),
    ("R", "UCT conversion rate", "pct1", "uct_rate"),
    ("S", "Unique Click Through Conv ROAS", "dec2", "ctc_roas"),
    ("T", "Unique Conversions", "int", "conv"),
    ("U", "Revenue (INR)", "int", "rev"),
    ("V", "Cost/Unique Conversion", "int", "cost_conv"),
    ("W", "Delivered Thru Conv. rate", "pct2", "del_rate"),
    ("X", "Overall ROAS", "roas", "overall_roas"),
)
_CONVERSION_NUMFMTS: tuple[tuple[str, int], ...] = (
    ("int", 3),
    ("dec2", 4),
    ("pct1", 164),
    ("pct2", 10),
    ("roas", 2),
)


def _numeric_copy(col: str) -> str:
    """Same-row copy of a stable pivot value; skip the header caption row."""
    return f'IF(ISNUMBER({col}10),{col}10,"")'


def _ratio_formula(numerator: str, denominator: str) -> str:
    den = f"{denominator}10"
    num = f"{numerator}10"
    return f'IF(NOT(ISNUMBER({den})),"",IF({den}=0,"",{num}/{den}))'


def _conversion_formula(kind: str) -> str:
    if kind == "uct":
        return _numeric_copy("I")
    if kind == "ctc_rev":
        return _numeric_copy("J")
    if kind == "cost_uct":
        return _ratio_formula("C", "I")
    if kind == "uct_rate":
        return _ratio_formula("I", "H")
    if kind == "ctc_roas":
        return _ratio_formula("J", "C")
    if kind == "conv":
        return _numeric_copy("K")
    if kind == "rev":
        return _numeric_copy("L")
    if kind == "cost_conv":
        return _ratio_formula("C", "K")
    if kind == "del_rate":
        return _ratio_formula("K", "F")
    if kind == "overall_roas":
        return _ratio_formula("L", "C")
    raise ValueError(f"unknown conversion formula {kind}")


def _append_conversion_number_styles(styles_xml: str) -> tuple[str, dict[str, int]]:
    """Append worksheet XFs for O:X so formats survive Refresh All (not General)."""
    marked = re.search(rf"{re.escape(CONVERSION_STYLE_MARK)}(\d+)-->", styles_xml)
    if marked:
        start = int(marked.group(1))
        return styles_xml, {
            name: start + index for index, (name, _fmt) in enumerate(_CONVERSION_NUMFMTS)
        }
    count_match = re.search(r'<cellXfs count="(\d+)"', styles_xml)
    if count_match is None:
        return styles_xml, {}
    start = int(count_match.group(1))
    extras = "".join(
        f'<xf numFmtId="{fmt}" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/>'
        for _name, fmt in _CONVERSION_NUMFMTS
    )
    styles_xml = styles_xml.replace(
        "<cellXfs ",
        f"{CONVERSION_STYLE_MARK}{start}--><cellXfs ",
        1,
    )
    styles_xml = re.sub(
        r'<cellXfs count="\d+"',
        f'<cellXfs count="{start + len(_CONVERSION_NUMFMTS)}"',
        styles_xml,
        count=1,
    )
    styles_xml = styles_xml.replace("</cellXfs>", extras + "</cellXfs>", 1)
    return styles_xml, {
        name: start + index for index, (name, _fmt) in enumerate(_CONVERSION_NUMFMTS)
    }


def _conversion_header_cells() -> str:
    return "".join(
        _inline(f"{col}9", measure_datafield_display_name(header))
        for col, header, _style, _kind in CONVERSION_SECTION_FIELDS
    )


def _style_attr(style: int | None) -> str:
    return "" if style is None else f' s="{style}"'


def _shared_master_cell(col: str, formula: str, si: int, style: int | None) -> str:
    last = CONVERSION_LAST_FORMULA_ROW
    start = CONVERSION_VALUE_ROW
    return (
        f'<c r="{col}{start}"{_style_attr(style)}>'
        f'<f t="shared" ref="{col}{start}:{col}{last}" si="{si}">'
        f"{_xml_text(formula)}</f></c>"
    )


def _shared_slot_cell(col: str, row: int, si: int, style: int | None) -> str:
    return f'<c r="{col}{row}"{_style_attr(style)}><f t="shared" si="{si}"/></c>'


def _conversion_fill_rows_xml(styles: dict[str, int]) -> str:
    """Stored same-row formulas for O:X. Shared formulas keep the package small."""
    last = CONVERSION_LAST_FORMULA_ROW
    rows = []
    master = "".join(
        _shared_master_cell(col, _conversion_formula(kind), index, styles.get(style_key))
        for index, (col, _header, style_key, kind) in enumerate(CONVERSION_SECTION_FIELDS)
    )
    rows.append(f'<row r="{CONVERSION_VALUE_ROW}" spans="15:24">{master}</row>')
    for row in range(CONVERSION_VALUE_ROW + 1, last + 1):
        cells = "".join(
            _shared_slot_cell(col, row, index, styles.get(style_key))
            for index, (col, _header, style_key, _kind) in enumerate(CONVERSION_SECTION_FIELDS)
        )
        rows.append(f'<row r="{row}" spans="15:24">{cells}</row>')
    return "".join(rows)


def _clear_non_sum_datafield_flags(xml: str) -> str:
    """Excel rejects Open when dataField='1' count != dataFields count."""
    match = re.search(r"<pivotFields\b[^>]*>.*?</pivotFields>", xml, re.DOTALL)
    if match is None:
        return xml
    keep = {cache_field_index(header) for header in sum_measure_headers()}
    block = match.group(0)
    fields = list(
        re.finditer(r"<pivotField\b[^>]*/>|<pivotField\b[^>]*>.*?</pivotField>", block, re.DOTALL)
    )
    pieces: list[str] = []
    last = 0
    for index, field in enumerate(fields):
        pieces.append(block[last : field.start()])
        tag = field.group(0)
        if index not in keep:
            tag = re.sub(r'\s*dataField="1"', "", tag, count=1)
        pieces.append(tag)
        last = field.end()
    pieces.append(block[last:])
    return xml[: match.start()] + "".join(pieces) + xml[match.end() :]


def _bind_refreshable_sum_pivot(xml: str) -> str:
    xml, replaced = re.subn(
        r"<dataFields\b.*?</dataFields>",
        sum_measure_data_fields_xml(),
        xml,
        count=1,
        flags=re.DOTALL,
    )
    if replaced != 1:
        return xml
    xml = _clear_non_sum_datafield_flags(xml)
    xml, _n = re.subn(
        r"<colItems\b.*?</colItems>",
        sum_measure_col_items_xml(),
        xml,
        count=1,
        flags=re.DOTALL,
    )
    xml = re.sub(
        r'<location ref="B9:[A-Z]+\d+"',
        f'<location ref="{CONVERSION_PIVOT_LOCATION}"',
        xml,
        count=1,
    )
    return xml


def _inject_row_cells(sheet_xml: str, row: int, cells: str, spans: str) -> str:
    self_close = re.search(rf'<row r="{row}"([^>]*)/>', sheet_xml)
    if self_close:
        attrs = self_close.group(1)
        if "spans=" not in attrs:
            attrs += f' spans="{spans}"'
        replacement = f'<row r="{row}"{attrs}>{cells}</row>'
        return sheet_xml[: self_close.start()] + replacement + sheet_xml[self_close.end() :]
    open_row = re.search(rf'<row r="{row}"[^>]*>', sheet_xml)
    if open_row:
        return sheet_xml[: open_row.end()] + cells + sheet_xml[open_row.end() :]
    return sheet_xml.replace(
        "</sheetData>",
        f'<row r="{row}" spans="{spans}">{cells}</row></sheetData>',
        1,
    )


def _inject_conversion_formulas(
    sheet_xml: str, styles: dict[str, int], fill: str | None = None
) -> str:
    if 't="shared"' in sheet_xml and "ISNUMBER(I10)" in sheet_xml:
        return sheet_xml
    sheet_xml = _inject_row_cells(sheet_xml, 9, _conversion_header_cells(), "15:24")
    fill_xml = fill if fill is not None else _conversion_fill_rows_xml(styles)
    if re.search(rf'<row r="{CONVERSION_VALUE_ROW}"', sheet_xml):
        sheet_xml = re.sub(
            rf'<row r="{CONVERSION_VALUE_ROW}"[^>]*>.*?</row>',
            "",
            sheet_xml,
            count=1,
            flags=re.DOTALL,
        )
    return sheet_xml.replace("</sheetData>", fill_xml + "</sheetData>", 1)


def apply_refreshable_conversion_layout_into(parts: dict[str, bytes]) -> None:
    """Keep conversion sections populated after Power Query Refresh All."""
    workbook_xml = parts["xl/workbook.xml"].decode("utf-8")
    rels_xml = parts["xl/_rels/workbook.xml.rels"].decode("utf-8")
    bindings = bindings_from_workbook(workbook_xml, rels_xml)
    styles_xml, conv_styles = _append_conversion_number_styles(
        parts["xl/styles.xml"].decode("utf-8")
    )
    parts["xl/styles.xml"] = styles_xml.encode("utf-8")
    fill = _conversion_fill_rows_xml(conv_styles)
    for binding in bindings:
        parts[binding.pivot_part] = _bind_refreshable_sum_pivot(
            parts[binding.pivot_part].decode("utf-8")
        ).encode("utf-8")
        parts[binding.sheet_part] = _inject_conversion_formulas(
            parts[binding.sheet_part].decode("utf-8"),
            conv_styles,
            fill,
        ).encode("utf-8")


def apply_refreshable_conversion_layout_xml(body: bytes) -> bytes:
    """Keep conversion sections populated after Power Query Refresh All.

    Calculated cache fields are dropped from DataFields when the shared cache
    rebuilds from ExternalData_1. Shrink each PivotTable to the 10 additive
    sums (C:L) and place Click-Through / Overall formulas in O:X.
    """
    return mutate_xlsx(body, apply_refreshable_conversion_layout_into)


def apply_page_filter_defaults_excel(path: str | Path) -> None:
    """Set native CurrentPage after cache population so the selection survives refresh.

    Excel rebuilds page items on the first populated refresh and resets dropdowns
    to (All) unless CurrentPage is applied on the live PivotField. Calculated
    fields are re-added because that same rebuild drops them from DataFields.
    """
    import time

    win32com_client = __import__("win32com.client", fromlist=["DispatchEx"])
    target = Path(path).resolve()
    excel = win32com_client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    excel.AskToUpdateLinks = False
    excel.EnableEvents = False
    workbook = None
    try:
        workbook = excel.Workbooks.Open(str(target), UpdateLinks=0, ReadOnly=False)
        workbook.Worksheets(REPORT_SHEET_NAMES[0]).PivotTables(1).PivotCache().Refresh()
        restore_report_pivots_excel(workbook)
        for contract in REPORT_CONTRACTS:
            pivot = workbook.Worksheets(contract.name).PivotTables(1)
            if contract.default_filter_logic_1 and pivot.PageFields.Count >= 1:
                try:
                    pivot.PageFields.Item(1).CurrentPage = contract.default_filter_logic_1
                except Exception:
                    pass
            elif contract.default_group and pivot.PageFields.Count >= 1:
                try:
                    pivot.PageFields.Item(1).CurrentPage = contract.default_group
                except Exception:
                    pass
        workbook.Save()
        workbook.Close(False)
        workbook = None
    finally:
        if workbook is not None:
            try:
                workbook.Close(False)
            except Exception:
                pass
        try:
            excel.Quit()
        except Exception:
            pass
        time.sleep(0.4)


def _caption_filter_xml(field_header: str, value: str, filter_id: int) -> str:
    fld = cache_field_index(field_header)
    return (
        f'<filter fld="{fld}" type="captionEqual" evalOrder="0" id="{filter_id}">'
        '<autoFilter ref="A1"><filterColumn colId="0"><filters>'
        f'<filter val="{_xml_attr(value)}"/>'
        "</filters></filterColumn></autoFilter></filter>"
    )


def _contract_caption_filters_xml(contract: ReportSheetContract) -> str:
    """Label filters for P10 page defaults. No shared-item indexes (crash-safe)."""
    parts: list[str] = []
    filter_id = 1
    if contract.default_filter_logic_1:
        parts.append(
            _caption_filter_xml("Filter Logic 1", contract.default_filter_logic_1, filter_id)
        )
        filter_id += 1
    if contract.default_group:
        parts.append(_caption_filter_xml("Filter Logic 1_2", contract.default_group, filter_id))
        filter_id += 1
    if not parts:
        return ""
    return f'<filters count="{len(parts)}">{"".join(parts)}</filters>'


def _slicer_cache_name(binding: ReportPivotBinding, field: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "_", field).strip("_")
    return f"Slicer_{slug}_{binding.pivot_name}"


def _slicer_object_name(binding: ReportPivotBinding, field: str) -> str:
    return f"{field} {binding.pivot_name}"


def _slicer_source_name(field: str) -> str:
    return PUBLISHED_HEADER.get(field, field)


def _slicer_caption(field: str) -> str:
    if field == "Filter Logic 1_2":
        return "Filter Logic 1_2"
    return field


def slicer_cache_xml(binding: ReportPivotBinding, field: str) -> str:
    name = _slicer_cache_name(binding, field)
    source = _slicer_source_name(field)
    uid = _uid(f"slicer-cache-{name}")
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
        f'<slicerCacheDefinition xmlns="{NS_SLICER_MAIN}" '
        f'xmlns:mc="{NS_MC}" mc:Ignorable="x xr10" xmlns:x="{NS_MAIN}" '
        'xmlns:xr10="http://schemas.microsoft.com/office/spreadsheetml/2016/revision10" '
        f'name="{_xml_attr(name)}" xr10:uid="{{{uid}}}" sourceName="{_xml_attr(source)}">'
        "<pivotTables>"
        f'<pivotTable tabId="{binding.sheet_id}" name="{binding.pivot_name}"/>'
        "</pivotTables>"
        f'<data><tabular pivotCacheId="{SLICER_PIVOT_CACHE_ID}"/>'
        "</data>"
        '<extLst><x:ext uri="{470722E0-AACD-4C17-9CDC-17EF765DBC7E}" '
        'xmlns:x15="http://schemas.microsoft.com/office/spreadsheetml/2010/11/main">'
        "<x15:slicerCacheHideItemsWithNoData/></x:ext></extLst>"
        "</slicerCacheDefinition>"
    )


def slicers_xml(binding: ReportPivotBinding) -> str:
    parts = []
    for field in binding.contract.slicers:
        name = _slicer_object_name(binding, field)
        cache = _slicer_cache_name(binding, field)
        caption = _slicer_caption(field)
        uid = _uid(f"slicer-{name}")
        cols = ' columnCount="2"' if field == "Filter Logic 1" else ""
        parts.append(
            f'<slicer name="{_xml_attr(name)}" xr10:uid="{{{uid}}}" '
            f'cache="{_xml_attr(cache)}" caption="{_xml_attr(caption)}"'
            f'{cols} rowHeight="241300"/>'
        )
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
        f'<slicers xmlns="{NS_SLICER_MAIN}" xmlns:mc="{NS_MC}" '
        f'mc:Ignorable="x xr10" xmlns:x="{NS_MAIN}" '
        'xmlns:xr10="http://schemas.microsoft.com/office/spreadsheetml/2016/revision10">'
        f"{''.join(parts)}</slicers>"
    )


# FY-2026 drawing twoCellAnchor boxes per report sheet, in slicer-field order.
# (from_col, from_row, to_col, to_row, from_colOff, from_rowOff, to_colOff, to_rowOff)
# Slicers sit to the right of the pivot (AA+). Do not park them in C2:X7.
SlicerBox = tuple[int, int, int, int, int, int, int, int]
REFERENCE_SLICER_GEOMETRY: dict[str, tuple[SlicerBox, ...]] = {
    OVERALL: (
        (30, 3, 41, 25, 49347, 43360, 49114, 159920),
        (26, 2, 29, 13, 380633, 185536, 397360, 95249),
        (26, 14, 29, 28, 304800, 133350, 485775, 133351),
    ),
    VERTICAL: (
        (30, 2, 41, 25, 49347, 157660, 49115, 83720),
        (26, 2, 29, 13, 542558, 185536, 559285, 95249),
        (33, 27, 36, 40, 235117, 128336, 221582, 170447),
        (30, 27, 33, 41, 172451, 165935, 158915, 25067),
        (26, 15, 29, 31, 323850, 28574, 504825, 99059),
    ),
    CHANNEL: (
        (30, 3, 41, 28, 49347, 62410, 49115, 133350),
        (26, 3, 29, 13, 352058, 4561, 368785, 104774),
        (26, 15, 29, 29, 323850, 28575, 504825, 28575),
    ),
    SUB_SPLIT: (
        (30, 3, 41, 28, 49347, 62410, 49115, 133350),
        (26, 3, 29, 13, 352058, 4561, 368785, 104774),
        (26, 15, 29, 29, 323850, 28575, 504825, 28575),
        (26, 30, 29, 44, 230505, 41910, 230505, 41910),
    ),
    AMC_DAYWISE: (
        (30, 3, 41, 25, 49347, 43360, 49115, 159920),
        (26, 2, 29, 13, 380633, 185536, 397358, 95249),
        (26, 14, 29, 28, 304800, 133350, 485773, 133351),
        (26, 29, 29, 43, 386862, 129149, 193429, 134278),
    ),
    AMC_SPLIT: (
        (30, 3, 41, 25, 49347, 43360, 49114, 159920),
        (26, 2, 29, 13, 380633, 185536, 397360, 95249),
        (26, 14, 29, 28, 304800, 133350, 485775, 133349),
        (27, 30, 30, 43, 145171, 31457, 156894, 164785),
    ),
    AMC_VERTICAL: (
        (30, 3, 41, 25, 49347, 43360, 49114, 145266),
        (26, 2, 29, 13, 380633, 185536, 397360, 95249),
        (26, 14, 29, 29, 304800, 133349, 485775, 132550),
        (26, 29, 29, 43, 370308, 170003, 382031, 135345),
    ),
    D2C: (
        (30, 3, 41, 25, 49347, 43360, 49114, 159920),
        (26, 2, 29, 13, 380633, 185536, 397360, 95249),
        (26, 13, 29, 29, 304800, 178172, 485775, 6330),
        (26, 29, 29, 43, 370308, 170003, 382031, 122538),
    ),
    SERVICE: (
        (31, 3, 42, 28, 49347, 62410, 49115, 144847),
        (27, 3, 30, 13, 352058, 4561, 368785, 104774),
        (27, 15, 30, 29, 323850, 28575, 504825, 40072),
        (27, 29, 30, 43, 428625, 95250, 428625, 81882),
    ),
}


def slicer_anchor_box(sheet_name: str, index: int) -> SlicerBox:
    """Return FY-2026 two-cell anchor box for one slicer on a report sheet."""
    boxes = REFERENCE_SLICER_GEOMETRY.get(sheet_name)
    if not boxes or index < 0 or index >= len(boxes):
        raise ValueError("slicer layout exceeds the reference drawing.")
    return boxes[index]


def _slicer_anchor(index: int, name: str, shape_id: int, sheet_name: str) -> str:
    col, row, to_col, to_row, from_col_off, from_row_off, to_col_off, to_row_off = (
        slicer_anchor_box(sheet_name, index)
    )
    uid = _uid(f"drawing-{name}")
    return (
        '<xdr:twoCellAnchor editAs="oneCell">'
        f"<xdr:from><xdr:col>{col}</xdr:col><xdr:colOff>{from_col_off}</xdr:colOff>"
        f"<xdr:row>{row}</xdr:row><xdr:rowOff>{from_row_off}</xdr:rowOff></xdr:from>"
        f"<xdr:to><xdr:col>{to_col}</xdr:col><xdr:colOff>{to_col_off}</xdr:colOff>"
        f"<xdr:row>{to_row}</xdr:row><xdr:rowOff>{to_row_off}</xdr:rowOff></xdr:to>"
        f'<mc:AlternateContent xmlns:mc="{NS_MC}" '
        'xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main">'
        '<mc:Choice Requires="a14">'
        '<xdr:graphicFrame macro="">'
        "<xdr:nvGraphicFramePr>"
        f'<xdr:cNvPr id="{shape_id}" name="{_xml_attr(name)}">'
        '<a:extLst><a:ext uri="{FF2B5EF4-FFF2-40B4-BE49-F238E27FC236}">'
        f'<a16:creationId xmlns:a16="http://schemas.microsoft.com/office/drawing/2014/main" '
        f'id="{{{uid}}}"/></a:ext></a:extLst></xdr:cNvPr>'
        "<xdr:cNvGraphicFramePr/></xdr:nvGraphicFramePr>"
        '<xdr:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/></xdr:xfrm>'
        '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/drawing/2010/slicer">'
        f'<sle:slicer xmlns:sle="http://schemas.microsoft.com/office/drawing/2010/slicer" '
        f'name="{_xml_attr(name)}"/>'
        "</a:graphicData></a:graphic></xdr:graphicFrame></mc:Choice>"
        '<mc:Fallback xmlns="">'
        '<xdr:sp macro="" textlink="">'
        "<xdr:nvSpPr>"
        f'<xdr:cNvPr id="{shape_id + 100}" name="{_xml_attr(name)} Fallback"/>'
        '<xdr:cNvSpPr><a:spLocks noTextEdit="1"/></xdr:cNvSpPr></xdr:nvSpPr>'
        '<xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1000000" cy="1000000"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        '<a:solidFill><a:prstClr val="white"/></a:solidFill></xdr:spPr>'
        '<xdr:txBody><a:bodyPr vertOverflow="clip" horzOverflow="clip"/>'
        '<a:lstStyle/><a:p><a:r><a:rPr lang="en-US" sz="1100"/>'
        f"<a:t>{_xml_text(_SLICER_FALLBACK)}</a:t></a:r></a:p></xdr:txBody>"
        "</xdr:sp></mc:Fallback></mc:AlternateContent><xdr:clientData/></xdr:twoCellAnchor>"
    )


def drawing_xml(binding: ReportPivotBinding) -> str:
    anchors = []
    for index, field in enumerate(binding.contract.slicers):
        name = _slicer_object_name(binding, field)
        anchors.append(_slicer_anchor(index, name, index + 2, binding.contract.name))
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
        f'<xdr:wsDr xmlns:xdr="{NS_XDR}" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        f'xmlns:mc="{NS_MC}">'
        f"{''.join(anchors)}</xdr:wsDr>"
    )


def pivot_sheet_xml(contract: ReportSheetContract) -> str:
    """Report chrome around a compact native PivotTable at B9. No UNIQUE formulas."""
    uct_start = _col(3 + 12)
    uct_end = _col(3 + 16)
    overall_start = _col(3 + 17)
    overall_end = _col(3 + 21)
    last_col_idx = max(24, CAPTION_NAME_COL)
    last_col = _col(last_col_idx)
    span = f"1:{last_col_idx}"
    hint = ZERO_VS_BLANK_HINT_VERTICAL if contract.device_product_filters else ZERO_VS_BLANK_HINT
    empty_formula = f'IF(IFERROR(COUNTA(PublishedFacts!A:A),0)<=1,"{EMPTY_STATE}","")'
    caption_by_row: dict[int, str] = {}
    for index, (key, caption) in enumerate(_caption_pairs(), start=1):
        caption_by_row[index] = _inline(f"{_col(CAPTION_KEY_COL)}{index}", key) + _inline(
            f"{_col(CAPTION_NAME_COL)}{index}", caption
        )

    def _row(number: int, inner: str, *, height: float | None = None) -> str:
        ht = "" if height is None else f' ht="{height}" customHeight="1"'
        extra = caption_by_row.get(number, "")
        return f'<row r="{number}" spans="{span}"{ht}>{inner}{extra}</row>'

    rows = [
        _row(1, _inline("A1", "Web Engage Daily Report", style=XF_TITLE), height=22),
        _row(2, _inline("A2", contract.purpose, style=XF_PURPOSE)),
        _row(3, ""),
        _row(4, ""),
        _row(5, _inline("A5", hint, style=XF_HINT), height=32),
        _row(6, _formula_cell("A6", empty_formula, XF_HINT)),
        _row(7, ""),
        _row(
            8,
            (
                f"{_inline(uct_start + '8', 'Click-Through Conversions', style=XF_GROUP_BANNER)}"
                f"{_inline(overall_start + '8', 'Overall Conversions', style=XF_GROUP_BANNER)}"
            ),
        ),
    ]
    merges = (
        f'<mergeCells count="2">'
        f'<mergeCell ref="{uct_start}8:{uct_end}8"/>'
        f'<mergeCell ref="{overall_start}8:{overall_end}8"/>'
        f"</mergeCells>"
    )
    uid = _uid(f"sheet-{contract.name}")
    cols = (
        "<cols>"
        '<col min="1" max="1" width="12" customWidth="1"/>'
        '<col min="2" max="2" width="28" customWidth="1"/>'
        '<col min="3" max="24" width="13" customWidth="1"/>'
        f'<col min="25" max="{CAPTION_KEY_COL - 1}" width="0" hidden="1" customWidth="1"/>'
        f'<col min="{CAPTION_KEY_COL}" max="{CAPTION_NAME_COL}" '
        'width="0" hidden="1" customWidth="1"/>'
        "</cols>"
    )
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
        f'<worksheet xmlns="{NS_MAIN}" xmlns:r="{NS_REL}" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'mc:Ignorable="x14ac xr xr2 xr3" '
        'xmlns:x14ac="http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac" '
        'xmlns:xr="http://schemas.microsoft.com/office/spreadsheetml/2014/revision" '
        'xmlns:xr2="http://schemas.microsoft.com/office/spreadsheetml/2015/revision2" '
        'xmlns:xr3="http://schemas.microsoft.com/office/spreadsheetml/2016/revision3" '
        f'xr:uid="{{{uid}}}">'
        f'<dimension ref="A1:{last_col}10"/>'
        '<sheetViews><sheetView workbookViewId="0" showGridLines="0" zoomScale="90">'
        '<selection activeCell="B10" sqref="B10"/>'
        "</sheetView></sheetViews>"
        '<sheetFormatPr defaultRowHeight="15" x14ac:dyDescent="0.3"/>'
        f"{cols}"
        f"<sheetData>{''.join(rows)}</sheetData>"
        f"{merges}"
        '<pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75"'
        ' header="0.3" footer="0.3"/>'
        '<drawing r:id="rId3"/>'
        f'<extLst><ext uri="{SHEET_SLICER_URI}" '
        'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
        '<x14:slicerList><x14:slicer r:id="rId2"/></x14:slicerList>'
        "</ext></extLst></worksheet>"
    )


def _formula_cell(ref: str, formula: str, style: int) -> str:
    return f'<c r="{ref}" s="{style}"><f>{_xml_text(formula)}</f></c>'


def worksheet_rels_xml(binding: ReportPivotBinding) -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
        f'<Relationships xmlns="{NS_PKG_REL}">'
        f'<Relationship Id="rId1" Type="{NS_PIVOT_TABLE_REL}" '
        f'Target="../pivotTables/pivotTable{binding.index}.xml"/>'
        f'<Relationship Id="rId2" Type="{NS_SLICER_REL}" '
        f'Target="../slicers/slicer{binding.index}.xml"/>'
        f'<Relationship Id="rId3" Type="{NS_DRAWING_REL}" '
        f'Target="../drawings/drawing{binding.index}.xml"/>'
        "</Relationships>"
    )


def pivot_table_rels_xml() -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
        f'<Relationships xmlns="{NS_PKG_REL}">'
        f'<Relationship Id="rId1" Type="{NS_PIVOT_CACHE_REL}" '
        'Target="../pivotCache/pivotCacheDefinition1.xml"/>'
        "</Relationships>"
    )


def pivot_cache_rels_xml() -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
        f'<Relationships xmlns="{NS_PKG_REL}">'
        f'<Relationship Id="rId1" Type="{NS_PIVOT_RECORDS_REL}" '
        'Target="pivotCacheRecords1.xml"/>'
        "</Relationships>"
    )


def _sheet_rels_name(part: str) -> str:
    parent, name = part.rsplit("/", 1)
    return f"{parent}/_rels/{name}.rels"


def _sheet_id_map(workbook_xml: str) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for tag in re.findall(r"<sheet\b[^/]*/>", workbook_xml):
        name = re.search(r'name="([^"]*)"', tag)
        sheet_id = re.search(r'sheetId="(\d+)"', tag)
        if name and sheet_id:
            mapping[name.group(1).replace("&amp;", "&").replace("&quot;", '"')] = int(
                sheet_id.group(1)
            )
    return mapping


def bindings_from_workbook(workbook_xml: str, rels_xml: str) -> list[ReportPivotBinding]:
    parts = _sheet_part_map(workbook_xml, rels_xml)
    ids = _sheet_id_map(workbook_xml)
    missing = [name for name in REPORT_SHEET_NAMES if name not in parts or name not in ids]
    if missing:
        raise ValueError(f"report sheets missing from workbook: {missing}")
    return [
        ReportPivotBinding(contract, parts[contract.name], ids[contract.name], index)
        for index, contract in enumerate(REPORT_CONTRACTS, start=1)
    ]


def _slicer_plan(
    bindings: Iterable[ReportPivotBinding],
) -> list[tuple[ReportPivotBinding, str, int]]:
    plan: list[tuple[ReportPivotBinding, str, int]] = []
    cache_n = 1
    for binding in bindings:
        for field in binding.contract.slicers:
            plan.append((binding, field, cache_n))
            cache_n += 1
    return plan


def build_pivot_parts(bindings: list[ReportPivotBinding]) -> dict[str, bytes]:
    files: dict[str, bytes] = {
        PIVOT_CACHE_PART: pivot_cache_definition_xml().encode("utf-8"),
        PIVOT_RECORDS_PART: pivot_cache_records_xml().encode("utf-8"),
        "xl/pivotCache/_rels/pivotCacheDefinition1.xml.rels": pivot_cache_rels_xml().encode(
            "utf-8"
        ),
    }
    for binding in bindings:
        files[binding.sheet_part] = pivot_sheet_xml(binding.contract).encode("utf-8")
        files[binding.pivot_part] = _pin_page_field_state(
            _excel_safe_pivot_xml(pivot_table_xml(binding)),
            binding.contract,
            PAGE_GROUP_ITEMS,
            (SERVICE_FILTER_LOGIC_1,),
        ).encode("utf-8")
        files[f"xl/pivotTables/_rels/pivotTable{binding.index}.xml.rels"] = (
            pivot_table_rels_xml().encode("utf-8")
        )
        files[binding.slicer_part] = slicers_xml(binding).encode("utf-8")
        files[binding.drawing_part] = drawing_xml(binding).encode("utf-8")
        files[_sheet_rels_name(binding.sheet_part)] = worksheet_rels_xml(binding).encode("utf-8")
    for binding, field, cache_n in _slicer_plan(bindings):
        files[f"xl/slicerCaches/slicerCache{cache_n}.xml"] = slicer_cache_xml(
            binding, field
        ).encode("utf-8")
    return files


def _strip_workbook_rels(xml: str) -> str:
    xml = re.sub(
        r"<Relationship\b[^>]*pivotCacheDefinition[^/]*/>",
        "",
        xml,
    )
    xml = re.sub(r"<Relationship\b[^>]*slicerCache[^/]*/>", "", xml)
    xml = re.sub(r"<Relationship\b[^>]*calcChain[^/]*/>", "", xml)
    return xml


def _patch_workbook_rels(xml: str, slicer_cache_count: int) -> tuple[str, str, list[str]]:
    xml = _strip_workbook_rels(xml)
    next_id = _next_relationship_id(xml)
    cache_rid = f"rId{next_id}"
    extras = [
        f'<Relationship Id="{cache_rid}" Type="{NS_PIVOT_CACHE_REL}" '
        'Target="pivotCache/pivotCacheDefinition1.xml"/>'
    ]
    slicer_rids: list[str] = []
    for offset in range(slicer_cache_count):
        rid = f"rId{next_id + 1 + offset}"
        slicer_rids.append(rid)
        extras.append(
            f'<Relationship Id="{rid}" Type="{NS_SLICER_CACHE_REL}" '
            f'Target="slicerCaches/slicerCache{offset + 1}.xml"/>'
        )
    xml = xml.replace("</Relationships>", "".join(extras) + "</Relationships>")
    return xml, cache_rid, slicer_rids


def _patch_workbook_xml(
    xml: str, cache_rid: str, slicer_rids: list[str], bindings: list[ReportPivotBinding]
) -> str:
    xml = re.sub(r"<pivotCaches>.*?</pivotCaches>", "", xml, flags=re.DOTALL)
    xml = re.sub(
        rf'<ext uri="{re.escape(SLICER_CACHES_URI)}".*?</ext>',
        "",
        xml,
        flags=re.DOTALL,
    )
    xml = re.sub(r'<definedName name="Slicer_[^"]*">#N/A</definedName>', "", xml)
    names = []
    seen: set[str] = set()
    for binding in bindings:
        for field in binding.contract.slicers:
            cache_name = _slicer_cache_name(binding, field)
            if cache_name in seen:
                continue
            seen.add(cache_name)
            names.append(f'<definedName name="{_xml_attr(cache_name)}">#N/A</definedName>')
    if "<definedNames>" in xml:
        xml = xml.replace("</definedNames>", "".join(names) + "</definedNames>", 1)
    else:
        xml = xml.replace(
            "<calcPr",
            "<definedNames>" + "".join(names) + "</definedNames><calcPr",
            1,
        )
    caches = (
        f'<pivotCaches><pivotCache cacheId="{PIVOT_CACHE_ID}" r:id="{cache_rid}"/></pivotCaches>'
    )
    if "<pivotCaches" not in xml:
        if "</calcPr>" in xml:
            xml = xml.replace("</calcPr>", "</calcPr>" + caches, 1)
        else:
            xml = xml.replace("<extLst>", caches + "<extLst>", 1)
    slicer_ext = (
        f'<ext uri="{SLICER_CACHES_URI}" '
        'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
        "<x14:slicerCaches>"
        + "".join(f'<x14:slicerCache r:id="{rid}"/>' for rid in slicer_rids)
        + "</x14:slicerCaches></ext>"
    )
    if "<extLst>" in xml:
        xml = xml.replace("<extLst>", "<extLst>" + slicer_ext, 1)
    else:
        xml = xml.replace("</workbook>", f"<extLst>{slicer_ext}</extLst></workbook>", 1)
    return xml


def _strip_content_types(xml: str) -> str:
    patterns = (
        r'<Override PartName="/xl/pivotCache/[^"]+"[^>]*>',
        r'<Override PartName="/xl/pivotTables/[^"]+"[^>]*>',
        r'<Override PartName="/xl/slicers/[^"]+"[^>]*>',
        r'<Override PartName="/xl/slicerCaches/[^"]+"[^>]*>',
        r'<Override PartName="/xl/drawings/drawing\d+\.xml"[^>]*>',
        r'<Override PartName="/xl/calcChain.xml"[^>]*>',
    )
    for pattern in patterns:
        xml = re.sub(pattern, "", xml)
    return xml


def _patch_content_types(
    xml: str, bindings: list[ReportPivotBinding], slicer_cache_count: int
) -> str:
    xml = _strip_content_types(xml)
    extras = [
        f'<Override PartName="/{PIVOT_CACHE_PART}" ContentType="{CT_CACHE}"/>',
        f'<Override PartName="/{PIVOT_RECORDS_PART}" ContentType="{CT_RECORDS}"/>',
    ]
    for binding in bindings:
        extras.append(f'<Override PartName="/{binding.pivot_part}" ContentType="{CT_PIVOT}"/>')
        extras.append(f'<Override PartName="/{binding.slicer_part}" ContentType="{CT_SLICER}"/>')
        extras.append(f'<Override PartName="/{binding.drawing_part}" ContentType="{CT_DRAWING}"/>')
    for index in range(1, slicer_cache_count + 1):
        extras.append(
            f'<Override PartName="/xl/slicerCaches/slicerCache{index}.xml" '
            f'ContentType="{CT_SLICER_CACHE}"/>'
        )
    return xml.replace("</Types>", "".join(extras) + "</Types>")


def _skip_original_part(name: str, generated: dict[str, bytes]) -> bool:
    if name in generated or name == "xl/calcChain.xml":
        return True
    return any(name.startswith(prefix) for prefix in _SKIP_ZIP_PREFIXES)


def apply_native_pivots(raw: bytes) -> bytes:
    """Rewrite report sheets as native PivotTables. Mashup/PublishedFacts stay."""
    with ZipFile(io.BytesIO(raw), "r") as original:
        names = set(original.namelist())
        for fake in FAKE_MASHUP_ZIP_PARTS:
            if fake in names:
                raise ValueError(f"refusing to patch a workbook that contains {fake}")
        workbook_xml = original.read("xl/workbook.xml").decode("utf-8")
        rels_xml = original.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        types_xml = original.read("[Content_Types].xml").decode("utf-8")
        styles_xml = original.read("xl/styles.xml").decode("utf-8")
        bindings = bindings_from_workbook(workbook_xml, rels_xml)
        generated = build_pivot_parts(bindings)
        slicer_count = len(_slicer_plan(bindings))
        rels_xml, cache_rid, slicer_rids = _patch_workbook_rels(rels_xml, slicer_count)
        rels_xml = _rels_with_metadata(rels_xml)
        workbook_xml = _patch_workbook_xml(workbook_xml, cache_rid, slicer_rids, bindings)
        types_xml = _types_with_metadata(_patch_content_types(types_xml, bindings, slicer_count))
        styles_xml = _ensure_report_styles(styles_xml)
        out = io.BytesIO()
        with ZipFile(out, "w") as written:
            for info in original.infolist():
                if _skip_original_part(info.filename, generated):
                    continue
                if info.filename == "xl/workbook.xml":
                    data = workbook_xml.encode("utf-8")
                elif info.filename == "xl/_rels/workbook.xml.rels":
                    data = rels_xml.encode("utf-8")
                elif info.filename == "[Content_Types].xml":
                    data = types_xml.encode("utf-8")
                elif info.filename == "xl/styles.xml":
                    data = styles_xml.encode("utf-8")
                else:
                    data = original.read(info.filename)
                written.writestr(_clone_zipinfo(info), data)
            for name, payload in generated.items():
                written.writestr(_new_sheet_zipinfo(name), payload)
    return apply_reference_formatting(out.getvalue())


def shared_pivot_cache_binding(workbook_xml: str, rels_xml: str) -> tuple[str, str, str]:
    """Return (cache_id, r:id, relationship target) for the one shared PivotCache.

    Excel Save may remap the numeric cacheId (1 → 0 or 12). The contract is
    exactly one workbook PivotCache whose r:id resolves to
    pivotCacheDefinition1.xml — not a specific number.
    """
    block = re.search(r"<pivotCaches>(.*?)</pivotCaches>", workbook_xml, flags=re.DOTALL)
    if block is None:
        raise ValueError("Client report is incomplete.")
    tags = re.findall(r"<pivotCache\b[^/]*/>", block.group(1))
    if len(tags) != 1:
        raise ValueError("Client report is incomplete.")
    cache_id = re.search(r'cacheId="(\d+)"', tags[0])
    rel_id = re.search(r'r:id="([^"]+)"', tags[0])
    if cache_id is None or rel_id is None:
        raise ValueError("Client report is incomplete.")
    rid = rel_id.group(1)
    target = None
    for relationship in re.findall(r"<Relationship\b[^>]*/?>", rels_xml):
        rel = re.search(r'\bId="([^"]+)"', relationship)
        href = re.search(r'\bTarget="([^"]+)"', relationship)
        if rel is not None and href is not None and rel.group(1) == rid:
            target = href.group(1)
            break
    if target is None or not target.replace("\\", "/").endswith(
        "pivotCache/pivotCacheDefinition1.xml"
    ):
        raise ValueError("Client report is incomplete.")
    return cache_id.group(1), rid, target


def _pivot_table_cache_id(table_xml: str) -> str:
    match = re.search(r"<pivotTableDefinition\b[^>]*cacheId=\"(\d+)\"", table_xml)
    if match is None:
        match = re.search(r'cacheId="(\d+)"', table_xml)
    if match is None:
        raise ValueError("Client report is incomplete.")
    return match.group(1)


def assert_published_facts_mashup(body: bytes) -> str:
    """Return DataMashup Section1.m. Raises if the native package is missing."""
    with ZipFile(io.BytesIO(body), "r") as archive:
        names = set(archive.namelist())
        if "xl/connections.xml" not in names:
            raise ValueError("Client report is incomplete.")
        connections = archive.read("xl/connections.xml").decode("utf-8")
        if "Query - PublishedFacts" not in connections:
            raise ValueError("Client report is incomplete.")
        try:
            section = extract_published_facts_section_from_package(archive)
        except ValueError as exc:
            raise ValueError("Client report is incomplete.") from exc
    if "shared PublishedFacts" not in section and "PublishedFacts =" not in section:
        raise ValueError("Client report is incomplete.")
    return section


def mashup_is_renewal_enabled(section_m: str) -> bool:
    return "/auth/refresh" in section_m and "publications/history/facts.csv" in section_m


def assert_renewal_enabled_mashup(body: bytes) -> None:
    """Require the Excel-authored /auth/refresh PublishedFacts formula."""
    section = assert_published_facts_mashup(body)
    if not mashup_is_renewal_enabled(section):
        raise ValueError("Client report is incomplete.")


def assert_native_pivot_package(body: bytes) -> None:
    """Raise ValueError if the nine native PivotTables / slicers are missing.

    Does not require a specific numeric workbook cacheId. Excel Save may remap
    1 → 0 or 1 → 12. All nine PivotTables must share the one workbook cache
    whose relationship still targets pivotCacheDefinition1.xml.
    """
    with ZipFile(io.BytesIO(body), "r") as archive:
        names = set(archive.namelist())
        if PIVOT_CACHE_PART not in names or PIVOT_RECORDS_PART not in names:
            raise ValueError("Client report is incomplete.")
        workbook_xml = archive.read("xl/workbook.xml").decode("utf-8")
        rels_xml = archive.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        cache_id, _rid, _target = shared_pivot_cache_binding(workbook_xml, rels_xml)
        bindings = bindings_from_workbook(workbook_xml, rels_xml)
        if len(bindings) != 9:
            raise ValueError("Client report is incomplete.")
        seen_names: set[str] = set()
        for binding in bindings:
            if binding.pivot_part not in names or binding.slicer_part not in names:
                raise ValueError("Client report is incomplete.")
            table = archive.read(binding.pivot_part).decode("utf-8")
            if f'name="{binding.pivot_name}"' not in table:
                raise ValueError("Client report is incomplete.")
            if _pivot_table_cache_id(table) != cache_id:
                raise ValueError("Client report is incomplete.")
            seen_names.add(binding.pivot_name)
        if seen_names != set(PIVOT_TABLE_NAMES.values()):
            raise ValueError("Client report is incomplete.")
        slicer_caches = [
            name for name in names if name.startswith("xl/slicerCaches/") and name.endswith(".xml")
        ]
        if len(slicer_caches) != 35:
            raise ValueError("Client report is incomplete.")
        if "x14:slicerCaches" not in workbook_xml:
            raise ValueError("Client report is incomplete.")
        cache_def = archive.read(PIVOT_CACHE_PART).decode("utf-8")
        if f'pivotCacheId="{SLICER_PIVOT_CACHE_ID}"' not in cache_def:
            raise ValueError("Client report is incomplete.")
    assert_published_facts_mashup(body)
