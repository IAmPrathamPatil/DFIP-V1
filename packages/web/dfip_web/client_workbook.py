"""Build the empty client workbook *structure* (Settings + Facts tables).

This module writes a spreadsheet template. It does **not** create Office
DataMashup, `xl/connections.xml`, query tables, or any Power Query package.
Python/openpyxl cannot manufacture a native-refresh workbook. Refresh All
requires Excel Desktop to load `excel/PublishedFacts.m` unchanged.

The M script on disk is the source of truth. Do not embed it as unregistered
zip parts (`xl/queryMashup/`, `customXml/powerQuery.xml`): those are not
DataMashup and caused Excel to open the file as Repaired.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[3]
EXCEL_DIR = ROOT / "excel"
M_PATH = EXCEL_DIR / "PublishedFacts.m"
XLSX_PATH = EXCEL_DIR / "Client_Report.xlsx"

# Python cannot author Microsoft DataMashup. Builder output is structure only.
BUILDER_CREATES_NATIVE_DATAMASHUP = False

SETTINGS_HEADERS = ("Parameter", "Value")
SETTINGS_ROWS = (
    ("ApiBaseUrl", "http://127.0.0.1:8000"),
    ("BearerToken", ""),
    ("ClientId", ""),
)

FACT_HEADERS = (
    "Filter Logic 1",
    "Filter Logic 2",
    "Template Status",
    "AMC Status - Filter Logic 3",
    "AMC Device Category -  Filter Logic 4",
    "AMC Product Cat -  Filter Logic 5",
    "Manual Or Automated",
    "Total Cost",
    "HHH",
    "Month",
    "Day",
    "Campaign Name",
    "Campaign ID",
    "Variation Name",
    "Variation ID",
    "Channel",
    "Type of Campaign",
    "Start Date",
    "Sent",
    "Failed",
    "Delivered",
    "Unique Impressions",
    "Unique Clicks",
    "Unique Conversions",
    "Unique Impression-Through Conversions",
    "Unique Click-Through Conversions",
    "Revenue (INR)",
    "Impression-Through Revenue (INR)",
    "Click-Through Revenue (INR)",
    "Template Name (WhatsApp)",
    "client_id",
    "variation_id_key",
    "month_start",
    "filter_logic_1_group",
    "label_match_status",
    "template_match_status",
    "rate_card_rule_id",
    "processing_run_id",
    "batch_id",
    "campaign_label_version_id",
    "template_label_version_id",
    "rate_card_version_id",
    "label_group_version_id",
    "first_seen_at",
    "last_seen_at",
)

FAKE_MASHUP_ZIP_PARTS = (
    "xl/queryMashup/PublishedFacts.m",
    "customXml/powerQuery.xml",
)


def mashup_text() -> str:
    return M_PATH.read_text(encoding="utf-8")


def build_client_report(path: Path | None = None) -> Path:
    """Write Settings + empty Facts tables. Not a native-refresh workbook."""
    target = path or XLSX_PATH
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Facts"
    bold = Font(bold=True)

    sheet["A1"] = SETTINGS_HEADERS[0]
    sheet["B1"] = SETTINGS_HEADERS[1]
    sheet["A1"].font = bold
    sheet["B1"].font = bold
    for index, (key, value) in enumerate(SETTINGS_ROWS, start=2):
        sheet.cell(index, 1, key)
        sheet.cell(index, 2, value)
    settings = Table(displayName="Settings", ref="A1:B4")
    settings.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(settings)

    header_row = 6
    for col, header in enumerate(FACT_HEADERS, start=1):
        cell = sheet.cell(header_row, col, header)
        cell.font = bold
        sheet.cell(header_row + 1, col, None)
    last_col = get_column_letter(len(FACT_HEADERS))
    facts = Table(
        displayName="Facts",
        ref=f"A{header_row}:{last_col}{header_row + 1}",
    )
    facts.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    sheet.add_table(facts)
    sheet.freeze_panes = "A7"
    workbook.save(target)
    return target


if __name__ == "__main__":
    built = build_client_report()
    print(built)
    print(
        "Wrote structure-only Settings + Facts. "
        "This is not Office DataMashup and cannot Refresh All. "
        "Load excel/PublishedFacts.m in Excel Desktop with an empty BearerToken."
    )
