"""RUN 012: Download Consolidated Company Workbook.

A third export mode beside the static Company Workbook and the Refreshable
Workbook. Static file, cumulative rows: the same template, renderer, sheet
names, formulas and pivots as the static Company Workbook, populated with
every successfully processed month currently published for the company, so
Excel Refresh All is never required.

The cumulative row set is the same newest-wins complete-snapshot union the
refresh CSV serves, so dedup, ordering and exclusion of incomplete runs are
inherited rather than reimplemented.
"""

from __future__ import annotations

import io
from datetime import UTC, date, datetime
from zipfile import ZipFile

import pytest
from dfip_api import publication_service as publication_service_module
from dfip_api.publication_service import (
    NO_PUBLICATION_WORKBOOK_MESSAGE,
    ROW_CAP_WORKBOOK_MESSAGE,
)
from dfip_api.schemas import HISTORY_FACTS_MAX_PAGE_LIMIT
from dfip_web.client_report_download import (
    ARTIFACT_CONSOLIDATED,
    ARTIFACT_REFRESHABLE,
    ARTIFACT_STATIC,
    CLIENT_REPORT_CONSOLIDATED_SUFFIX,
    QUERY_TABLE_PART,
    client_report_download_filename,
)
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from test_company_workbook import _assert_static_and_native, _published_headers
from test_p5_api import AUTH, CLIENT_ID
from test_p7_publication import _error, publisher_app
from test_p8_client_report import _campaign_ids, _session_jwt_stamped, _settings_value
from test_r10_published_history import _fact, _publish_snapshot

CLIENT_B = "a0000000-0000-4000-8000-000000000002"
CONSOLIDATED_PATH = "/api/v1/publications/current/consolidated-client-report.xlsx"
STATIC_PATH = "/api/v1/publications/current/client-report.xlsx"
REFRESHABLE_PATH = "/api/v1/publications/current/refreshable-client-report.xlsx"

RUN_M1 = "a0000000-0000-4000-8000-0000000000a1"
RUN_M2 = "a0000000-0000-4000-8000-0000000000a2"
RUN_M3 = "a0000000-0000-4000-8000-0000000000a3"
RUN_M1_REDO = "a0000000-0000-4000-8000-0000000000b1"
RUN_FAILED = "a0000000-0000-4000-8000-0000000000f1"

MONTHS = {
    "Jan-25": date(2025, 1, 6),
    "Feb-25": date(2025, 2, 6),
    "Mar-25": date(2025, 3, 6),
}


def _add_run(ingest, run_id: str, *, status: str = "succeeded") -> None:
    from dfip_core.ingest.store import ProcessingRunRecord

    ingest.processing_runs[run_id] = ProcessingRunRecord(
        id=run_id,
        batch_id="a0000000-0000-4000-8000-0000000000b0",
        campaign_label_version_id="a0000000-0000-4000-8000-000000000022",
        template_label_version_id="a0000000-0000-4000-8000-000000000034",
        rate_card_version_id="a0000000-0000-4000-8000-000000000012",
        label_group_version_id="a0000000-0000-4000-8000-000000000041",
        engine_version="0.4.0",
        started_at=datetime(2025, 1, 1, tzinfo=UTC),
        finished_at=datetime(2025, 1, 1, 0, 5, tzinfo=UTC),
        status=status,
        qa_verdict="pass" if status == "succeeded" else "fail",
    )


def _month_facts(month_label: str, *, sent: int, clicks: int, client_id: str = CLIENT_ID):
    day = MONTHS[month_label]
    return [
        _fact(
            client_id=client_id,
            campaign_id=f"camp-{month_label}",
            day=day,
            month_label=month_label,
            sent=sent,
            unique_conversions=clicks,
            unique_click_through_conversions=clicks,
        )
    ]


def _measure(body: bytes, header: str) -> float:
    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=False)
    sheet = workbook["PublishedFacts"]
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    index = list(rows[0]).index(header)
    return sum(float(row[index] or 0) for row in rows[1:] if row[index] not in {None, ""})


def _months(body: bytes) -> set[str]:
    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=False)
    sheet = workbook["PublishedFacts"]
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    index = list(rows[0]).index("Month")
    return {str(row[index]) for row in rows[1:] if row[index] not in {None, ""}}


def _app_with_months(months: list[str], *, other_company: bool = False):
    app, ingest, facts, store = publisher_app()
    for index, month_label in enumerate(months, start=1):
        run_id = f"a0000000-0000-4000-8000-0000000000a{index}"
        _add_run(ingest, run_id)
        _publish_snapshot(
            store,
            _month_facts(month_label, sent=10, clicks=2),
            run_id=run_id,
        )
    if other_company:
        _add_run(ingest, "a0000000-0000-4000-8000-0000000000c1")
        _publish_snapshot(
            store,
            _month_facts("Jan-25", sent=999, clicks=999, client_id=CLIENT_B),
            run_id="a0000000-0000-4000-8000-0000000000c1",
        )
    return TestClient(app), ingest, facts, store


def test_month_one_appears_in_consolidated_workbook() -> None:
    http, *_rest = _app_with_months(["Jan-25"])
    response = http.get(CONSOLIDATED_PATH, headers=AUTH, params={"client_id": CLIENT_ID})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert CLIENT_REPORT_CONSOLIDATED_SUFFIX in response.headers["content-disposition"]
    assert _months(response.content) == {"Jan-25"}


def test_two_months_both_appear() -> None:
    http, *_rest = _app_with_months(["Jan-25", "Feb-25"])
    response = http.get(CONSOLIDATED_PATH, headers=AUTH, params={"client_id": CLIENT_ID})
    assert response.status_code == 200
    assert _months(response.content) == {"Jan-25", "Feb-25"}


def test_three_months_all_appear_in_one_static_file() -> None:
    http, *_rest = _app_with_months(["Jan-25", "Feb-25", "Mar-25"])
    response = http.get(CONSOLIDATED_PATH, headers=AUTH, params={"client_id": CLIENT_ID})
    assert response.status_code == 200
    assert _months(response.content) == {"Jan-25", "Feb-25", "Mar-25"}


def test_republished_month_is_not_duplicated() -> None:
    app, ingest, facts, store = publisher_app()
    http = TestClient(app)
    _add_run(ingest, RUN_M1)
    _add_run(ingest, RUN_M1_REDO)
    _publish_snapshot(store, _month_facts("Jan-25", sent=10, clicks=2), run_id=RUN_M1)
    # Same reporting period republished with different values.
    _publish_snapshot(store, _month_facts("Jan-25", sent=55, clicks=7), run_id=RUN_M1_REDO)
    response = http.get(CONSOLIDATED_PATH, headers=AUTH, params={"client_id": CLIENT_ID})
    assert response.status_code == 200
    # One period, not two, and the newest publication wins.
    assert _months(response.content) == {"Jan-25"}
    assert len(_campaign_ids(response.content)) == 1
    assert _measure(response.content, "Sent") == 55


def test_failed_or_unpublished_run_is_excluded() -> None:
    app, ingest, facts, store = publisher_app()
    http = TestClient(app)
    _add_run(ingest, RUN_M1)
    _add_run(ingest, RUN_FAILED, status="failed")
    _publish_snapshot(store, _month_facts("Jan-25", sent=10, clicks=2), run_id=RUN_M1)
    # Failed run data stays in the working set only: never published.
    facts.upsert(
        _fact(
            client_id=CLIENT_ID,
            campaign_id="camp-failed",
            day=date(2025, 3, 6),
            month_label="Mar-25",
            sent=777,
            unique_conversions=777,
            unique_click_through_conversions=777,
            processing_run_id=RUN_FAILED,
        )
    )
    response = http.get(CONSOLIDATED_PATH, headers=AUTH, params={"client_id": CLIENT_ID})
    assert response.status_code == 200
    assert _months(response.content) == {"Jan-25"}
    assert "camp-failed" not in _campaign_ids(response.content)
    assert _measure(response.content, "Sent") == 10


def test_another_company_never_leaks_into_the_workbook() -> None:
    http, *_rest = _app_with_months(["Jan-25"], other_company=True)
    response = http.get(CONSOLIDATED_PATH, headers=AUTH, params={"client_id": CLIENT_ID})
    assert response.status_code == 200
    assert "camp-Jan-25" in _campaign_ids(response.content)
    assert _measure(response.content, "Sent") == 10


def test_cross_company_download_is_forbidden() -> None:
    from test_p7_publication import jwt_app

    app, _store, _facts, headers = jwt_app(role="reader", client_id=CLIENT_B)
    http = TestClient(app)
    denied = http.get(CONSOLIDATED_PATH, headers=headers, params={"client_id": CLIENT_ID})
    assert denied.status_code == 403


def test_client_may_download_own_consolidated_workbook() -> None:
    from test_p7_publication import jwt_app

    app, store, _facts, headers = jwt_app(role="client", client_id=CLIENT_ID)
    _add_run(app.state.ingest_store, RUN_M1)
    _publish_snapshot(store, _month_facts("Jan-25", sent=10, clicks=2), run_id=RUN_M1)
    http = TestClient(app)
    response = http.get(CONSOLIDATED_PATH, headers=headers)
    assert response.status_code == 200
    assert _months(response.content) == {"Jan-25"}


def test_unauthenticated_download_is_401() -> None:
    http, *_rest = _app_with_months(["Jan-25"])
    assert http.get(CONSOLIDATED_PATH).status_code == 401


def test_no_history_returns_404_instead_of_an_empty_workbook() -> None:
    http, *_rest = _app_with_months([])
    response = http.get(CONSOLIDATED_PATH, headers=AUTH, params={"client_id": CLIENT_ID})
    assert response.status_code == 404
    assert _error(response)["message"] == NO_PUBLICATION_WORKBOOK_MESSAGE


def test_row_cap_is_422_and_never_truncates(monkeypatch: pytest.MonkeyPatch) -> None:
    """Over the cap is an error, never a partial workbook."""
    http, *_rest = _app_with_months(["Jan-25", "Feb-25", "Mar-25"])
    # Three monthly rows are staged; cap them below the dataset size.
    monkeypatch.setattr(publication_service_module, "HISTORY_FACTS_MAX_PAGE_LIMIT", 2)
    response = http.get(CONSOLIDATED_PATH, headers=AUTH, params={"client_id": CLIENT_ID})
    assert response.status_code == 422
    assert _error(response)["message"] == ROW_CAP_WORKBOOK_MESSAGE
    # Back under the cap the same company downloads every month.
    monkeypatch.setattr(
        publication_service_module,
        "HISTORY_FACTS_MAX_PAGE_LIMIT",
        HISTORY_FACTS_MAX_PAGE_LIMIT,
    )
    ok = http.get(CONSOLIDATED_PATH, headers=AUTH, params={"client_id": CLIENT_ID})
    assert ok.status_code == 200
    assert _months(ok.content) == {"Jan-25", "Feb-25", "Mar-25"}


def test_aggregated_measures_accumulate_across_months() -> None:
    app, ingest, facts, store = publisher_app()
    http = TestClient(app)
    _add_run(ingest, RUN_M1)
    _add_run(ingest, RUN_M2)
    _add_run(ingest, RUN_M3)
    _publish_snapshot(store, _month_facts("Jan-25", sent=10, clicks=2), run_id=RUN_M1)
    _publish_snapshot(store, _month_facts("Feb-25", sent=20, clicks=4), run_id=RUN_M2)
    _publish_snapshot(store, _month_facts("Mar-25", sent=30, clicks=6), run_id=RUN_M3)
    response = http.get(CONSOLIDATED_PATH, headers=AUTH, params={"client_id": CLIENT_ID})
    assert response.status_code == 200
    body = response.content
    assert _measure(body, "Sent") == 60
    assert _measure(body, "Unique Conversions") == 12
    # Derived measures follow the existing DFIP definitions over aggregated values.
    assert _measure(body, "Unique Click-Through Conversions") == 12


def test_structure_matches_the_static_company_template() -> None:
    http, *_rest = _app_with_months(["Jan-25", "Feb-25", "Mar-25"])
    response = http.get(CONSOLIDATED_PATH, headers=AUTH, params={"client_id": CLIENT_ID})
    assert response.status_code == 200
    body = response.content
    _assert_static_and_native(body)
    assert _published_headers(body) == _published_headers(
        http.get(STATIC_PATH, headers=AUTH, params={"client_id": CLIENT_ID}).content
    )


def test_consolidated_never_requires_excel_refresh() -> None:
    http, *_rest = _app_with_months(["Jan-25", "Feb-25"])
    response = http.get(CONSOLIDATED_PATH, headers=AUTH, params={"client_id": CLIENT_ID})
    assert response.status_code == 200
    body = response.content
    with ZipFile(io.BytesIO(body)) as archive:
        names = archive.namelist()
    # No query table, so Excel has nothing to refresh, and no grant is embedded.
    assert QUERY_TABLE_PART not in names
    assert not _session_jwt_stamped(_settings_value(body, "BearerToken"))
    assert _settings_value(body, "ApiBaseUrl") in {None, ""}


def test_existing_static_company_workbook_is_unchanged() -> None:
    http, *_rest = _app_with_months(["Jan-25", "Feb-25", "Mar-25"])
    response = http.get(STATIC_PATH, headers=AUTH, params={"client_id": CLIENT_ID})
    assert response.status_code == 200
    body = response.content
    _assert_static_and_native(body)
    assert "Client_Report.xlsx" in response.headers["content-disposition"]
    assert ARTIFACT_STATIC not in response.headers["content-disposition"]


def test_existing_refreshable_workbook_is_unchanged() -> None:
    http, *_rest = _app_with_months(["Jan-25", "Feb-25"])
    response = http.get(REFRESHABLE_PATH, headers=AUTH, params={"client_id": CLIENT_ID})
    assert response.status_code == 200
    body = response.content
    assert body.startswith(b"PK")
    assert "Client_Report_Refreshable.xlsm" in response.headers["content-disposition"]
    assert response.headers["content-type"].startswith(
        "application/vnd.ms-excel.sheet.macroEnabled"
    )
    # Refreshable still writes ApiBaseUrl, which consolidated and static do not.
    assert _settings_value(body, "ApiBaseUrl")


def test_consolidated_filename_is_distinct_from_the_other_two() -> None:
    from datetime import UTC, datetime

    published = datetime(2026, 9, 1, 12, 0, tzinfo=UTC)
    assert (
        client_report_download_filename("default", published, artifact=ARTIFACT_CONSOLIDATED)
        == "DFIP_default_2026-09-01_Client_Report_Consolidated.xlsx"
    )
    assert client_report_download_filename("default", published, artifact=ARTIFACT_STATIC) == (
        "DFIP_default_2026-09-01_Client_Report.xlsx"
    )
    assert client_report_download_filename("default", published, artifact=ARTIFACT_REFRESHABLE) == (
        "DFIP_default_2026-09-01_Client_Report_Refreshable.xlsm"
    )


def test_frontend_exposes_a_third_download_button() -> None:
    from pathlib import Path

    web_static = Path(__file__).resolve().parents[1] / "apps" / "web" / "static"
    views = (web_static / "js" / "views.js").read_text(encoding="utf-8")
    app_js = (web_static / "js" / "app.js").read_text(encoding="utf-8")
    client_js = (web_static / "js" / "api-client.js").read_text(encoding="utf-8")

    assert "Download Consolidated Company Workbook" in views
    assert "data-download-consolidated-client-report" in views
    # The two existing buttons are untouched.
    assert "Download Company Workbook" in views
    assert "Download Refreshable Workbook" in views
    # UI text explains the static, no-refresh, all-months behaviour.
    assert "never needs Excel Refresh All" in views
    assert "successfully processed month" in views
    assert "downloadConsolidatedClientReport" in client_js
    assert "consolidated-client-report.xlsx" in client_js
    assert "data-download-consolidated-client-report" in app_js


def test_consolidated_download_is_read_only() -> None:
    """Consolidated must not create, move, or republish publication_current."""
    http, _ingest, _facts, store = _app_with_months(["Jan-25", "Feb-25"])
    before = list(store.publications)
    response = http.get(CONSOLIDATED_PATH, headers=AUTH, params={"client_id": CLIENT_ID})
    assert response.status_code == 200
    assert _months(response.content) == {"Jan-25", "Feb-25"}
    assert list(store.publications) == before
    current, pointer = store.get_current(CLIENT_ID)
    assert pointer is not None
    assert current.id == before[-1]
