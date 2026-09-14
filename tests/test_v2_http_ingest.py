"""HTTP workbook ingest and published-slice download (in-memory stores)."""

from __future__ import annotations

import csv
import io
import tempfile
from datetime import date
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from dfip_api.app import create_app
from dfip_api.published_download import FACT_DOWNLOAD_COLUMNS
from dfip_api.schemas import MAX_PAGE_LIMIT
from dfip_core.ingest.store import InMemoryIngestStore
from dfip_core.transform.fact import FactRecord
from dfip_core.transform.store import InMemoryFactStore
from dfip_web.api_client import DfipApiClient
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from http_ingest_support import source_row, upload_workbook, workbook_bytes
from test_p5_api import (
    AUTH,
    CLIENT_ID,
    DEV_TOKEN,
    JWT_SECRET,
    _encode_jwt,
    inspector_settings,
    make_settings,
)
from test_p7_publication import publisher_settings

CLIENT_B = "a0000000-0000-4000-8000-000000000002"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _error(response) -> dict:
    payload = response.json()
    assert "error" in payload
    return payload["error"]


def _publisher_app(**overrides):
    return create_app(
        settings=publisher_settings(**overrides),
        ingest_store=InMemoryIngestStore(),
        fact_store=InMemoryFactStore(),
    )


def _jwt_app(role: str, client_id: str | None):
    return create_app(
        settings=make_settings(dfip_auth_mode="jwt", dfip_auth_secret=JWT_SECRET),
        ingest_store=InMemoryIngestStore(),
        fact_store=InMemoryFactStore(),
    ), {"Authorization": f"Bearer {_encode_jwt(role=role, client_id=client_id)}"}


def _upload(
    client: TestClient,
    content: bytes,
    filename: str,
    *,
    headers=AUTH,
    wait: bool = True,
    timeout: float = 60.0,
    **form,
):
    return upload_workbook(
        client,
        content,
        filename,
        headers=headers,
        wait=wait,
        timeout=timeout,
        **form,
    )


def test_unauthenticated_upload_is_401(tmp_path: Path) -> None:
    app = _publisher_app()
    content = workbook_bytes(tmp_path / "a.xlsx", [source_row()])
    response = _upload(TestClient(app), content, "a.xlsx", headers={})
    assert response.status_code == 401
    assert _error(response)["code"] == "AUTHENTICATION_FAILED"


def test_reader_cannot_upload(tmp_path: Path) -> None:
    app = create_app(
        settings=make_settings(dfip_dev_auth_role="reader"),
        ingest_store=InMemoryIngestStore(),
        fact_store=InMemoryFactStore(),
    )
    content = workbook_bytes(tmp_path / "a.xlsx", [source_row()])
    response = _upload(TestClient(app), content, "a.xlsx", client_id=CLIENT_ID)
    assert response.status_code == 403
    assert _error(response)["code"] == "AUTHORIZATION_FAILED"


def test_upload_ingests_and_transforms_without_publishing(tmp_path: Path) -> None:
    app = _publisher_app()
    http = TestClient(app)
    content = workbook_bytes(tmp_path / "a.xlsx", [source_row()])
    response = _upload(http, content, "client-a.xlsx", client_id=CLIENT_ID)
    assert response.status_code == 201
    body = response.json()
    assert body["published"] is False
    assert body["replayed"] is False
    assert body["client_id"] == CLIENT_ID
    assert body["original_filename"] == "client-a.xlsx"
    assert body["processing_run"]["status"] == "succeeded"
    assert body["transform"]["transformed"] == 1
    assert body["batch"]["status"] == "processed"
    assert tempfile.gettempdir() not in response.text
    assert "dfip-upload" not in response.text
    working = http.get("/api/v1/facts?limit=200", headers=AUTH).json()
    assert working["pagination"]["total"] == 1
    assert working["items"][0]["unique_clicks"] == 2
    assert working["items"][0]["total_cost"] == "15.0000"
    published = http.get(
        "/api/v1/publications/current/facts",
        headers=AUTH,
        params={"client_id": CLIENT_ID},
    ).json()
    assert published["items"] == []
    downloaded = http.get(
        "/api/v1/publications/current/facts.csv",
        headers=AUTH,
        params={"client_id": CLIENT_ID},
    )
    assert downloaded.status_code == 200
    rows = list(csv.DictReader(io.StringIO(downloaded.text)))
    assert rows == []


def test_jwt_client_cannot_upload_for_another_client(tmp_path: Path) -> None:
    app, headers = _jwt_app("publisher", CLIENT_ID)
    content = workbook_bytes(tmp_path / "a.xlsx", [source_row()])
    response = _upload(TestClient(app), content, "a.xlsx", headers=headers, client_id=CLIENT_B)
    assert response.status_code == 403
    assert _error(response)["code"] == "AUTHORIZATION_FAILED"


def test_unsupported_file_type_is_rejected(tmp_path: Path) -> None:
    app = _publisher_app()
    response = _upload(TestClient(app), b"not-xlsx", "notes.csv", client_id=CLIENT_ID)
    assert response.status_code == 422
    assert _error(response)["code"] == "VALIDATION_ERROR"
    assert "xlsx" in _error(response)["message"].lower()


def test_path_traversal_filename_is_rejected(tmp_path: Path) -> None:
    app = _publisher_app()
    content = workbook_bytes(tmp_path / "a.xlsx", [source_row()])
    response = _upload(TestClient(app), content, "../../etc/passwd.xlsx", client_id=CLIENT_ID)
    assert response.status_code == 422
    assert _error(response)["code"] == "VALIDATION_ERROR"
    assert tempfile.gettempdir() not in response.text


def test_oversized_upload_is_413() -> None:
    app = _publisher_app(dfip_upload_max_bytes=64)
    payload = b"PK" + b"A" * 200
    response = _upload(TestClient(app), payload, "big.xlsx", client_id=CLIENT_ID)
    assert response.status_code == 413
    assert _error(response)["code"] == "PAYLOAD_TOO_LARGE"


def test_malformed_workbook_is_rejected() -> None:
    app = _publisher_app()
    payload = b"PK\x03\x04this-is-not-a-zip-archive"
    response = _upload(TestClient(app), payload, "bad.xlsx", client_id=CLIENT_ID)
    assert response.status_code == 422
    assert _error(response)["code"] == "VALIDATION_ERROR"
    assert tempfile.gettempdir() not in response.text


def test_invalid_headers_fail_before_processing_run(tmp_path: Path) -> None:
    app = _publisher_app()
    content = workbook_bytes(
        tmp_path / "bad-headers.xlsx",
        [source_row()],
        header_override={12: "Campaign"},
    )
    response = _upload(TestClient(app), content, "bad-headers.xlsx", client_id=CLIENT_ID)
    assert response.status_code == 201
    body = response.json()
    assert body["processing_run"] is None
    assert body["published"] is False
    assert body["batch"]["status"] == "failed"
    assert body["transform"] is None
    codes = {item["reason_code"] for item in body["rejections"]}
    summary = body["batch"]["error_summary"] or ""
    assert codes.intersection({"HEADER_CONTRACT", "DUPLICATE_HEADERS"}) or (
        "HEADER_CONTRACT" in summary or "DUPLICATE_HEADERS" in summary
    )
    facts = TestClient(app).get("/api/v1/facts", headers=AUTH).json()
    assert facts["pagination"]["total"] == 0


def test_invalid_numeric_preserves_reason_code(tmp_path: Path) -> None:
    app = _publisher_app()
    content = workbook_bytes(tmp_path / "num.xlsx", [source_row(**{"Sent": "twelve"})])
    response = _upload(TestClient(app), content, "num.xlsx", client_id=CLIENT_ID)
    assert response.status_code == 201
    body = response.json()
    assert body["processing_run"]["status"] == "failed"
    assert any(item["reason_code"] == "INVALID_NUMERIC" for item in body["rejections"])
    assert body["transform"]["transformed"] == 0
    facts = TestClient(app).get("/api/v1/facts", headers=AUTH).json()
    assert facts["pagination"]["total"] == 0


def test_duplicate_grain_preserves_reason_code(tmp_path: Path) -> None:
    app = _publisher_app()
    rows = [source_row(), source_row(**{"Sent": 90})]
    content = workbook_bytes(tmp_path / "dup.xlsx", rows)
    response = _upload(TestClient(app), content, "dup.xlsx", client_id=CLIENT_ID)
    assert response.status_code == 201
    body = response.json()
    assert body["processing_run"]["status"] == "succeeded"
    assert body["transform"]["transformed"] == 1
    assert any(item["reason_code"] == "DUPLICATE_GRAIN_KEY" for item in body["rejections"])


def test_publish_then_download_matches_published_facts(tmp_path: Path) -> None:
    app = _publisher_app()
    http = TestClient(app)
    content = workbook_bytes(tmp_path / "a.xlsx", [source_row()])
    uploaded = _upload(http, content, "a.xlsx", client_id=CLIENT_ID).json()
    run_id = uploaded["processing_run"]["processing_run_id"]
    published = http.post(
        "/api/v1/publications",
        headers=AUTH,
        json={"client_id": CLIENT_ID, "processing_run_id": run_id},
    )
    assert published.status_code == 201
    facts = http.get(
        "/api/v1/publications/current/facts",
        headers=AUTH,
        params={"client_id": CLIENT_ID, "limit": 200},
    ).json()
    csv_response = http.get(
        "/api/v1/publications/current/facts.csv",
        headers=AUTH,
        params={"client_id": CLIENT_ID},
    )
    assert csv_response.status_code == 200
    assert csv_response.headers["content-type"].startswith("text/csv")
    assert "published-facts-" in csv_response.headers["content-disposition"]
    rows = list(csv.DictReader(io.StringIO(csv_response.content.decode("utf-8"))))
    assert list(rows[0].keys()) == list(FACT_DOWNLOAD_COLUMNS)
    assert len(rows) == 1
    assert rows[0]["campaign_id"] == facts["items"][0]["campaign_id"]
    assert rows[0]["unique_clicks"] == str(facts["items"][0]["unique_clicks"])
    assert rows[0]["total_cost"] == facts["items"][0]["total_cost"]
    xlsx_response = http.get(
        "/api/v1/publications/current/facts.xlsx",
        headers=AUTH,
        params={"client_id": CLIENT_ID},
    )
    assert xlsx_response.status_code == 200
    assert xlsx_response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(BytesIO(xlsx_response.content), read_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    values = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
    workbook.close()
    assert headers == list(FACT_DOWNLOAD_COLUMNS)
    assert values[headers.index("unique_clicks")] == "2"
    assert values[headers.index("total_cost")] == "15.0000"


def test_download_does_not_use_working_set_or_raise_page_limit(tmp_path: Path) -> None:
    app = _publisher_app()
    http = TestClient(app)
    content = workbook_bytes(tmp_path / "a.xlsx", [source_row()])
    uploaded = _upload(http, content, "a.xlsx", client_id=CLIENT_ID).json()
    extra = FactRecord(
        client_id=CLIENT_ID,
        campaign_id="camp-unpublished",
        variation_id="var-x",
        variation_id_key="var-x",
        day=date(2025, 8, 2),
        processing_run_id="d0000000-0000-4000-8000-000000000099",
        unique_clicks=99,
    )
    app.state.fact_store.facts[extra.key] = extra
    working = http.get("/api/v1/facts?limit=200", headers=AUTH).json()
    assert working["pagination"]["total"] == 2
    http.post(
        "/api/v1/publications",
        headers=AUTH,
        json={
            "client_id": CLIENT_ID,
            "processing_run_id": uploaded["processing_run"]["processing_run_id"],
        },
    )
    rejected = http.get(
        "/api/v1/publications/current/facts",
        headers=AUTH,
        params={"client_id": CLIENT_ID, "limit": MAX_PAGE_LIMIT + 1},
    )
    assert rejected.status_code == 422
    assert MAX_PAGE_LIMIT == 200
    csv_response = http.get(
        "/api/v1/publications/current/facts.csv",
        headers=AUTH,
        params={"client_id": CLIENT_ID},
    )
    rows = list(csv.DictReader(io.StringIO(csv_response.text)))
    assert len(rows) == 1
    assert rows[0]["campaign_id"] == "camp-a"
    assert "camp-unpublished" not in csv_response.text


def test_cross_client_download_is_forbidden(tmp_path: Path) -> None:
    ingest = InMemoryIngestStore()
    facts = InMemoryFactStore()
    publisher = create_app(
        settings=publisher_settings(),
        ingest_store=ingest,
        fact_store=facts,
    )
    reader = create_app(
        settings=make_settings(dfip_auth_mode="jwt", dfip_auth_secret=JWT_SECRET),
        ingest_store=ingest,
        fact_store=facts,
        publication_store=publisher.state.publication_store,
    )
    http = TestClient(publisher)
    content = workbook_bytes(tmp_path / "a.xlsx", [source_row()])
    uploaded = _upload(http, content, "a.xlsx", client_id=CLIENT_ID).json()
    http.post(
        "/api/v1/publications",
        headers=AUTH,
        json={
            "client_id": CLIENT_ID,
            "processing_run_id": uploaded["processing_run"]["processing_run_id"],
        },
    )
    other = TestClient(reader)
    headers = {"Authorization": f"Bearer {_encode_jwt(role='client', client_id=CLIENT_B)}"}
    scoped = other.get("/api/v1/publications/current/facts.csv", headers=headers)
    assert scoped.status_code == 200
    assert list(csv.DictReader(io.StringIO(scoped.text))) == []
    denied = other.get(
        "/api/v1/publications/current/facts.csv",
        headers=headers,
        params={"client_id": CLIENT_ID},
    )
    assert denied.status_code == 403
    working = other.get("/api/v1/facts", headers=headers)
    assert working.status_code == 403


def test_reader_cannot_inspect_working_set_after_upload(tmp_path: Path) -> None:
    app = _publisher_app()
    content = workbook_bytes(tmp_path / "a.xlsx", [source_row()])
    _upload(TestClient(app), content, "a.xlsx", client_id=CLIENT_ID)
    reader = create_app(
        settings=make_settings(),
        ingest_store=app.state.ingest_store,
        fact_store=app.state.fact_store,
        publication_store=app.state.publication_store,
    )
    response = TestClient(reader).get("/api/v1/facts", headers=AUTH)
    assert response.status_code == 403


def test_replay_does_not_publish(tmp_path: Path) -> None:
    app = _publisher_app()
    http = TestClient(app)
    content = workbook_bytes(tmp_path / "a.xlsx", [source_row()])
    first = _upload(http, content, "a.xlsx", client_id=CLIENT_ID)
    second = _upload(http, content, "a.xlsx", client_id=CLIENT_ID)
    assert first.status_code == 201
    assert second.status_code == 200
    body = second.json()
    assert body["replayed"] is True
    assert body["published"] is False
    assert (
        body["processing_run"]["processing_run_id"]
        == first.json()["processing_run"]["processing_run_id"]
    )


def test_python_client_upload_and_download(tmp_path: Path) -> None:
    app = _publisher_app()
    content = workbook_bytes(tmp_path / "a.xlsx", [source_row()])
    with DfipApiClient("http://testserver", token=DEV_TOKEN, http_client=TestClient(app)) as api:
        uploaded = api.upload_workbook(content, "a.xlsx", client_id=CLIENT_ID)
        created = api.create_publication(
            {
                "client_id": CLIENT_ID,
                "processing_run_id": uploaded["processing_run"]["processing_run_id"],
            }
        )
        payload = api.download_published_facts(kind="csv", client_id=CLIENT_ID)
    assert uploaded["transform"]["transformed"] == 1
    assert (
        created["publication"]["processing_run_id"]
        == uploaded["processing_run"]["processing_run_id"]
    )
    rows = list(csv.DictReader(io.StringIO(payload.decode("utf-8"))))
    assert rows[0]["campaign_id"] == "camp-a"


def test_zip_without_workbook_parts_is_rejected() -> None:
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("readme.txt", "not a workbook")
    app = _publisher_app()
    response = _upload(TestClient(app), buffer.getvalue(), "fake.xlsx", client_id=CLIENT_ID)
    assert response.status_code == 422
    assert _error(response)["code"] == "VALIDATION_ERROR"


def test_page_limit_constant_is_unchanged() -> None:
    assert MAX_PAGE_LIMIT == 200
    settings = inspector_settings()
    assert settings.dfip_upload_max_bytes == 64 * 1024 * 1024
    assert settings.dfip_upload_max_files == 5
    assert settings.dfip_upload_max_total_bytes == 128 * 1024 * 1024


def test_dev_token_upload_requires_client_id(tmp_path: Path) -> None:
    app = _publisher_app()
    content = workbook_bytes(tmp_path / "a.xlsx", [source_row()])
    response = _upload(TestClient(app), content, "a.xlsx")
    assert response.status_code == 422
    assert _error(response)["code"] == "VALIDATION_ERROR"


def test_unauthenticated_download_is_401() -> None:
    app = _publisher_app()
    http = TestClient(app)
    assert http.get("/api/v1/publications/current/facts.csv").status_code == 401
    assert http.get("/api/v1/publications/current/facts.xlsx").status_code == 401


def test_cors_exposes_content_disposition_for_published_download() -> None:
    app = _publisher_app()
    http = TestClient(app)
    response = http.get(
        "/api/v1/publications/current/facts.csv",
        headers={**AUTH, "Origin": "http://127.0.0.1:3000"},
        params={"client_id": CLIENT_ID},
    )
    assert response.status_code == 200
    expose = response.headers.get("access-control-expose-headers", "")
    assert "content-disposition" in expose.lower()
    assert response.headers.get("access-control-allow-origin") == "http://127.0.0.1:3000"
