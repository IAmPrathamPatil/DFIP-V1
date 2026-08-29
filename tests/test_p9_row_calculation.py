"""P9: each hierarchy row gets its own KPI spill, not only the first detail row."""

from __future__ import annotations

import re
from xml.etree import ElementTree

from dfip_web.client_workbook import XLSX_PATH
from dfip_web.daily_report import (
    MEASURE_OPS,
    NS_MAIN,
    REPORT_CONTRACTS,
    additive_formula,
    contract_uses_case_identity,
    derived_formula,
    report_formula,
    worksheet_xml,
)
from openpyxl.utils import get_column_letter

from test_p7_publication import FORBIDDEN_TEMPLATE_TOKENS, WORKING_SET_FACTS_PATH


def _measure_start(contract) -> int:
    return 2 + len(contract.display_fields)


def test_hierarchy_formula_still_unique_filters_all_group_columns() -> None:
    for contract in REPORT_CONTRACTS:
        text = report_formula(contract)
        assert "UNIQUE(" in text
        assert "FILTER(" in text
        if contract_uses_case_identity(contract):
            assert "LOWER(" in text
        else:
            assert "LOWER(" not in text
        assert "INDEX($B10#" not in text
        assert "LAMBDA" not in text
        assert "GROUPBY" not in text
        assert WORKING_SET_FACTS_PATH not in text
        xml = worksheet_xml(contract)
        last_dim = get_column_letter(1 + len(contract.display_fields))
        assert f't="array" ref="B10:{last_dim}10"' in xml


def test_sumifs_criteria_are_the_full_hierarchy_spill_not_first_row_only() -> None:
    for contract in REPORT_CONTRACTS:
        additive = additive_formula(contract, "Sent")
        display_count = len(contract.display_fields)
        assert "INDEX($B10#,0,1)" in additive
        for index in range(1, display_count + 1):
            assert f"INDEX($B10#,0,{index})" in additive
        assert re.search(r"SUMIFS\([^)]*,\$B10\)", additive) is None
        assert "SUMIFS(" in additive
        ratio = derived_formula(contract, ("Delivery Rate", "ratio", "Delivered", "Sent"))
        assert "IFERROR" in ratio
        sent = get_column_letter(_measure_start(contract) + 1)
        delivered = get_column_letter(_measure_start(contract) + 5)
        assert f"{delivered}10#/{sent}10#" in ratio


def test_measure_cells_store_single_cell_array_formulas_so_they_spill() -> None:
    """Excel Formula2 save uses t='array' ref='E10'. Without it, E11+ stay blank."""
    for contract in REPORT_CONTRACTS:
        xml = worksheet_xml(contract)
        root = ElementTree.fromstring(xml)
        ns = {"m": NS_MAIN}
        start = _measure_start(contract)
        for index, op in enumerate(MEASURE_OPS):
            cell = f"{get_column_letter(start + index)}10"
            found = None
            for node in root.findall("m:sheetData/m:row/m:c", ns):
                if node.get("r") == cell:
                    found = node
                    break
            assert found is not None, cell
            assert found.get("cm") == "1", cell
            formula = found.find("m:f", ns)
            assert formula is not None, cell
            assert formula.get("t") == "array", (contract.name, cell, op[0])
            assert formula.get("ref") == cell, (contract.name, cell)
            body = formula.text or ""
            if op[1] == "sum":
                assert "SUMIFS(" in body
                assert "INDEX(_xlfn.ANCHORARRAY($B10),0," in body
            else:
                assert "ANCHORARRAY" in body


def test_three_independent_dates_are_separate_sumifs_criteria_slots() -> None:
    overall = REPORT_CONTRACTS[0]
    assert overall.display_fields == ("Month", "Day")
    additive = additive_formula(overall, "Sent")
    assert "INDEX($B10#,0,1)" in additive
    assert "INDEX($B10#,0,2)" in additive
    assert additive.index("INDEX($B10#,0,1)") < additive.index("INDEX($B10#,0,2)")


def test_tracked_report_sheets_have_spilling_measure_array_flags() -> None:
    """P9 array flags remain on reconstruction XML; delivered sheets are PivotTables."""
    from dfip_web.daily_report import worksheet_xml
    from dfip_web.pivot_report import assert_native_pivot_package

    overall = worksheet_xml(REPORT_CONTRACTS[0])
    assert 't="array" ref="E10"' in overall
    assert 't="array" ref="B10:C10"' in overall
    body = XLSX_PATH.read_bytes()
    assert_native_pivot_package(body)
    text = body.decode("utf-8", errors="ignore")
    for token in FORBIDDEN_TEMPLATE_TOKENS:
        assert token not in text
    assert "DFIP_DEV_AUTH_TOKEN" not in text
