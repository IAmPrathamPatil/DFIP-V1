"""P3 multi-file + incremental monthly workflow (in-memory). Does not touch August."""

from __future__ import annotations

import hashlib
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from dfip_api.app import create_app
from dfip_api.source_storage import InMemorySourceObjectStore
from dfip_core.ingest.store import InMemoryIngestStore
from dfip_core.transform.store import InMemoryFactStore
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from http_ingest_support import (
    source_row,
    upload_workbook,
    upload_workbooks,
    workbook_bytes,
)
from test_p5_api import AUTH, CLIENT_ID
from test_p7_publication import _publish, publisher_settings

CLIENT_B = "a0000000-0000-4000-8000-000000000002"
JAN_1 = date(2026, 1, 1)
JAN_4 = date(2026, 1, 4)
JAN_5 = date(2026, 1, 5)
JAN_7 = date(2026, 1, 7)
JAN_8 = date(2026, 1, 8)
JAN_10 = date(2026, 1, 10)
JAN_11 = date(2026, 1, 11)
JAN_20 = date(2026, 1, 20)
JAN_21 = date(2026, 1, 21)
JAN_31 = date(2026, 1, 31)
FEB_1 = date(2026, 2, 1)
FEB_10 = date(2026, 2, 10)
MAR_1 = date(2026, 3, 1)
MAR_5 = date(2026, 3, 5)
ROOT = Path(__file__).resolve().parents[1]
WEB_JS = ROOT / "apps" / "web" / "static" / "js"


def _days(start: date, end: date) -> list[date]:
    values = []
    current = start
    while current <= end:
        values.append(current)
        current += timedelta(days=1)
    return values


def _iso(start: date, end: date) -> set[str]:
    return {day.isoformat() for day in _days(start, end)}


def _day_dt(value: date) -> datetime:
    return datetime(value.year, value.month, value.day)


def _with_note_sheet(content: bytes, note: str) -> bytes:
    workbook = load_workbook(BytesIO(content))
    sheet = workbook.create_sheet("archive-note")
    sheet["A1"] = note
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _range_bytes(
    tmp_path: Path,
    filename: str,
    start: date,
    end: date,
    *,
    sent: int | None = None,
    campaign_id: str = "camp-a",
    note: str | None = None,
) -> bytes:
    rows = []
    for day in _days(start, end):
        rows.append(
            source_row(
                Day=_day_dt(day),
                Sent=day.day if sent is None else sent,
                **{"Campaign ID": campaign_id},
            )
        )
    content = workbook_bytes(tmp_path / filename, rows)
    if note is not None:
        content = _with_note_sheet(content, note)
    return content


def _app(source_store=None):
    return create_app(
        settings=publisher_settings(),
        ingest_store=InMemoryIngestStore(),
        fact_store=InMemoryFactStore(),
        source_store=source_store or InMemorySourceObjectStore(),
    )


def _upload(http: TestClient, content: bytes, filename: str, *, client_id: str = CLIENT_ID):
    return upload_workbook(http, content, filename, headers=AUTH, client_id=client_id)


def _upload_many(http: TestClient, parts: list[tuple[str, bytes]], *, client_id: str = CLIENT_ID):
    return upload_workbooks(http, parts, headers=AUTH, client_id=client_id)


def _publish_current(http: TestClient, run_id: str, **extra: object):
    return _publish(
        http,
        processing_run_id=run_id,
        fact_scope="client_current",
        **extra,
    )


def _published(http: TestClient, *, client_id: str = CLIENT_ID, **params: object) -> dict:
    query = {"client_id": client_id, "limit": 200, **params}
    response = http.get("/api/v1/publications/current/facts", headers=AUTH, params=query)
    assert response.status_code == 200
    return response.json()


def _working(http: TestClient, *, client_id: str = CLIENT_ID) -> dict:
    response = http.get(
        "/api/v1/facts",
        headers=AUTH,
        params={"client_id": client_id, "limit": 200},
    )
    assert response.status_code == 200
    return response.json()


def _history(http: TestClient, *, client_id: str = CLIENT_ID) -> dict:
    response = http.get(
        "/api/v1/facts/history",
        headers=AUTH,
        params={"client_id": client_id, "limit": 200},
    )
    assert response.status_code == 200
    return response.json()


def _by_day(items: list[dict]) -> dict[str, dict]:
    return {item["day"]: item for item in items}


def _grains(items: list[dict]) -> list[tuple[str, str, str]]:
    return [(item["campaign_id"], item["variation_id_key"], item["day"]) for item in items]


def _assert_unique_grains(items: list[dict]) -> None:
    keys = _grains(items)
    assert len(keys) == len(set(keys))


def _run_id(body: dict) -> str:
    run = body.get("processing_run") or {}
    run_id = run.get("processing_run_id")
    assert run_id
    return run_id


def _process_and_publish(
    http: TestClient,
    content: bytes,
    filename: str,
    *,
    client_id: str = CLIENT_ID,
) -> dict:
    uploaded = _upload(http, content, filename, client_id=client_id)
    assert uploaded.status_code in {200, 201}
    body = uploaded.json()
    assert body["published"] is False
    published = _publish_current(http, _run_id(body), client_id=client_id)
    assert published.status_code == 201
    return body


def test_single_file_upload_contract_is_unchanged(tmp_path: Path) -> None:
    http = TestClient(_app())
    content = _range_bytes(tmp_path, "Jan_01_10.xlsx", JAN_1, JAN_10)
    response = _upload(http, content, "Jan_01_10.xlsx")
    assert response.status_code == 201
    body = response.json()
    assert "items" not in body
    assert "file_count" not in body
    assert body["original_filename"] == "Jan_01_10.xlsx"
    assert body["published"] is False
    assert body["processing_run"]["status"] == "succeeded"
    assert body["transform"]["transformed"] == 10


def test_incremental_january_february_march_sequence(tmp_path: Path) -> None:
    store = InMemorySourceObjectStore()
    app = _app(store)
    http = TestClient(app)

    jan_early = _range_bytes(tmp_path, "Jan_01_10.xlsx", JAN_1, JAN_10)
    first = _process_and_publish(http, jan_early, "Jan_01_10.xlsx")
    body = _published(http)
    assert {item["day"] for item in body["items"]} == _iso(JAN_1, JAN_10)
    assert body["pagination"]["total"] == 10
    _assert_unique_grains(body["items"])
    source = app.state.ingest_store.get_source_file(first["source_file_id"])
    assert source is not None
    assert source.storage_uri is not None
    assert store.object_count() == 1

    jan_mid = _range_bytes(tmp_path, "Jan_11_20.xlsx", JAN_11, JAN_20)
    _process_and_publish(http, jan_mid, "Jan_11_20.xlsx")
    body = _published(http)
    assert {item["day"] for item in body["items"]} == _iso(JAN_1, JAN_20)
    assert body["pagination"]["total"] == 20
    _assert_unique_grains(body["items"])

    jan_late = _range_bytes(tmp_path, "Jan_21_31.xlsx", JAN_21, JAN_31)
    _process_and_publish(http, jan_late, "Jan_21_31.xlsx")
    body = _published(http)
    assert {item["day"] for item in body["items"]} == _iso(JAN_1, JAN_31)
    assert body["pagination"]["total"] == 31
    _assert_unique_grains(body["items"])
    working = _working(http)
    assert working["pagination"]["total"] == 31

    feb = _range_bytes(tmp_path, "Feb_01_10.xlsx", FEB_1, FEB_10)
    _process_and_publish(http, feb, "Feb_01_10.xlsx")
    body = _published(http)
    assert {item["day"] for item in body["items"]} == _iso(JAN_1, JAN_31) | _iso(FEB_1, FEB_10)
    assert body["pagination"]["total"] == 41
    _assert_unique_grains(body["items"])
    working = _working(http)
    assert working["pagination"]["total"] == 41
    assert {item["day"] for item in working["items"]} == _iso(JAN_1, JAN_31) | _iso(FEB_1, FEB_10)

    march = _range_bytes(tmp_path, "Mar_01_05.xlsx", MAR_1, MAR_5)
    _process_and_publish(http, march, "Mar_01_05.xlsx")
    body = _published(http)
    expected = _iso(JAN_1, JAN_31) | _iso(FEB_1, FEB_10) | _iso(MAR_1, MAR_5)
    assert {item["day"] for item in body["items"]} == expected
    assert body["pagination"]["total"] == 46
    _assert_unique_grains(body["items"])
    assert store.object_count() == 5
    assert _history(http)["pagination"]["total"] == 0


def test_jan_5_7_correction_restates_only_those_grains(tmp_path: Path) -> None:
    http = TestClient(_app())
    _process_and_publish(http, _range_bytes(tmp_path, "Jan_01_31.xlsx", JAN_1, JAN_31), "Jan_01_31.xlsx")
    _process_and_publish(http, _range_bytes(tmp_path, "Feb_01_10.xlsx", FEB_1, FEB_10), "Feb_01_10.xlsx")
    before = _by_day(_published(http)["items"])
    jan_untouched = {day.isoformat(): before[day.isoformat()]["sent"] for day in _days(JAN_1, JAN_4)}
    jan_untouched.update(
        {day.isoformat(): before[day.isoformat()]["sent"] for day in _days(JAN_8, JAN_31)}
    )
    feb_before = {day.isoformat(): before[day.isoformat()]["sent"] for day in _days(FEB_1, FEB_10)}

    correction = _range_bytes(tmp_path, "Jan_05_07.xlsx", JAN_5, JAN_7, sent=900)
    uploaded = _process_and_publish(http, correction, "Jan_05_07.xlsx")
    after = _published(http)
    assert after["pagination"]["total"] == 41
    _assert_unique_grains(after["items"])
    by_day = _by_day(after["items"])
    for day in _days(JAN_5, JAN_7):
        assert by_day[day.isoformat()]["sent"] == 900
        assert by_day[day.isoformat()]["processing_run_id"] == _run_id(uploaded)
    for day, sent in jan_untouched.items():
        assert by_day[day]["sent"] == sent
        assert by_day[day]["processing_run_id"] != _run_id(uploaded)
    for day, sent in feb_before.items():
        assert by_day[day]["sent"] == sent
    history = _history(http)
    assert history["pagination"]["total"] == 3
    hist_days = {item["day"] for item in history["items"]}
    assert hist_days == _iso(JAN_5, JAN_7)
    for item in history["items"]:
        assert item["sent"] == date.fromisoformat(item["day"]).day
        assert item["superseded_by_run_id"] == _run_id(uploaded)
    working = _working(http)
    assert working["pagination"]["total"] == 41


def test_full_january_replacement_does_not_duplicate_or_touch_february(tmp_path: Path) -> None:
    http = TestClient(_app())
    _process_and_publish(http, _range_bytes(tmp_path, "Jan_01_31.xlsx", JAN_1, JAN_31), "Jan_01_31.xlsx")
    _process_and_publish(http, _range_bytes(tmp_path, "Feb_01_10.xlsx", FEB_1, FEB_10), "Feb_01_10.xlsx")
    replacement = _range_bytes(tmp_path, "Jan_01_31_corrected.xlsx", JAN_1, JAN_31, sent=50)
    uploaded = _process_and_publish(http, replacement, "Jan_01_31_corrected.xlsx")
    body = _published(http)
    assert body["pagination"]["total"] == 41
    _assert_unique_grains(body["items"])
    by_day = _by_day(body["items"])
    for day in _days(JAN_1, JAN_31):
        assert by_day[day.isoformat()]["sent"] == 50
        assert by_day[day.isoformat()]["processing_run_id"] == _run_id(uploaded)
    for day in _days(FEB_1, FEB_10):
        assert by_day[day.isoformat()]["sent"] == day.day
        assert by_day[day.isoformat()]["processing_run_id"] != _run_id(uploaded)
    working = _working(http)
    assert working["pagination"]["total"] == 41
    history = _history(http)
    assert history["pagination"]["total"] == 31


def test_multifile_order_and_combined_match_sequential(tmp_path: Path) -> None:
    files = {
        "Jan_01_10.xlsx": _range_bytes(tmp_path, "Jan_01_10.xlsx", JAN_1, JAN_10),
        "Jan_11_20.xlsx": _range_bytes(tmp_path, "Jan_11_20.xlsx", JAN_11, JAN_20),
        "Jan_21_31.xlsx": _range_bytes(tmp_path, "Jan_21_31.xlsx", JAN_21, JAN_31),
    }

    sequential = TestClient(_app())
    for name in ("Jan_01_10.xlsx", "Jan_11_20.xlsx", "Jan_21_31.xlsx"):
        _process_and_publish(sequential, files[name], name)
    sequential_days = {
        item["day"]: item["sent"] for item in _published(sequential)["items"]
    }
    assert set(sequential_days) == _iso(JAN_1, JAN_31)

    reversed_http = TestClient(_app())
    for name in ("Jan_21_31.xlsx", "Jan_11_20.xlsx", "Jan_01_10.xlsx"):
        uploaded = _upload(reversed_http, files[name], name)
        assert uploaded.status_code == 201
        _publish_current(reversed_http, _run_id(uploaded.json()))
    reversed_days = {
        item["day"]: item["sent"] for item in _published(reversed_http)["items"]
    }
    assert reversed_days == sequential_days

    together = TestClient(_app())
    parts = [
        ("Jan_21_31.xlsx", files["Jan_21_31.xlsx"]),
        ("Jan_01_10.xlsx", files["Jan_01_10.xlsx"]),
        ("Jan_11_20.xlsx", files["Jan_11_20.xlsx"]),
    ]
    grouped = _upload_many(together, parts)
    assert grouped.status_code == 201
    body = grouped.json()
    assert body["file_count"] == 3
    assert body["staged_row_count"] == 31
    assert body["fact_count"] == 31
    assert body["duration_ms"] is not None
    assert len(body["items"]) == 3
    names = [item["original_filename"] for item in body["items"]]
    assert names == ["Jan_01_10.xlsx", "Jan_11_20.xlsx", "Jan_21_31.xlsx"]
    source_ids = {item["source_file_id"] for item in body["items"]}
    assert len(source_ids) == 3
    shas = {item["sha256"] for item in body["items"]}
    assert len(shas) == 3
    last_run = _run_id(body["items"][-1])
    published = _publish_current(together, last_run)
    assert published.status_code == 201
    together_days = {item["day"]: item["sent"] for item in _published(together)["items"]}
    assert together_days == sequential_days
    _assert_unique_grains(_working(together)["items"])


def test_duplicate_file_in_group_does_not_duplicate_facts(tmp_path: Path) -> None:
    store = InMemorySourceObjectStore()
    http = TestClient(_app(store))
    content = _range_bytes(tmp_path, "Jan_01_10.xlsx", JAN_1, JAN_10)
    grouped = _upload_many(
        http,
        [("Jan_01_10.xlsx", content), ("Jan_01_10.xlsx", content)],
    )
    assert grouped.status_code == 201
    body = grouped.json()
    assert body["file_count"] == 2
    assert {item["source_file_id"] for item in body["items"]} == {body["items"][0]["source_file_id"]}
    working = _working(http)
    assert working["pagination"]["total"] == 10
    _assert_unique_grains(working["items"])
    assert store.object_count() == 1
    replay = _upload(http, content, "Jan_01_10.xlsx")
    assert replay.status_code == 200
    assert replay.json()["replayed"] is True
    assert replay.json()["source_file_id"] == body["items"][0]["source_file_id"]
    assert _working(http)["pagination"]["total"] == 10
    assert store.object_count() == 1


def test_overlapping_files_in_group_resolve_deterministically(tmp_path: Path) -> None:
    first = _range_bytes(tmp_path, "Jan_01_10.xlsx", JAN_1, JAN_10)
    overlap = _range_bytes(tmp_path, "Jan_05_07.xlsx", JAN_5, JAN_7, sent=900)

    def _run(parts: list[tuple[str, bytes]]) -> dict[str, int]:
        http = TestClient(_app())
        grouped = _upload_many(http, parts)
        assert grouped.status_code == 201
        body = grouped.json()
        _publish_current(http, _run_id(body["items"][-1]))
        return {item["day"]: item["sent"] for item in _published(http)["items"]}

    forward = _run([("Jan_01_10.xlsx", first), ("Jan_05_07.xlsx", overlap)])
    reverse = _run([("Jan_05_07.xlsx", overlap), ("Jan_01_10.xlsx", first)])
    assert forward == reverse
    assert set(forward) == _iso(JAN_1, JAN_10)
    for day in _days(JAN_5, JAN_7):
        assert forward[day.isoformat()] == 900
    assert forward[JAN_1.isoformat()] == 1
    assert forward[JAN_10.isoformat()] == 10


def test_same_sha_replay_is_idempotent(tmp_path: Path) -> None:
    store = InMemorySourceObjectStore()
    http = TestClient(_app(store))
    content = _range_bytes(tmp_path, "Jan_01_10.xlsx", JAN_1, JAN_10)
    first = _upload(http, content, "Jan_01_10.xlsx")
    assert first.status_code == 201
    first_body = first.json()
    _publish_current(http, _run_id(first_body))
    replay = _upload(http, content, "Jan_01_10.xlsx")
    assert replay.status_code == 200
    replayed = replay.json()
    assert replayed["replayed"] is True
    assert replayed["source_file_id"] == first_body["source_file_id"]
    assert replayed["sha256"] == first_body["sha256"]
    assert hashlib.sha256(content).hexdigest() == first_body["sha256"]
    working = _working(http)
    assert working["pagination"]["total"] == 10
    _assert_unique_grains(working["items"])
    assert store.object_count() == 1
    sources = http.get("/api/v1/source-files", headers=AUTH, params={"client_id": CLIENT_ID}).json()
    assert sources["pagination"]["total"] == 1


def test_new_sha_same_business_values_does_not_duplicate_grains(tmp_path: Path) -> None:
    store = InMemorySourceObjectStore()
    http = TestClient(_app(store))
    original = _range_bytes(tmp_path, "Jan_01_10.xlsx", JAN_1, JAN_10)
    first = _process_and_publish(http, original, "Jan_01_10.xlsx")
    retouched = _range_bytes(
        tmp_path, "Jan_01_10_note.xlsx", JAN_1, JAN_10, note=str(uuid4())
    )
    assert hashlib.sha256(retouched).hexdigest() != first["sha256"]
    second = _upload(http, retouched, "Jan_01_10_note.xlsx")
    assert second.status_code == 201
    second_body = second.json()
    assert second_body["source_file_id"] != first["source_file_id"]
    assert second_body["sha256"] != first["sha256"]
    _publish_current(http, _run_id(second_body))
    working = _working(http)
    assert working["pagination"]["total"] == 10
    _assert_unique_grains(working["items"])
    published = _published(http)
    assert published["pagination"]["total"] == 10
    _assert_unique_grains(published["items"])
    assert store.object_count() == 2
    sources = http.get("/api/v1/source-files", headers=AUTH, params={"client_id": CLIENT_ID}).json()
    assert sources["pagination"]["total"] == 2


def test_processing_run_records_logic_and_labels_versions(tmp_path: Path) -> None:
    http = TestClient(_app())
    first = _upload(http, _range_bytes(tmp_path, "Jan_01_10.xlsx", JAN_1, JAN_10), "Jan_01_10.xlsx")
    second = _upload(http, _range_bytes(tmp_path, "Jan_11_20.xlsx", JAN_11, JAN_20), "Jan_11_20.xlsx")
    grouped = _upload_many(
        TestClient(_app()),
        [
            ("Jan_01_10.xlsx", _range_bytes(tmp_path, "g1.xlsx", JAN_1, JAN_10)),
            ("Jan_11_20.xlsx", _range_bytes(tmp_path, "g2.xlsx", JAN_11, JAN_20)),
        ],
    )
    runs = [
        first.json()["processing_run"],
        second.json()["processing_run"],
        *([item["processing_run"] for item in grouped.json()["items"]]),
    ]
    for run in runs:
        assert run["campaign_label_version_id"]
        assert run["label_group_version_id"]
        assert run["template_label_version_id"]
        assert run["rate_card_version_id"]
    first_ids = (
        first.json()["processing_run"]["campaign_label_version_id"],
        first.json()["processing_run"]["label_group_version_id"],
    )
    second_ids = (
        second.json()["processing_run"]["campaign_label_version_id"],
        second.json()["processing_run"]["label_group_version_id"],
    )
    assert first_ids == second_ids
    group_ids = {
        (
            item["processing_run"]["campaign_label_version_id"],
            item["processing_run"]["label_group_version_id"],
        )
        for item in grouped.json()["items"]
    }
    assert group_ids == {first_ids}


def test_client_b_never_enters_client_a_publication(tmp_path: Path) -> None:
    http = TestClient(_app())
    _process_and_publish(http, _range_bytes(tmp_path, "Jan_01_10.xlsx", JAN_1, JAN_10), "Jan_01_10.xlsx")
    _process_and_publish(
        http,
        _range_bytes(tmp_path, "Feb_B.xlsx", FEB_1, FEB_10, campaign_id="camp-b"),
        "Feb_B.xlsx",
        client_id=CLIENT_B,
    )
    published_a = _published(http, client_id=CLIENT_ID)
    assert {item["day"] for item in published_a["items"]} == _iso(JAN_1, JAN_10)
    assert all(item["client_id"] == CLIENT_ID for item in published_a["items"])
    assert all(item["campaign_id"] != "camp-b" for item in published_a["items"])
    working_a = _working(http, client_id=CLIENT_ID)
    assert working_a["pagination"]["total"] == 10
    working_b = _working(http, client_id=CLIENT_B)
    assert working_b["pagination"]["total"] == 10
    assert all(item["campaign_id"] == "camp-b" for item in working_b["items"])


def test_client_current_period_clip_still_works(tmp_path: Path) -> None:
    http = TestClient(_app())
    _process_and_publish(http, _range_bytes(tmp_path, "Jan_01_31.xlsx", JAN_1, JAN_31), "Jan_01_31.xlsx")
    feb = _upload(http, _range_bytes(tmp_path, "Feb_01_10.xlsx", FEB_1, FEB_10), "Feb_01_10.xlsx")
    clipped = _publish_current(
        http,
        _run_id(feb.json()),
        period_start=FEB_1.isoformat(),
        period_end=FEB_10.isoformat(),
    )
    assert clipped.status_code == 201
    body = _published(http)
    assert {item["day"] for item in body["items"]} == _iso(FEB_1, FEB_10)
    assert body["pagination"]["total"] == 10
    working = _working(http)
    assert working["pagination"]["total"] == 41
    unclipped = _publish_current(http, _run_id(feb.json()))
    assert unclipped.status_code == 201
    full = _published(http)
    assert full["pagination"]["total"] == 41


def test_spa_exposes_multifile_and_cumulative_publish() -> None:
    views = (WEB_JS / "views.js").read_text(encoding="utf-8")
    app_js = (WEB_JS / "app.js").read_text(encoding="utf-8")
    client_js = (WEB_JS / "api-client.js").read_text(encoding="utf-8")
    assert "multiple" in views
    assert 'name: "files"' in views or 'name="files"' in views
    assert "client_current" in views
    assert "fact_scope" in app_js
    assert 'body.append("files"' in client_js
    assert "The website stays usable while this file is processed." in views
