"""P7 publication API, client scoping, and Excel/Power Query template tests."""

from __future__ import annotations

from datetime import date
from pathlib import Path
from uuid import UUID
from zipfile import ZipFile

from dfip_api.app import create_app
from dfip_api.publication_store import InMemoryPublicationStore
from dfip_api.roles import can_publish
from dfip_config.settings import Settings
from dfip_core.transform.fact import FactRecord
from dfip_web.api_client import DfipApiClient
from dfip_web.client_workbook import (
    BUILDER_CREATES_NATIVE_DATAMASHUP,
    FACT_HEADERS,
    FAKE_MASHUP_ZIP_PARTS,
    M_PATH,
    SETTINGS_ROWS,
    XLSX_PATH,
    build_client_report,
    mashup_text,
)
from dfip_web.daily_report import CLIENT_WORKBOOK_SHEET_NAMES
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from test_p5_api import (
    AUTH,
    CLIENT_ID,
    DEV_TOKEN,
    JWT_SECRET,
    MISSING_ID,
    RUN_A,
    RUN_B,
    RUN_C,
    _encode_jwt,
    make_settings,
    seed_stores,
)

ROOT = Path(__file__).resolve().parents[1]
CLIENT_B = "a0000000-0000-4000-8000-000000000002"
WEB_STATIC = ROOT / "apps" / "web" / "static"
PUBLISHED_FACTS_PATH = "/api/v1/publications/current/facts"
WORKING_SET_FACTS_PATH = "/api/v1/facts"
FORBIDDEN_TEMPLATE_TOKENS = (
    "DATABASE_URL",
    "DFIP_AUTH_SECRET",
    "SUPABASE_SERVICE_ROLE_KEY",
    "postgresql://",
    "stg_source_row",
    "staged-rows",
    "/source-files",
    "storage_uri",
)


def _omits_working_set_facts_url(text: str) -> bool:
    return WORKING_SET_FACTS_PATH not in text.replace(PUBLISHED_FACTS_PATH, "")


def assert_no_fake_mashup_parts(xlsx_path: Path) -> tuple[set[str], str, str]:
    with ZipFile(xlsx_path) as archive:
        names = set(archive.namelist())
        connections = (
            archive.read("xl/connections.xml").decode("utf-8")
            if "xl/connections.xml" in names
            else ""
        )
        item_props = (
            archive.read("customXml/itemProps1.xml").decode("utf-8")
            if "customXml/itemProps1.xml" in names
            else ""
        )
    for part in FAKE_MASHUP_ZIP_PARTS:
        assert part not in names
    assert not any(name.startswith("xl/queryMashup/") for name in names)
    return names, connections, item_props


def publisher_settings(**overrides: object) -> Settings:
    values: dict[str, object] = {"dfip_dev_auth_role": "publisher"}
    values.update(overrides)
    return make_settings(**values)


def publisher_app(**settings_overrides: object):
    ingest, facts = seed_stores()
    store = InMemoryPublicationStore()
    app = create_app(
        settings=publisher_settings(**settings_overrides),
        ingest_store=ingest,
        fact_store=facts,
        publication_store=store,
    )
    return app, ingest, facts, store


def jwt_app(role: str = "reader", client_id: str | None = CLIENT_ID):
    ingest, facts = seed_stores()
    store = InMemoryPublicationStore()
    app = create_app(
        settings=make_settings(dfip_auth_mode="jwt", dfip_auth_secret=JWT_SECRET),
        ingest_store=ingest,
        fact_store=facts,
        publication_store=store,
    )
    token = _encode_jwt(role=role, client_id=client_id)
    return app, store, facts, {"Authorization": f"Bearer {token}"}


def _error(response) -> dict:
    payload = response.json()
    assert "error" in payload
    return payload["error"]


def _publish(
    client: TestClient,
    *,
    headers: dict[str, str] | None = None,
    client_id: str = CLIENT_ID,
    processing_run_id: str = RUN_A,
    **extra: object,
):
    body = {"client_id": client_id, "processing_run_id": processing_run_id, **extra}
    return client.post("/api/v1/publications", headers=headers or AUTH, json=body)


def test_admin_and_publisher_can_publish() -> None:
    assert can_publish("admin")
    assert can_publish("publisher")
    assert not can_publish("reader")
    assert not can_publish("client")
    assert not can_publish(None)


def test_empty_current_publication_is_safe() -> None:
    app, _ingest, _facts, store = publisher_app()
    client = TestClient(app)
    current = client.get(
        "/api/v1/publications/current",
        headers=AUTH,
        params={"client_id": CLIENT_ID},
    )
    assert current.status_code == 200
    assert current.json() == {"publication": None, "current": None}
    facts = client.get(
        "/api/v1/publications/current/facts",
        headers=AUTH,
        params={"client_id": CLIENT_ID},
    )
    assert facts.status_code == 200
    body = facts.json()
    assert body["items"] == []
    assert body["pagination"] == {"limit": 50, "offset": 0, "total": 0}
    assert store.get_current(CLIENT_ID) == (None, None)


def test_reader_cannot_publish() -> None:
    ingest, facts = seed_stores()
    app = create_app(
        settings=make_settings(dfip_dev_auth_role="reader"),
        ingest_store=ingest,
        fact_store=facts,
    )
    response = TestClient(app).post(
        "/api/v1/publications",
        headers=AUTH,
        json={"client_id": CLIENT_ID, "processing_run_id": RUN_A},
    )
    assert response.status_code == 403
    assert _error(response)["code"] == "AUTHORIZATION_FAILED"


def test_client_role_cannot_publish() -> None:
    app, _store, _facts, headers = jwt_app(role="client", client_id=CLIENT_ID)
    response = TestClient(app).post(
        "/api/v1/publications",
        headers=headers,
        json={"client_id": CLIENT_ID, "processing_run_id": RUN_A},
    )
    assert response.status_code == 403


def test_publisher_creates_publication_and_current_pointer() -> None:
    app, _ingest, _facts, store = publisher_app()
    response = _publish(TestClient(app))
    assert response.status_code == 201
    body = response.json()
    publication = body["publication"]
    current = body["current"]
    assert publication["client_id"] == CLIENT_ID
    assert publication["processing_run_id"] == RUN_A
    assert publication["fact_scope"] == "processing_run"
    assert publication["period_start"] is None
    assert current["publication_id"] == publication["publication_id"]
    assert current["client_id"] == CLIENT_ID
    record, pointer = store.get_current(CLIENT_ID)
    assert record is not None
    assert pointer is not None
    assert record.id == publication["publication_id"]


def test_nonexistent_processing_run_cannot_be_published() -> None:
    app, *_rest = publisher_app()
    response = _publish(TestClient(app), processing_run_id=MISSING_ID)
    assert response.status_code == 404
    assert _error(response)["code"] == "NOT_FOUND"


def test_failed_processing_run_cannot_be_published() -> None:
    app, *_rest = publisher_app()
    response = _publish(TestClient(app), processing_run_id=RUN_B)
    assert response.status_code == 422
    assert _error(response)["code"] == "VALIDATION_ERROR"
    pending = _publish(TestClient(app), processing_run_id=RUN_C)
    assert pending.status_code == 422


def test_invalid_period_is_rejected() -> None:
    app, *_rest = publisher_app()
    response = _publish(
        TestClient(app),
        period_start="2025-08-10",
        period_end="2025-08-01",
    )
    assert response.status_code == 422
    assert _error(response)["code"] == "VALIDATION_ERROR"


def test_republish_replaces_current_and_keeps_history() -> None:
    app, _ingest, _facts, store = publisher_app()
    client = TestClient(app)
    first = _publish(client, notes="first")
    second = _publish(client, notes="second", period_start="2025-08-01")
    assert first.status_code == 201
    assert second.status_code == 201
    first_id = first.json()["publication"]["publication_id"]
    second_id = second.json()["publication"]["publication_id"]
    assert first_id != second_id
    current = client.get(
        "/api/v1/publications/current",
        headers=AUTH,
        params={"client_id": CLIENT_ID},
    ).json()
    assert current["publication"]["publication_id"] == second_id
    assert current["publication"]["notes"] == "second"
    history = store.list_for_client(CLIENT_ID)
    assert [item.id for item in history] == [first_id, second_id]
    listed = client.get(
        "/api/v1/publications",
        headers=AUTH,
        params={"client_id": CLIENT_ID},
    )
    assert listed.status_code == 200
    ids = [item["publication_id"] for item in listed.json()["items"]]
    assert ids == [second_id, first_id]
    assert listed.json()["pagination"]["total"] == 2


def test_published_facts_match_the_published_run_only() -> None:
    app, _ingest, facts, _store = publisher_app()
    other = FactRecord(
        client_id=CLIENT_ID,
        campaign_id="camp-other-run",
        variation_id="var-z",
        variation_id_key="var-z",
        day=date(2025, 8, 4),
        processing_run_id=RUN_C,
        total_cost=None,
        template_status="",
    )
    facts.facts[other.key] = other
    client = TestClient(app)
    working = client.get("/api/v1/facts", headers=AUTH).json()
    assert working["pagination"]["total"] == 4
    _publish(client)
    published = client.get(
        "/api/v1/publications/current/facts",
        headers=AUTH,
        params={"client_id": CLIENT_ID},
    ).json()
    assert published["pagination"]["total"] == 3
    assert {item["campaign_id"] for item in published["items"]} == {"camp-1", "camp-2", "camp-3"}
    assert all(item["processing_run_id"] == RUN_A for item in published["items"])
    still_working = client.get("/api/v1/facts", headers=AUTH).json()
    assert still_working["pagination"]["total"] == 4


def test_published_facts_preserve_null_empty_and_decimal_strings() -> None:
    app, *_rest = publisher_app()
    client = TestClient(app)
    _publish(client)
    items = client.get(
        "/api/v1/publications/current/facts",
        headers=AUTH,
        params={"client_id": CLIENT_ID, "limit": 200},
    ).json()["items"]
    camp1 = next(item for item in items if item["campaign_id"] == "camp-1")
    assert camp1["filter_logic_1"] is None
    assert camp1["template_status"] == ""
    assert camp1["total_cost"] == "2.50"
    assert isinstance(camp1["total_cost"], str)
    camp2 = next(item for item in items if item["campaign_id"] == "camp-2")
    assert camp2["variation_id"] is None
    camp3 = next(item for item in items if item["campaign_id"] == "camp-3")
    assert camp3["total_cost"] is None


def test_period_filtering_on_published_facts() -> None:
    app, *_rest = publisher_app()
    client = TestClient(app)
    _publish(client, period_start="2025-08-01", period_end="2025-08-01")
    body = client.get(
        "/api/v1/publications/current/facts",
        headers=AUTH,
        params={"client_id": CLIENT_ID},
    ).json()
    assert body["pagination"]["total"] == 1
    assert body["items"][0]["campaign_id"] == "camp-1"
    assert body["items"][0]["day"] == "2025-08-01"


def test_published_facts_paginate_without_raising_global_limit() -> None:
    app, _ingest, facts, _store = publisher_app()
    for index in range(201):
        record = FactRecord(
            client_id=CLIENT_ID,
            campaign_id=f"extra-{index:03d}",
            variation_id="var-x",
            variation_id_key="var-x",
            day=date(2025, 9, 1),
            processing_run_id=RUN_A,
        )
        facts.facts[record.key] = record
    client = TestClient(app)
    _publish(client)
    rejected = client.get(
        "/api/v1/publications/current/facts",
        headers=AUTH,
        params={"client_id": CLIENT_ID, "limit": 201},
    )
    assert rejected.status_code == 422
    first = client.get(
        "/api/v1/publications/current/facts",
        headers=AUTH,
        params={"client_id": CLIENT_ID, "limit": 200, "offset": 0},
    ).json()
    second = client.get(
        "/api/v1/publications/current/facts",
        headers=AUTH,
        params={"client_id": CLIENT_ID, "limit": 200, "offset": 200},
    ).json()
    assert first["pagination"]["limit"] == 200
    assert first["pagination"]["total"] == 204
    assert len(first["items"]) == 200
    assert len(second["items"]) == 4
    working = client.get("/api/v1/facts?limit=200", headers=AUTH).json()
    assert working["pagination"]["total"] == 204


def test_jwt_client_cannot_read_another_clients_publication() -> None:
    ingest, facts = seed_stores()
    store = InMemoryPublicationStore()
    publisher = create_app(
        settings=publisher_settings(),
        ingest_store=ingest,
        fact_store=facts,
        publication_store=store,
    )
    reader = create_app(
        settings=make_settings(dfip_auth_mode="jwt", dfip_auth_secret=JWT_SECRET),
        ingest_store=ingest,
        fact_store=facts,
        publication_store=store,
    )
    created = _publish(TestClient(publisher))
    assert created.status_code == 201
    token = _encode_jwt(role="client", client_id=CLIENT_B)
    headers = {"Authorization": f"Bearer {token}"}
    http = TestClient(reader)
    scoped = http.get("/api/v1/publications/current/facts", headers=headers)
    assert scoped.status_code == 200
    assert scoped.json()["items"] == []
    assert scoped.json()["pagination"]["total"] == 0
    denied = http.get(
        "/api/v1/publications/current/facts",
        headers=headers,
        params={"client_id": CLIENT_ID},
    )
    assert denied.status_code == 403


def test_jwt_with_client_id_scopes_published_facts() -> None:
    app, _store, _facts, headers = jwt_app(role="publisher", client_id=CLIENT_ID)
    http = TestClient(app)
    created = http.post(
        "/api/v1/publications",
        headers=headers,
        json={"client_id": CLIENT_ID, "processing_run_id": RUN_A},
    )
    assert created.status_code == 201
    published = http.get("/api/v1/publications/current/facts", headers=headers).json()
    assert published["pagination"]["total"] == 3
    mismatch = http.post(
        "/api/v1/publications",
        headers=headers,
        json={"client_id": CLIENT_B, "processing_run_id": RUN_A},
    )
    assert mismatch.status_code == 403


def test_dev_token_requires_explicit_client_id_and_is_not_rls() -> None:
    app, *_rest = publisher_app()
    client = TestClient(app)
    empty = client.get("/api/v1/publications/current/facts", headers=AUTH).json()
    assert empty["pagination"]["total"] == 0
    _publish(client)
    filled = client.get(
        "/api/v1/publications/current/facts",
        headers=AUTH,
        params={"client_id": CLIENT_ID},
    ).json()
    assert filled["pagination"]["total"] == 3


def test_python_api_client_covers_publication_routes() -> None:
    app, *_rest = publisher_app()
    with DfipApiClient("http://testserver", token=DEV_TOKEN, http_client=TestClient(app)) as api:
        created = api.create_publication({"client_id": CLIENT_ID, "processing_run_id": RUN_A})
        current = api.get_current_publication(client_id=CLIENT_ID)
        facts = api.list_published_facts(client_id=CLIENT_ID, limit=200)
        working = api.list_facts()
    assert created["publication"]["processing_run_id"] == RUN_A
    assert current["publication"]["publication_id"] == created["publication"]["publication_id"]
    assert facts["pagination"]["total"] == 3
    assert working["pagination"]["total"] == 3


def test_publication_routes_require_auth() -> None:
    app, *_rest = publisher_app()
    client = TestClient(app)
    assert client.post("/api/v1/publications", json={}).status_code == 401
    assert client.get("/api/v1/publications").status_code == 401
    assert client.get("/api/v1/publications/current").status_code == 401
    assert client.get("/api/v1/publications/current/facts").status_code == 401


def test_frontend_publish_control_is_enabled() -> None:
    views = (WEB_STATIC / "js" / "views.js").read_text(encoding="utf-8")
    app_js = (WEB_STATIC / "js" / "app.js").read_text(encoding="utf-8")
    client_js = (WEB_STATIC / "js" / "api-client.js").read_text(encoding="utf-8")
    assert 'data-publish-form="true"' in views
    assert "Publish (not implemented)" not in views
    assert "createPublication" in app_js
    assert "createPublication" in client_js
    assert "listPublications" in client_js
    assert "/publications" in client_js
    assert "JSON.stringify" in client_js


def test_power_query_pages_published_facts_and_avoids_staging() -> None:
    assert M_PATH.is_file()
    mashup = mashup_text()
    disk = M_PATH.read_text(encoding="utf-8")
    assert mashup == disk
    assert PUBLISHED_FACTS_PATH in mashup
    assert _omits_working_set_facts_url(mashup)
    assert "200" in mashup
    assert "offset" in mashup
    assert "Authorization" in mashup
    assert "Bearer" in mashup
    assert "stg_source_row" not in mashup
    assert "staged-rows" not in mashup
    for token in FORBIDDEN_TEMPLATE_TOKENS:
        assert token not in mashup
    assert "not RLS" in mashup


def test_client_workbook_is_native_v2x_without_fake_mashup() -> None:
    """Tracked workbook is Desktop-native V2-X, not the Python structure-only scaffold."""
    assert M_PATH.is_file()
    assert XLSX_PATH.is_file()
    names, connections, item_props = assert_no_fake_mashup_parts(XLSX_PATH)
    assert "xl/connections.xml" in names
    assert any(name.startswith("xl/queryTables/") for name in names)
    assert "Microsoft.Mashup.OleDb.1" in connections
    assert "Location=PublishedFacts" in connections
    assert "http://schemas.microsoft.com/DataMashup" in item_props
    workbook = load_workbook(XLSX_PATH, read_only=True, data_only=False)
    assert workbook.sheetnames == list(CLIENT_WORKBOOK_SHEET_NAMES)
    sheet = workbook["Facts"]
    assert sheet["A1"].value == "Parameter"
    assert sheet["B1"].value == "Value"
    for index, (key, value) in enumerate(SETTINGS_ROWS, start=2):
        assert sheet.cell(index, 1).value == key
        actual = sheet.cell(index, 2).value
        if key == "BearerToken":
            assert actual in {None, ""}
        elif key == "ClientId":
            if actual not in {None, ""}:
                UUID(str(actual))
        else:
            assert actual == value
    headers = [sheet.cell(6, column).value for column in range(1, len(FACT_HEADERS) + 1)]
    assert headers == list(FACT_HEADERS)
    workbook.close()
    raw = XLSX_PATH.read_bytes()
    assert _omits_working_set_facts_url(raw.decode("latin-1"))
    for token in FORBIDDEN_TEMPLATE_TOKENS:
        assert token.encode("utf-8") not in raw
    assert b"p5-test-jwt-secret" not in raw
    assert b"p5-test-dev-token" not in raw
    assert b"SUPABASE_SERVICE_ROLE_KEY" not in raw


def test_builder_does_not_embed_fake_mashup(tmp_path: Path) -> None:
    assert BUILDER_CREATES_NATIVE_DATAMASHUP is False
    target = tmp_path / "Client_Report.xlsx"
    built = build_client_report(target)
    assert built == target
    assert_no_fake_mashup_parts(target)


def test_p7_does_not_add_kpi_qa_or_rls() -> None:
    api_root = ROOT / "packages" / "api" / "dfip_api"
    combined = ""
    for path in api_root.rglob("*.py"):
        combined += path.read_text(encoding="utf-8").lower()
    assert "kpi engine" not in combined
    assert "qa engine" not in combined
    assert "enable row level security" not in combined
    assert "calculate_total_cost" not in combined
    mashup = mashup_text().lower()
    assert "new logic" not in mashup
    assert "vlookup" not in mashup
