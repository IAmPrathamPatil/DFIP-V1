"""P8 authenticated nine-sheet Client Report download and historical recovery."""

from __future__ import annotations

import io
import time
from dataclasses import replace
from datetime import date
from decimal import Decimal
from zipfile import ZipFile

from dfip_api.app import create_app
from dfip_api.publication_store import SNAPSHOT_STATUS_COMPLETE, InMemoryPublicationStore
from dfip_core.ingest.store import BatchRecord, ProcessingRunRecord
from dfip_core.transform.fact import FactRecord
from dfip_web.api_client import DfipApiClient
from dfip_web.client_report_download import (
    CLIENT_REPORT_DOWNLOAD_NAME,
    CLIENT_REPORT_MEDIA_TYPE,
    render_client_report_xlsx,
)
from dfip_web.client_workbook import FAKE_MASHUP_ZIP_PARTS, XLSX_PATH
from dfip_web.daily_report import (
    CLIENT_WORKBOOK_SHEET_NAMES,
    REPORT_SHEET_NAMES,
    _sheet_part_map,
)
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from test_p5_api import (
    AUTH,
    BATCH_A,
    CLIENT_ID,
    DEV_TOKEN,
    FILE_A,
    JWT_SECRET,
    RUN_A,
    _encode_jwt,
    _ts,
    make_settings,
    seed_stores,
)
from test_p7_publication import (
    CLIENT_B,
    FORBIDDEN_TEMPLATE_TOKENS,
    WEB_STATIC,
    _error,
    _publish,
    jwt_app,
    publisher_app,
    publisher_settings,
)

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
RUN_D = "d0000000-0000-4000-8000-00000000000d"
BATCH_CLIENT_B = "c0000000-0000-4000-8000-0000000000b1"
RUN_CLIENT_B = "d0000000-0000-4000-8000-0000000000b1"
SECRET_BYTES = (*FORBIDDEN_TEMPLATE_TOKENS, DEV_TOKEN, JWT_SECRET, "DFIP_DEV_AUTH_TOKEN")


def _client_report_path(publication_id: str | None = None) -> str:
    if publication_id:
        return f"/api/v1/publications/{publication_id}/client-report.xlsx"
    return "/api/v1/publications/current/client-report.xlsx"


def _sheet_names(body: bytes) -> list[str]:
    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=False)
    names = list(workbook.sheetnames)
    workbook.close()
    return names


def _campaign_ids(body: bytes) -> list[str]:
    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=False)
    sheet = workbook["PublishedFacts"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    index = headers.index("Campaign ID")
    values = [
        str(row[index])
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if row[index] not in {None, ""}
    ]
    workbook.close()
    return values


def _settings_value(body: bytes, parameter: str) -> object:
    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=False)
    sheet = workbook["Facts"]
    found = None
    for row in sheet.iter_rows(min_row=1, max_row=5, max_col=2, values_only=True):
        if row[0] == parameter:
            found = row[1]
            break
    workbook.close()
    return found


def _scan_workbook_text(body: bytes) -> str:
    chunks: list[str] = []
    with ZipFile(io.BytesIO(body)) as archive:
        for name in archive.namelist():
            chunks.append(archive.read(name).decode("utf-8", errors="ignore"))
    return "\n".join(chunks)


def _add_later_run(ingest, facts) -> None:
    ingest.processing_runs[RUN_D] = ProcessingRunRecord(
        id=RUN_D,
        batch_id=BATCH_A,
        campaign_label_version_id="a0000000-0000-4000-8000-000000000022",
        template_label_version_id="a0000000-0000-4000-8000-000000000034",
        rate_card_version_id="a0000000-0000-4000-8000-000000000012",
        label_group_version_id="a0000000-0000-4000-8000-000000000041",
        engine_version="0.4.0",
        started_at=_ts(5),
        finished_at=_ts(6),
        status="succeeded",
        qa_verdict="pass",
    )
    later = FactRecord(
        client_id=CLIENT_ID,
        campaign_id="camp-hist-b",
        variation_id="var-b",
        variation_id_key="var-b",
        day=date(2025, 9, 1),
        campaign_name="later-publication",
        processing_run_id=RUN_D,
        batch_id=BATCH_A,
        total_cost=Decimal("9.99"),
        first_seen_at=_ts(5),
        last_seen_at=_ts(5),
    )
    facts.facts[later.key] = later


def _seed_client_b(ingest, facts) -> None:
    ingest.batches[BATCH_CLIENT_B] = BatchRecord(
        id=BATCH_CLIENT_B,
        source_file_id=FILE_A,
        client_id=CLIENT_B,
        status="processed",
        row_count_declared=1,
        row_count_staged=1,
        row_count_rejected=0,
        observed_day_min=date(2025, 8, 1),
        observed_day_max=date(2025, 8, 1),
        created_at=_ts(7),
        completed_at=_ts(8),
        error_summary=None,
        worksheet_name="Web-Engage Raw",
        header_row=1,
        source_start_column="K",
        empty_row_count=0,
    )
    ingest.processing_runs[RUN_CLIENT_B] = ProcessingRunRecord(
        id=RUN_CLIENT_B,
        batch_id=BATCH_CLIENT_B,
        campaign_label_version_id="a0000000-0000-4000-8000-000000000022",
        template_label_version_id="a0000000-0000-4000-8000-000000000034",
        rate_card_version_id="a0000000-0000-4000-8000-000000000012",
        label_group_version_id="a0000000-0000-4000-8000-000000000041",
        engine_version="0.4.0",
        started_at=_ts(7),
        finished_at=_ts(8),
        status="succeeded",
        qa_verdict="pass",
    )
    foreign = FactRecord(
        client_id=CLIENT_B,
        campaign_id="camp-b-only",
        variation_id="var-b",
        variation_id_key="var-b",
        day=date(2025, 8, 1),
        campaign_name="client-b",
        processing_run_id=RUN_CLIENT_B,
        batch_id=BATCH_CLIENT_B,
        total_cost=Decimal("77.00"),
        first_seen_at=_ts(7),
        last_seen_at=_ts(7),
    )
    facts.facts[foreign.key] = foreign


def _two_publications():
    ingest, facts = seed_stores()
    store = InMemoryPublicationStore()
    app = create_app(
        settings=publisher_settings(),
        ingest_store=ingest,
        fact_store=facts,
        publication_store=store,
    )
    http = TestClient(app)
    first = _publish(http)
    assert first.status_code == 201
    pub_a = first.json()["publication"]["publication_id"]
    _add_later_run(ingest, facts)
    second = _publish(http, processing_run_id=RUN_D)
    assert second.status_code == 201
    pub_b = second.json()["publication"]["publication_id"]
    return http, ingest, facts, store, pub_a, pub_b


def test_authenticated_client_downloads_current_nine_sheet_report() -> None:
    app, _store, _facts, headers = jwt_app(role="client", client_id=CLIENT_ID)
    http = TestClient(app)
    created = http.post(
        "/api/v1/publications",
        headers={"Authorization": f"Bearer {_encode_jwt(role='publisher', client_id=CLIENT_ID)}"},
        json={"client_id": CLIENT_ID, "processing_run_id": RUN_A},
    )
    assert created.status_code == 201
    started = time.perf_counter()
    response = http.get(_client_report_path(), headers=headers)
    elapsed_ms = (time.perf_counter() - started) * 1000
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(XLSX_TYPE)
    assert 'attachment; filename="Client_Report.xlsx"' in response.headers["content-disposition"]
    body = response.content
    assert _sheet_names(body) == list(CLIENT_WORKBOOK_SHEET_NAMES)
    assert _sheet_names(body)[2:] == list(REPORT_SHEET_NAMES)
    campaigns = _campaign_ids(body)
    assert campaigns == ["camp-1", "camp-2", "camp-3"]
    assert elapsed_ms < 15_000
    assert len(body) > 10_000
    assert len(campaigns) == 3


def test_unauthenticated_client_report_is_rejected() -> None:
    app, *_rest = publisher_app()
    http = TestClient(app)
    _publish(http)
    response = http.get(_client_report_path(), params={"client_id": CLIENT_ID})
    assert response.status_code == 401
    assert _error(response)["code"] == "AUTHENTICATION_FAILED"
    assert not response.content.startswith(b"PK")


def test_client_a_cannot_download_client_b_current_or_historical_report() -> None:
    ingest, facts = seed_stores()
    _seed_client_b(ingest, facts)
    store = InMemoryPublicationStore()
    publisher = create_app(
        settings=publisher_settings(),
        ingest_store=ingest,
        fact_store=facts,
        publication_store=store,
    )
    pub_http = TestClient(publisher)
    created_a = _publish(pub_http)
    created_b = _publish(pub_http, client_id=CLIENT_B, processing_run_id=RUN_CLIENT_B)
    assert created_a.status_code == 201
    assert created_b.status_code == 201
    pub_b = created_b.json()["publication"]["publication_id"]
    reader = create_app(
        settings=make_settings(dfip_auth_mode="jwt", dfip_auth_secret=JWT_SECRET),
        ingest_store=ingest,
        fact_store=facts,
        publication_store=store,
    )
    headers_a = {"Authorization": f"Bearer {_encode_jwt(role='client', client_id=CLIENT_ID)}"}
    http = TestClient(reader)
    denied_current = http.get(
        _client_report_path(),
        headers=headers_a,
        params={"client_id": CLIENT_B},
    )
    assert denied_current.status_code == 403
    denied_hist = http.get(_client_report_path(pub_b), headers=headers_a)
    assert denied_hist.status_code == 404
    assert _error(denied_hist)["code"] == "NOT_FOUND"
    assert not denied_hist.content.startswith(b"PK")
    own = http.get(_client_report_path(), headers=headers_a)
    assert own.status_code == 200
    assert "camp-b-only" not in _campaign_ids(own.content)


def test_historical_report_stays_bound_after_later_publish() -> None:
    http, ingest, facts, store, pub_a, pub_b = _two_publications()
    current = store.get_current(CLIENT_ID)[0]
    assert current is not None
    assert current.id == pub_b
    params = {"client_id": CLIENT_ID}
    hist = http.get(_client_report_path(pub_a), headers=AUTH, params=params)
    latest = http.get(_client_report_path(pub_b), headers=AUTH, params=params)
    current_dl = http.get(_client_report_path(), headers=AUTH, params=params)
    assert hist.status_code == 200
    assert latest.status_code == 200
    assert current_dl.status_code == 200
    hist_ids = _campaign_ids(hist.content)
    latest_ids = _campaign_ids(latest.content)
    assert hist_ids == ["camp-1", "camp-2", "camp-3"]
    assert latest_ids == ["camp-hist-b"]
    assert _campaign_ids(current_dl.content) == latest_ids
    assert store.get_current(CLIENT_ID)[0].id == pub_b
    again = http.get(_client_report_path(pub_a), headers=AUTH, params=params)
    assert again.content == hist.content


def test_complete_snapshot_is_used_for_historical_report(monkeypatch) -> None:
    http, ingest, facts, store, pub_a, pub_b = _two_publications()
    record = store.get(pub_a)
    assert record is not None
    assert record.snapshot_status == SNAPSHOT_STATUS_COMPLETE
    live = facts.facts
    for key, item in list(live.items()):
        if item.processing_run_id == RUN_A:
            live[key] = replace(item, campaign_name="mutated-live", total_cost=Decimal("999"))

    def boom(*_args, **_kwargs):
        raise AssertionError("historical complete report must not read the live slice")

    monkeypatch.setattr(facts, "list_published_slice", boom)
    response = http.get(
        _client_report_path(pub_a),
        headers=AUTH,
        params={"client_id": CLIENT_ID},
    )
    assert response.status_code == 200
    assert _campaign_ids(response.content) == ["camp-1", "camp-2", "camp-3"]
    assert "mutated-live" not in _scan_workbook_text(response.content)
    assert "camp-hist-b" not in _campaign_ids(response.content)


def test_download_does_not_mutate_publication_facts_or_qa() -> None:
    http, ingest, facts, store, pub_a, pub_b = _two_publications()
    before_current = store.get_current(CLIENT_ID)
    before_snaps = {key: list(value) for key, value in store.snapshots.items()}
    before_facts = {key: item for key, item in facts.facts.items()}
    before_runs = {key: (run.status, run.qa_verdict) for key, run in ingest.processing_runs.items()}
    params = {"client_id": CLIENT_ID}
    assert http.get(_client_report_path(), headers=AUTH, params=params).status_code == 200
    assert http.get(_client_report_path(pub_a), headers=AUTH, params=params).status_code == 200
    assert store.get_current(CLIENT_ID) == before_current
    assert {key: list(value) for key, value in store.snapshots.items()} == before_snaps
    assert facts.facts == before_facts
    assert {
        key: (run.status, run.qa_verdict) for key, run in ingest.processing_runs.items()
    } == before_runs
    assert pub_b == before_current[0].id


def test_download_does_not_reprocess_and_fact_csv_xlsx_remain() -> None:
    http, ingest, _facts, _store, pub_a, _pub_b = _two_publications()
    run_ids = set(ingest.processing_runs)
    params = {"client_id": CLIENT_ID}
    report = http.get(_client_report_path(), headers=AUTH, params=params)
    csv_body = http.get("/api/v1/publications/current/facts.csv", headers=AUTH, params=params)
    xlsx_body = http.get("/api/v1/publications/current/facts.xlsx", headers=AUTH, params=params)
    facts_json = http.get("/api/v1/publications/current/facts", headers=AUTH, params=params)
    hist_json = http.get(f"/api/v1/publications/{pub_a}/facts", headers=AUTH, params=params)
    assert report.status_code == 200
    assert csv_body.status_code == 200
    assert xlsx_body.status_code == 200
    assert facts_json.status_code == 200
    assert hist_json.status_code == 200
    assert "text/csv" in csv_body.headers["content-type"]
    assert "published-facts-" in csv_body.headers["content-disposition"]
    assert csv_body.text.splitlines()[0].startswith("client_id,")
    assert set(ingest.processing_runs) == run_ids


def test_workbook_has_required_sheets_and_no_auth_secrets() -> None:
    http, *_rest, pub_a, _pub_b = _two_publications()
    response = http.get(
        _client_report_path(pub_a),
        headers=AUTH,
        params={"client_id": CLIENT_ID},
    )
    body = response.content
    assert _sheet_names(body) == list(CLIENT_WORKBOOK_SHEET_NAMES)
    with ZipFile(io.BytesIO(body)) as archive:
        names = set(archive.namelist())
        for fake in FAKE_MASHUP_ZIP_PARTS:
            assert fake not in names
        connections = archive.read("xl/connections.xml").decode("utf-8")
    assert 'keepAlive="0"' in connections
    assert 'refreshOnLoad="0"' in connections
    text = _scan_workbook_text(body)
    for token in SECRET_BYTES:
        assert token not in text
    assert _settings_value(body, "BearerToken") in {None, ""}
    assert _settings_value(body, "ClientId") in {None, ""}
    assert _settings_value(body, "ApiBaseUrl") in {None, ""}
    published = _settings_value(body, "Published")
    assert published not in {None, ""}
    workbook_xml = ZipFile(io.BytesIO(body)).read("xl/workbook.xml").decode("utf-8")
    for name in REPORT_SHEET_NAMES:
        assert f'name="{name}"' in workbook_xml


def test_failed_generate_does_not_return_partial_workbook(monkeypatch) -> None:
    app, *_rest = publisher_app()
    http = TestClient(app)
    _publish(http)

    def boom(*_args, **_kwargs):
        raise ValueError("broken template")

    monkeypatch.setattr(
        "dfip_api.publication_service.render_client_report_xlsx",
        boom,
    )
    response = http.get(
        _client_report_path(),
        headers=AUTH,
        params={"client_id": CLIENT_ID},
    )
    assert response.status_code == 500
    assert response.headers["content-type"].startswith("application/json")
    assert not response.content.startswith(b"PK")
    assert _error(response)["code"] == "INTERNAL_ERROR"
    assert "broken template" not in response.text


def test_python_and_spa_client_report_wiring() -> None:
    app, *_rest = publisher_app()
    with DfipApiClient("http://testserver", token=DEV_TOKEN, http_client=TestClient(app)) as api:
        created = api.create_publication({"client_id": CLIENT_ID, "processing_run_id": RUN_A})
        body = api.download_client_report(client_id=CLIENT_ID)
        hist = api.download_client_report(
            publication_id=created["publication"]["publication_id"],
            client_id=CLIENT_ID,
        )
    assert _sheet_names(body) == list(CLIENT_WORKBOOK_SHEET_NAMES)
    assert hist == body
    views = (WEB_STATIC / "js" / "views.js").read_text(encoding="utf-8")
    app_js = (WEB_STATIC / "js" / "app.js").read_text(encoding="utf-8")
    client_js = (WEB_STATIC / "js" / "api-client.js").read_text(encoding="utf-8")
    assert "data-download-client-report" in views
    assert "downloadClientReport" in app_js
    assert "downloadClientReport" in client_js
    assert "client-report.xlsx" in client_js
    assert "excel export" not in views.lower()
    assert "excel export" not in app_js.lower()
    assert "excel export" not in client_js.lower()
    assert "DFIP_DEV_AUTH_TOKEN" not in views
    assert "DFIP_DEV_AUTH_TOKEN" not in app_js
    assert "DFIP_DEV_AUTH_TOKEN" not in client_js


def test_render_preserves_report_formulas_and_is_deterministic() -> None:
    first = render_client_report_xlsx([], published_at=None)
    second = render_client_report_xlsx([], published_at=None)
    assert first == second
    assert _sheet_names(first) == list(CLIENT_WORKBOOK_SHEET_NAMES)
    with ZipFile(io.BytesIO(first)) as archive:
        workbook_xml = archive.read("xl/workbook.xml").decode("utf-8")
        rels_xml = archive.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        parts = _sheet_part_map(workbook_xml, rels_xml)
        overall = archive.read(parts[REPORT_SHEET_NAMES[0]]).decode("utf-8")
        cache = archive.read("xl/pivotCache/pivotCacheDefinition1.xml").decode("utf-8")
        table = archive.read("xl/pivotTables/pivotTable1.xml").decode("utf-8")
    assert "UNIQUE(" not in overall
    assert 'cacheSource type="worksheet"' in cache
    assert 'sheet="PublishedFacts"' in cache
    assert 'name="PivotOverall"' in table
    assert 'outline="1"' in table
    assert XLSX_PATH.is_file()
    assert CLIENT_REPORT_DOWNLOAD_NAME == "Client_Report.xlsx"
    assert CLIENT_REPORT_MEDIA_TYPE == XLSX_TYPE
