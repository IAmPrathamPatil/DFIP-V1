"""RUN 009: controlled extra-column / schema extensibility.

The 57 Web Engage headers (K:BO) remain the locked source contract.
FACT_HEADERS stays the frozen 45-column Excel/API reporting contract.

Approved extras (currently Campaign Objective) may appear after the 57-window.
They are stored on stg_source_row.raw only. They are not facts, not
FactResponse fields, not PivotCache fields, and not rpt_* columns.
"""

from __future__ import annotations

import io
import re
import time
from datetime import date, datetime
from pathlib import Path
from zipfile import ZipFile

import pytest
from dfip_api.schemas import FactResponse, StagedRowResponse
from dfip_api.service import fact_to_response, staged_row_to_response
from dfip_core.ingest import InMemoryIngestStore, ingest_workbook, inspect_workbook
from dfip_core.ingest.headers import HeaderContractError, classify_trailing_headers
from dfip_core.ingest.store import StagedRowRecord
from dfip_core.transform import ConfigBinder, transform_row
from dfip_db.catalog import (
    APPROVED_EXTRA_SOURCE_COLUMNS,
    APPROVED_EXTRA_SOURCE_HEADERS,
    SOURCE_COLUMNS,
    WEB_ENGAGE_SOURCE_HEADERS,
)
from dfip_web.client_report_download import (
    ARTIFACT_REFRESHABLE,
    ARTIFACT_STATIC,
    FACT_VALUE_FIELDS,
    query_table_field_names,
    render_client_report_xlsx,
)
from dfip_web.client_workbook import FACT_HEADERS
from dfip_web.daily_report import REPORT_SHEET_NAMES
from dfip_web.pivot_report import PIVOT_CACHE_PART, PIVOT_TABLE_NAMES, assert_native_pivot_package
from dfip_web.report_format import RUN006_FONT
from dfip_web.slicer_defaults import assert_slicer_package
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from test_p4_transform import CLIENT_ID, make_raw
from test_p8_client_report import _campaign_ids
from test_p11_pivot_report import _published_fact

DERIVED_HEADERS = tuple(
    col.excel_header for col in SOURCE_COLUMNS if col.source_role == "derived_excel"
)
APPROVED_EXTRA = "Campaign Objective"
FIELDN = re.compile(r"^Field\d+$")
COMPANY_1 = "a0000000-0000-4000-8000-000000000001"
COMPANY_2 = "b0000000-0000-4000-8000-000000000002"
NEW_COMPANY = "c0000000-0000-4000-8000-000000000003"


def _write_source(
    path: Path,
    rows: list[dict[str, object]],
    *,
    layout: str = "legacy",
    extra_headers: tuple[str, ...] = (),
    header_override: dict[int, object] | None = None,
    blank_before_extra: bool = False,
) -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Web-Engage Raw"
    headers = list(WEB_ENGAGE_SOURCE_HEADERS)
    start = 11 if layout == "legacy" else 1
    if layout == "legacy":
        for index, header in enumerate(DERIVED_HEADERS, start=1):
            worksheet.cell(1, index, header)
    for index, header in enumerate(headers, start=start):
        worksheet.cell(1, index, header)
    extra_col = start + 57
    if blank_before_extra:
        extra_col += 1
    for offset, header in enumerate(extra_headers):
        worksheet.cell(1, extra_col + offset, header)
    if header_override:
        for col, value in header_override.items():
            worksheet.cell(1, col, value)
    for offset, payload in enumerate(rows):
        excel_row = 2 + offset
        for index, header in enumerate(headers):
            if header in payload:
                worksheet.cell(excel_row, start + index, payload[header])
        for extra_offset, header in enumerate(extra_headers):
            if header in payload:
                worksheet.cell(excel_row, extra_col + extra_offset, payload[header])
    workbook.save(path)
    workbook.close()
    return path


def _sample_row(**overrides: object) -> dict[str, object]:
    row: dict[str, object] = {
        "Day": datetime(2025, 8, 1),
        "Campaign Name": "Alpha",
        "Campaign ID": "camp-1",
        "Variation ID": "var-1",
        "Channel": "WhatsApp",
        "Sent": 10,
        "Delivered": 8,
    }
    row.update(overrides)
    return row


def _ingest(path: Path, *, client_id: str = COMPANY_1) -> tuple:
    store = InMemoryIngestStore()
    result = ingest_workbook(path, store, client_id=client_id)
    return store, result


def _workbook_bytes(rows, *, artifact: str = ARTIFACT_STATIC, client_id: str, code: str):
    return render_client_report_xlsx(
        rows,
        published_at=None,
        client_id=client_id,
        client_name=code,
        client_code=code,
        publication_id=f"pub-{code}",
        artifact=artifact,
        api_base_url="http://127.0.0.1:8000" if artifact == ARTIFACT_REFRESHABLE else None,
    )


def _cache_field_names(body: bytes) -> list[str]:
    with ZipFile(io.BytesIO(body)) as archive:
        cache = archive.read(PIVOT_CACHE_PART).decode("utf-8")
    return re.findall(r'<cacheField name="([^"]+)"', cache)


def _assert_frozen_excel_contract(body: bytes, *, refreshable: bool = False) -> None:
    assert_native_pivot_package(body)
    assert_slicer_package(body)
    names = _cache_field_names(body)
    assert names[:45] == list(FACT_HEADERS)
    assert len(FACT_HEADERS) == 45
    assert APPROVED_EXTRA not in FACT_HEADERS
    assert APPROVED_EXTRA not in names
    assert not any(FIELDN.match(name) for name in names[:45])
    with ZipFile(io.BytesIO(body)) as archive:
        slicers = [
            n for n in archive.namelist() if n.startswith("xl/slicerCaches/") and n.endswith(".xml")
        ]
        styles = archive.read("xl/styles.xml").decode("utf-8")
        pivots = [n for n in archive.namelist() if n.startswith("xl/pivotTables/pivotTable")]
    assert len(slicers) == 35
    assert len(pivots) == 9
    assert RUN006_FONT in styles
    if refreshable:
        assert query_table_field_names(body) == FACT_HEADERS
    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=False)
    facts = workbook["PublishedFacts"]
    headers = [facts.cell(1, column).value for column in range(1, len(FACT_HEADERS) + 1)]
    assert headers == list(FACT_HEADERS)
    assert facts.cell(1, 46).value is None
    workbook.close()


def test_approved_extra_is_not_a_canonical_source_or_excel_field() -> None:
    assert len(SOURCE_COLUMNS) == 67
    assert len(WEB_ENGAGE_SOURCE_HEADERS) == 57
    assert APPROVED_EXTRA_SOURCE_HEADERS == (APPROVED_EXTRA,)
    extra = APPROVED_EXTRA_SOURCE_COLUMNS[0]
    assert extra.required is False
    assert extra.excel_exposure == "none"
    assert extra.powerbi_exposure == "none"
    assert extra.db_column == "campaign_objective"
    assert extra.value_kind == "text"
    assert APPROVED_EXTRA not in {col.excel_header for col in SOURCE_COLUMNS}
    assert APPROVED_EXTRA not in FACT_HEADERS
    assert extra.db_column not in FACT_VALUE_FIELDS
    assert extra.db_column not in FactResponse.model_fields
    assert tuple(FactResponse.model_fields) == FACT_VALUE_FIELDS


def test_fixture_a_existing_schema_still_ingests(tmp_path: Path) -> None:
    path = _write_source(tmp_path / "a.xlsx", [_sample_row()])
    store, result = _ingest(path)
    assert result.batch.status == "staged"
    assert result.batch.error_summary is None
    layout = inspect_workbook(path)
    assert layout.match.headers == WEB_ENGAGE_SOURCE_HEADERS
    assert layout.match.extra_headers == ()
    raw = store.staged_for_batch(result.batch.id)[0].raw
    assert list(raw) == list(WEB_ENGAGE_SOURCE_HEADERS)
    assert APPROVED_EXTRA not in raw


def test_fixture_b_approved_extra_is_staged_not_transformed(tmp_path: Path) -> None:
    path = _write_source(
        tmp_path / "b.xlsx",
        [_sample_row(**{APPROVED_EXTRA: "Awareness"})],
        extra_headers=(APPROVED_EXTRA,),
    )
    store, result = _ingest(path)
    assert result.batch.status == "staged"
    layout = inspect_workbook(path)
    assert layout.match.extra_headers == (APPROVED_EXTRA,)
    staged = store.staged_for_batch(result.batch.id)[0]
    assert list(staged.raw)[:57] == list(WEB_ENGAGE_SOURCE_HEADERS)
    assert staged.raw[APPROVED_EXTRA] == "Awareness"
    assert len(staged.raw) == 58
    outcome = transform_row(
        staged,
        ConfigBinder(),
        client_id=COMPANY_1,
        batch_id=result.batch.id,
        processing_run_id="run-1",
    )
    assert outcome.rejection is None
    assert outcome.fact is not None
    assert not hasattr(outcome.fact, "campaign_objective")
    response = fact_to_response(outcome.fact)
    dumped = response.model_dump()
    assert APPROVED_EXTRA not in dumped
    assert "campaign_objective" not in dumped
    staged_api = staged_row_to_response(staged)
    assert staged_api.raw[APPROVED_EXTRA] == "Awareness"
    with pytest.raises(ValidationError):
        FactResponse.model_validate({**dumped, "campaign_objective": "Awareness"})


def test_fixture_b_missing_optional_extra_is_fixture_a(tmp_path: Path) -> None:
    path = _write_source(tmp_path / "e.xlsx", [_sample_row()])
    layout = inspect_workbook(path)
    assert layout.match.extra_headers == ()
    _, result = _ingest(path)
    assert result.batch.status == "staged"


def test_fixture_c_unknown_column_fails_closed(tmp_path: Path) -> None:
    path = _write_source(
        tmp_path / "c.xlsx",
        [_sample_row(**{"Campaign Goal": "X"})],
        extra_headers=("Campaign Goal",),
    )
    store, result = _ingest(path)
    assert result.batch.status == "failed"
    assert result.batch.error_summary is not None
    assert result.batch.error_summary.startswith(
        "UNAPPROVED_SOURCE_COLUMN: Unapproved source column: Campaign Goal"
    )
    rejected = store.rejected_for_batch(result.batch.id)
    assert rejected[0].reason_code == "UNAPPROVED_SOURCE_COLUMN"
    assert rejected[0].reason_detail == "Unapproved source column: Campaign Goal"
    with pytest.raises(HeaderContractError) as caught:
        inspect_workbook(path)
    assert caught.value.reason_code == "UNAPPROVED_SOURCE_COLUMN"
    assert str(caught.value) == "UNAPPROVED_SOURCE_COLUMN: Unapproved source column: Campaign Goal"


def test_fixture_d_missing_required_canonical_column(tmp_path: Path) -> None:
    path = _write_source(
        tmp_path / "d.xlsx",
        [_sample_row()],
        header_override={11: "Business Day"},
    )
    store, result = _ingest(path)
    assert result.batch.status == "failed"
    assert result.batch.error_summary is not None
    assert "HEADER_CONTRACT" in result.batch.error_summary
    assert store.rejected_for_batch(result.batch.id)[0].reason_code == "HEADER_CONTRACT"


def test_fixture_f_renamed_canonical_column(tmp_path: Path) -> None:
    path = _write_source(
        tmp_path / "f.xlsx",
        [_sample_row()],
        header_override={12: "Campaign"},
    )
    result = ingest_workbook(path, InMemoryIngestStore())
    assert result.batch.status == "failed"
    assert result.batch.error_summary is not None
    assert "HEADER_CONTRACT" in result.batch.error_summary


def test_fixture_g_duplicate_headers(tmp_path: Path) -> None:
    inside = _write_source(
        tmp_path / "g-inside.xlsx",
        [_sample_row()],
        header_override={12: "Day"},
    )
    clash = _write_source(
        tmp_path / "g-clash.xlsx",
        [_sample_row()],
        extra_headers=("Campaign Name",),
    )
    extra_dup = _write_source(
        tmp_path / "g-extra.xlsx",
        [_sample_row()],
        extra_headers=(APPROVED_EXTRA, APPROVED_EXTRA),
    )
    for path in (inside, clash, extra_dup):
        result = ingest_workbook(path, InMemoryIngestStore())
        assert result.batch.status == "failed"
        assert result.batch.error_summary is not None
        assert result.batch.error_summary.startswith("DUPLICATE_HEADERS:")


def test_fixture_i_empty_column_name(tmp_path: Path) -> None:
    path = _write_source(
        tmp_path / "i.xlsx",
        [_sample_row(**{APPROVED_EXTRA: "Awareness"})],
        extra_headers=(APPROVED_EXTRA,),
        blank_before_extra=True,
    )
    result = ingest_workbook(path, InMemoryIngestStore())
    assert result.batch.status == "failed"
    assert result.batch.error_summary is not None
    assert result.batch.error_summary.startswith("EMPTY_COLUMN_NAME:")


def test_reordered_canonical_headers_are_not_identity(tmp_path: Path) -> None:
    path = _write_source(tmp_path / "reorder.xlsx", [_sample_row()])
    workbook = load_workbook(path)
    sheet = workbook.active
    sheet.cell(1, 12, "Campaign ID")
    sheet.cell(1, 13, "Campaign Name")
    workbook.save(path)
    workbook.close()
    result = ingest_workbook(path, InMemoryIngestStore())
    assert result.batch.status == "failed"
    assert result.batch.error_summary is not None
    assert "HEADER_CONTRACT" in result.batch.error_summary


def test_fixture_h_type_incompatible_existing_field_is_invalid_numeric() -> None:
    raw = make_raw(**{"Sent": "not-a-number"})
    staged = StagedRowRecord(
        id="stg-h",
        batch_id="batch-1",
        source_row_number=2,
        raw=raw,
        campaign_id=raw["Campaign ID"],
        variation_id=raw["Variation ID"],
        day=date(2025, 8, 1),
    )
    outcome = transform_row(
        staged,
        ConfigBinder(),
        client_id=CLIENT_ID,
        batch_id="batch-1",
        processing_run_id="run-1",
    )
    assert outcome.fact is None
    assert outcome.rejection is not None
    assert outcome.rejection.reason_code == "INVALID_NUMERIC"


def test_approved_extra_does_not_change_numeric_coercion() -> None:
    raw = make_raw()
    raw[APPROVED_EXTRA] = "Awareness"
    raw["Delivered"] = "1.5"
    staged = StagedRowRecord(
        id="stg-type",
        batch_id="batch-1",
        source_row_number=2,
        raw=raw,
        campaign_id=raw["Campaign ID"],
        variation_id=raw["Variation ID"],
        day=date(2025, 8, 1),
    )
    outcome = transform_row(
        staged,
        ConfigBinder(),
        client_id=CLIENT_ID,
        batch_id="batch-1",
        processing_run_id="run-1",
    )
    assert outcome.rejection is not None
    assert outcome.rejection.reason_code == "INVALID_NUMERIC"


def test_source_position_is_not_identity_for_approved_extra() -> None:
    observed = list(WEB_ENGAGE_SOURCE_HEADERS) + [None, None, APPROVED_EXTRA]
    with pytest.raises(HeaderContractError) as caught:
        classify_trailing_headers(observed, 0)
    assert caught.value.reason_code == "EMPTY_COLUMN_NAME"
    extras = classify_trailing_headers(list(WEB_ENGAGE_SOURCE_HEADERS) + [APPROVED_EXTRA], 0)
    assert extras == (APPROVED_EXTRA,)


def test_native_layout_approved_extra(tmp_path: Path) -> None:
    path = _write_source(
        tmp_path / "native.xlsx",
        [_sample_row(**{APPROVED_EXTRA: "Retention"})],
        layout="native",
        extra_headers=(APPROVED_EXTRA,),
    )
    store, result = _ingest(path)
    assert result.batch.status == "staged"
    assert inspect_workbook(path).match.source_kind == "native_export"
    assert store.staged_for_batch(result.batch.id)[0].raw[APPROVED_EXTRA] == "Retention"


def test_company_isolation_keeps_extras_in_tenant_raw(tmp_path: Path) -> None:
    path_a = _write_source(
        tmp_path / "iso-a.xlsx",
        [_sample_row(**{"Campaign ID": "camp-a", APPROVED_EXTRA: "A-only"})],
        extra_headers=(APPROVED_EXTRA,),
    )
    path_b = _write_source(
        tmp_path / "iso-b.xlsx",
        [_sample_row(**{"Campaign ID": "camp-b", APPROVED_EXTRA: "B-only"})],
        extra_headers=(APPROVED_EXTRA,),
    )
    path_n = _write_source(
        tmp_path / "iso-n.xlsx",
        [_sample_row(**{"Campaign ID": "camp-n", APPROVED_EXTRA: "N-only"})],
        extra_headers=(APPROVED_EXTRA,),
    )
    store_a, result_a = _ingest(path_a, client_id=COMPANY_1)
    store_b, result_b = _ingest(path_b, client_id=COMPANY_2)
    store_n, result_n = _ingest(path_n, client_id=NEW_COMPANY)
    raw_a = store_a.staged_for_batch(result_a.batch.id)[0].raw
    raw_b = store_b.staged_for_batch(result_b.batch.id)[0].raw
    raw_n = store_n.staged_for_batch(result_n.batch.id)[0].raw
    assert result_a.batch.client_id == COMPANY_1
    assert result_b.batch.client_id == COMPANY_2
    assert result_n.batch.client_id == NEW_COMPANY
    assert raw_a[APPROVED_EXTRA] == "A-only"
    assert raw_b[APPROVED_EXTRA] == "B-only"
    assert raw_n[APPROVED_EXTRA] == "N-only"
    assert "client_id" not in raw_a
    assert raw_a["Campaign ID"] == "camp-a"
    assert raw_b["Campaign ID"] == "camp-b"
    assert raw_n["Campaign ID"] == "camp-n"


def test_legacy_and_new_workbooks_keep_frozen_excel_contract() -> None:
    rows = [_published_fact(campaign_id="camp-a")]
    static = _workbook_bytes(rows, client_id=COMPANY_1, code="C1")
    refreshable = _workbook_bytes(
        rows, artifact=ARTIFACT_REFRESHABLE, client_id=COMPANY_1, code="C1"
    )
    _assert_frozen_excel_contract(static)
    _assert_frozen_excel_contract(refreshable, refreshable=True)
    assert _campaign_ids(static) == ["camp-a"]
    assert APPROVED_EXTRA not in _campaign_ids(static)


def test_company_workbooks_do_not_gain_extra_excel_fields() -> None:
    company_a = _workbook_bytes(
        [_published_fact(campaign_id="camp-a")], client_id=COMPANY_1, code="C1"
    )
    company_b = _workbook_bytes(
        [_published_fact(campaign_id="camp-b")], client_id=COMPANY_2, code="C2"
    )
    new_company = _workbook_bytes(
        [_published_fact(campaign_id="camp-n")], client_id=NEW_COMPANY, code="NEWCO"
    )
    for body, campaign in (
        (company_a, "camp-a"),
        (company_b, "camp-b"),
        (new_company, "camp-n"),
    ):
        _assert_frozen_excel_contract(body)
        assert _campaign_ids(body) == [campaign]


def test_staged_row_api_allows_approved_extra_inside_raw() -> None:
    raw = make_raw()
    raw[APPROVED_EXTRA] = "Awareness"
    staged = StagedRowRecord(
        id="stg-api",
        batch_id="batch-1",
        source_row_number=2,
        raw=raw,
        campaign_id="camp-1",
        variation_id="var-1",
        day=date(2025, 8, 1),
    )
    body = staged_row_to_response(staged).model_dump()
    parsed = StagedRowResponse.model_validate(body)
    assert parsed.raw[APPROVED_EXTRA] == "Awareness"


def test_ingest_performance_is_comparable_with_one_approved_extra(tmp_path: Path) -> None:
    canonical = _write_source(tmp_path / "perf-a.xlsx", [_sample_row()])
    extended = _write_source(
        tmp_path / "perf-b.xlsx",
        [_sample_row(**{APPROVED_EXTRA: "Awareness"})],
        extra_headers=(APPROVED_EXTRA,),
    )
    started = time.perf_counter()
    ingest_workbook(canonical, InMemoryIngestStore())
    baseline = time.perf_counter() - started
    started = time.perf_counter()
    ingest_workbook(extended, InMemoryIngestStore())
    extended_elapsed = time.perf_counter() - started
    assert baseline < 5
    assert extended_elapsed < 5


def _dispatch_excel():
    win32com_client = pytest.importorskip("win32com.client")
    try:
        excel = win32com_client.DispatchEx("Excel.Application")
    except Exception as exc:
        pytest.skip(f"Excel COM not available: {exc}")
    excel.Visible = False
    excel.DisplayAlerts = False
    excel.AskToUpdateLinks = False
    excel.EnableEvents = False
    return excel


def test_desktop_excel_legacy_contract_after_approved_extra(tmp_path: Path) -> None:
    """Legacy 45-field workbook: 9 pivots, 35 slicers, no FieldN. Not live Refresh All."""
    xl_sheet_hidden = 0
    rows = [_published_fact(campaign_id="camp-a")]
    static_path = tmp_path / "run009_legacy_static.xlsx"
    refresh_path = tmp_path / "run009_legacy_refreshable.xlsx"
    static_path.write_bytes(_workbook_bytes(rows, client_id=COMPANY_1, code="C1"))
    refresh_path.write_bytes(
        _workbook_bytes(rows, artifact=ARTIFACT_REFRESHABLE, client_id=COMPANY_1, code="C1")
    )
    excel = _dispatch_excel()
    try:
        for path in (static_path, refresh_path):
            workbook = excel.Workbooks.Open(str(path.resolve()), UpdateLinks=0, ReadOnly=False)
            excel.ActiveWindow.SelectedSheets.Item(1).Select()
            assert workbook.Worksheets.Count == 11
            assert workbook.Worksheets("PublishedFacts").Visible == xl_sheet_hidden
            assert workbook.SlicerCaches.Count == 35
            assert sum(int(ws.PivotTables().Count) for ws in workbook.Worksheets) == 9
            overall = workbook.Worksheets(REPORT_SHEET_NAMES[0])
            overall.PivotTables(1).PivotCache().Refresh()
            field_names = [str(field.Name) for field in overall.PivotTables(1).PivotFields()]
            assert APPROVED_EXTRA not in field_names
            assert not any(FIELDN.match(name) for name in field_names)
            assert str(overall.PivotTables(1).Name) in PIVOT_TABLE_NAMES.values()
            workbook.Close(False)
    finally:
        try:
            excel.Quit()
        except Exception:
            pass
        time.sleep(0.4)
