"""Nine-sheet Daily Report contract, workbook integrity, and KPI reconciliation."""

from __future__ import annotations

import re
from decimal import Decimal
from pathlib import Path
from zipfile import ZipFile

from dfip_analytics.kpis import spec_by_slug
from dfip_web.client_workbook import (
    FACT_HEADERS,
    FAKE_MASHUP_ZIP_PARTS,
    M_PATH,
    XLSX_PATH,
    mashup_text,
)
from dfip_web.daily_report import (
    AMC_GROUP,
    CLIENT_WORKBOOK_SHEET_NAMES,
    D2C_GROUP,
    EMPTY_STATE,
    FL4_HEADER,
    FL5_HEADER,
    MEASURE_HEADERS,
    MEASURE_OPS,
    NATIVE_SHEET_NAMES,
    REPORT_CONTRACTS,
    REPORT_KPI_SLUGS,
    REPORT_SHEET_NAMES,
    SERVICE_FILTER_LOGIC_1,
    SERVICE_GROUP,
    SOURCE_WORKBOOK,
    VERTICAL,
    additive_formula,
    aggregate_published_rows,
    attach_daily_report_sheets,
    group_display_name,
    report_formula,
    worksheet_xml,
)
from openpyxl import load_workbook

from test_p7_publication import (
    FORBIDDEN_TEMPLATE_TOKENS,
    PUBLISHED_FACTS_PATH,
    WORKING_SET_FACTS_PATH,
    _omits_working_set_facts_url,
    assert_no_fake_mashup_parts,
)

ROOT = Path(__file__).resolve().parents[1]


def test_source_contract_uses_recovered_sheet_names() -> None:
    assert SOURCE_WORKBOOK == "Web Engage - Daily Report - FY-2026.xlsx"
    assert REPORT_SHEET_NAMES[0] == "Overall Daywise Report "
    assert REPORT_SHEET_NAMES[7] == "D2C Vertical "
    assert REPORT_SHEET_NAMES[8] == "Service Campaigns"
    assert len(REPORT_CONTRACTS) == 9
    assert len(set(item.name for item in REPORT_CONTRACTS)) == 9


def test_channel_wise_uses_channel_as_filter_not_row_field() -> None:
    channel = next(item for item in REPORT_CONTRACTS if item.name == "Channel Wise")
    assert "Channel" not in channel.grain
    assert "Channel" in channel.slicers
    assert channel.grain == ("Month", "Filter Logic 1_2", "Filter Logic 1")


def test_amc_d2c_service_defaults_match_source_semantics() -> None:
    by_name = {item.name: item for item in REPORT_CONTRACTS}
    assert by_name["AMC DayWise"].default_group == AMC_GROUP
    assert by_name["AMC Split"].default_group == AMC_GROUP
    assert by_name["AMC Vertical Monthly Split"].default_group == AMC_GROUP
    assert by_name["D2C Vertical "].default_group == D2C_GROUP
    assert by_name["Service Campaigns"].default_filter_logic_1 == SERVICE_FILTER_LOGIC_1
    assert by_name["Service Campaigns"].default_group == SERVICE_GROUP
    assert by_name["Overall Daywise Report "].default_group == ""


def test_report_formulas_sum_then_divide_and_read_publishedfacts() -> None:
    overall = next(item for item in REPORT_CONTRACTS if item.name == "Overall Daywise Report ")
    formula = report_formula(overall)
    assert "PublishedFacts" in formula
    assert "UNIQUE(" in formula
    assert "FILTER(" in formula
    assert "SUMIFS(" in additive_formula(overall, "Sent")
    assert "GROUPBY" not in formula
    assert "LAMBDA" not in formula
    assert "HSTACK" not in formula
    assert "CHOOSECOLS" not in formula
    assert "DROP(" not in formula
    assert "GET /api/v1/facts" not in formula
    assert WORKING_SET_FACTS_PATH not in formula
    assert "VLOOKUP" not in formula
    assert '"Delivered"' in additive_formula(overall, "Delivered")
    assert "filter_logic_1_group" in formula
    assert "logic1Filter" in formula
    assert re.search(r"\bf1\s*,", formula) is None
    assert re.search(r"\bflt1\s*,", formula) is None
    assert len(formula) < 8192
    assert tuple(op[0] for op in MEASURE_OPS) == MEASURE_HEADERS
    cell_like = re.compile(r"^[A-Za-z]{1,3}\d+$")
    for contract in REPORT_CONTRACTS:
        text = report_formula(contract)
        for token in re.findall(r"(?:LET\(|,)([A-Za-z_][A-Za-z0-9_]*)\s*,", text):
            assert not cell_like.match(token), token


def test_client_kpi_formulas_are_the_report_layer() -> None:
    assert spec_by_slug("client", "ctr_del_to_clicks").formula == "'Unique Clicks'/Delivered"
    assert spec_by_slug("client", "delivery_rate").formula == "Delivered/Sent"
    assert spec_by_slug("client", "overall_roas").formula == "'Revenue (INR)'/'Total Cost'"
    assert spec_by_slug("client", "failed_rate_sm").formula == "Failed/Sent"
    for slug in REPORT_KPI_SLUGS:
        spec = spec_by_slug("client", slug)
        assert spec.namespace == "client"


def test_aggregate_reconciles_clicks_delivered_ctr() -> None:
    overall = next(item for item in REPORT_CONTRACTS if item.name == "Overall Daywise Report ")
    rows = [
        {
            "Channel": "WhatsApp",
            "Month": "Apr-25",
            "Filter Logic 1": "Service | Campaigns",
            "filter_logic_1_group": "Group2",
            "Sent": 200,
            "Failed": 20,
            "Delivered": 80,
            "Unique Impressions": 50,
            "Unique Clicks": 30,
            "Unique Conversions": 4,
            "Unique Click-Through Conversions": 3,
            "Total Cost": "10",
            "Revenue (INR)": "40",
            "Click-Through Revenue (INR)": "20",
        },
        {
            "Channel": "WhatsApp",
            "Month": "Apr-25",
            "Filter Logic 1": "Service | Campaigns",
            "filter_logic_1_group": "Group2",
            "Sent": 50,
            "Failed": 5,
            "Delivered": 20,
            "Unique Impressions": 10,
            "Unique Clicks": 10,
            "Unique Conversions": 1,
            "Unique Click-Through Conversions": 1,
            "Total Cost": "5",
            "Revenue (INR)": "10",
            "Click-Through Revenue (INR)": "5",
        },
    ]
    result = aggregate_published_rows(rows, overall)
    assert result["ctr_del_to_clicks"] == Decimal("0.400000")
    assert result["delivery_rate"] == Decimal("0.400000")
    assert result["actual_sent"] == Decimal("225")


def test_amc_default_filter_excludes_other_groups() -> None:
    amc = next(item for item in REPORT_CONTRACTS if item.name == "AMC DayWise")
    rows = [
        {
            "filter_logic_1_group": "Group7",
            "Delivered": 10,
            "Unique Clicks": 4,
            "Sent": 10,
            "Failed": 0,
            "Unique Impressions": 8,
            "Unique Conversions": 1,
            "Unique Click-Through Conversions": 1,
            "Total Cost": "2",
            "Revenue (INR)": "4",
            "Click-Through Revenue (INR)": "3",
        },
        {
            "filter_logic_1_group": "Group5",
            "Delivered": 100,
            "Unique Clicks": 40,
            "Sent": 100,
            "Failed": 0,
            "Unique Impressions": 80,
            "Unique Conversions": 10,
            "Unique Click-Through Conversions": 10,
            "Total Cost": "20",
            "Revenue (INR)": "40",
            "Click-Through Revenue (INR)": "30",
        },
    ]
    result = aggregate_published_rows(rows, amc)
    assert result["ctr_del_to_clicks"] == Decimal("0.400000")
    assert result["delivery_rate"] == Decimal("1.000000")


def test_empty_published_rows_do_not_fabricate_zeros() -> None:
    overall = next(item for item in REPORT_CONTRACTS if item.name == "Overall Daywise Report ")
    result = aggregate_published_rows([], overall)
    assert result["ctr_del_to_clicks"] is None
    assert result["overall_roas"] is None
    assert EMPTY_STATE == "No published data available."


def test_tracked_workbook_has_nine_report_sheets_and_native_v2x() -> None:
    names, connections, item_props = assert_no_fake_mashup_parts(XLSX_PATH)
    assert "xl/connections.xml" in names
    assert any(item.startswith("xl/queryTables/") for item in names)
    assert "Microsoft.Mashup.OleDb.1" in connections
    assert "Location=PublishedFacts" in connections
    assert "http://schemas.microsoft.com/DataMashup" in item_props
    for part in FAKE_MASHUP_ZIP_PARTS:
        assert part not in names
    workbook = load_workbook(XLSX_PATH, read_only=True, data_only=False)
    assert workbook.sheetnames == list(CLIENT_WORKBOOK_SHEET_NAMES)
    for name in REPORT_SHEET_NAMES:
        sheet = workbook[name]
        assert sheet["A1"].value == "Web Engage Daily Report"
        assert sheet["B10"].value in {None, ""}
        assert sheet["AZ1"].value == "Group1"
        assert sheet["BA1"].value == "API DATA - No Cost"
        assert sheet["AZ7"].value == "Group7"
        assert sheet["BA7"].value == "D2C AMC Vertical"
    from dfip_web.pivot_report import PIVOT_TABLE_NAMES, assert_native_pivot_package

    assert_native_pivot_package(XLSX_PATH.read_bytes())
    with ZipFile(XLSX_PATH) as archive:
        names = archive.namelist()
        report_xml = b"".join(
            archive.read(name)
            for name in names
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
        ).decode("utf-8", errors="ignore")
        pivots = "".join(
            archive.read(f"xl/pivotTables/pivotTable{i}.xml").decode("utf-8") for i in range(1, 10)
        )
        cache = archive.read("xl/pivotCache/pivotCacheDefinition1.xml").decode("utf-8")
        slicers = "".join(
            archive.read(f"xl/slicers/slicer{i}.xml").decode("utf-8") for i in range(1, 10)
        )
    assert "UNIQUE(" not in report_xml
    assert "GROUPBY" not in report_xml
    assert "LAMBDA" not in report_xml
    assert 'outline="1"' in pivots
    for pivot_name in PIVOT_TABLE_NAMES.values():
        assert f'name="{pivot_name}"' in pivots
    assert 'name="Filter Logic 1"' in cache
    assert "Service | FMS" in cache
    assert "Group7" in cache
    assert "minValue" not in cache
    assert "filter_logic_1_group" in cache
    assert FL4_HEADER in slicers
    assert FL5_HEADER in slicers
    assert "Filter Logic 1_2" in slicers
    facts = workbook["Facts"]
    headers = [facts.cell(6, column).value for column in range(1, len(FACT_HEADERS) + 1)]
    assert headers == list(FACT_HEADERS)
    assert facts["A3"].value == "BearerToken"
    assert facts["B3"].value in {None, ""}
    workbook.close()
    raw = XLSX_PATH.read_bytes()
    assert _omits_working_set_facts_url(raw.decode("latin-1"))
    for token in FORBIDDEN_TEMPLATE_TOKENS:
        assert token.encode("utf-8") not in raw
    assert b"xl/queryMashup/" not in raw
    mashup = mashup_text()
    assert mashup == M_PATH.read_text(encoding="utf-8")
    assert PUBLISHED_FACTS_PATH in mashup
    assert WORKING_SET_FACTS_PATH not in mashup.replace(PUBLISHED_FACTS_PATH, "")


def test_attach_rewrites_existing_report_formulas(tmp_path: Path) -> None:
    dest = tmp_path / "rewritten.xlsx"
    attach_daily_report_sheets(XLSX_PATH, dest)
    from dfip_web.pivot_report import assert_native_pivot_package

    assert_native_pivot_package(dest.read_bytes())
    with ZipFile(dest) as archive:
        names = archive.namelist()
        joined = "\n".join(
            archive.read(name).decode("utf-8", errors="ignore")
            for name in names
            if name.endswith(".xml")
        )
        sheet_xml = "".join(
            archive.read(name).decode("utf-8")
            for name in names
            if name.startswith("xl/worksheets/") and name.endswith(".xml")
        )
    assert "D2C AMC Vertical" in joined
    assert FL4_HEADER in joined
    assert "xl/pivotTables/pivotTable1.xml" in names
    assert "xl/slicerCaches/slicerCache35.xml" in names
    assert "UNIQUE(" not in sheet_xml
    assert "GROUPBY" not in joined
    assert "LAMBDA" not in joined
    assert "flt1" not in joined
    assert "xl/metadata.xml" in names


def test_zip_attach_is_idempotent(tmp_path: Path) -> None:
    first = tmp_path / "daily.xlsx"
    attach_daily_report_sheets(XLSX_PATH, first)
    with ZipFile(first) as archive:
        names = archive.namelist()
        workbook_xml = archive.read("xl/workbook.xml").decode("utf-8")
        connections = archive.read("xl/connections.xml").decode("utf-8")
        item_props = archive.read("customXml/itemProps1.xml").decode("utf-8")
    assert "xl/connections.xml" in names
    assert "Microsoft.Mashup.OleDb.1" in connections
    assert "http://schemas.microsoft.com/DataMashup" in item_props
    for name in REPORT_SHEET_NAMES:
        assert name in workbook_xml or _escape_present(name, workbook_xml)
    second = tmp_path / "daily2.xlsx"
    attach_daily_report_sheets(first, second)
    with ZipFile(second) as archive:
        again = archive.read("xl/workbook.xml").decode("utf-8")
    assert again.count("Overall Daywise Report") == workbook_xml.count("Overall Daywise Report")


def _escape_present(name: str, xml: str) -> bool:
    return name.replace("&", "&amp;") in xml


def _kpi_row(**overrides: object) -> dict[str, object]:
    row: dict[str, object] = {
        "Sent": 10,
        "Failed": 0,
        "Delivered": 10,
        "Unique Impressions": 8,
        "Unique Clicks": 4,
        "Unique Conversions": 1,
        "Unique Click-Through Conversions": 1,
        "Total Cost": "2",
        "Revenue (INR)": "4",
        "Click-Through Revenue (INR)": "3",
    }
    row.update(overrides)
    return row


def test_fl1_2_display_captions_do_not_change_filter_identity() -> None:
    by_name = {item.name: item for item in REPORT_CONTRACTS}
    assert group_display_name(AMC_GROUP) == "D2C AMC Vertical"
    assert group_display_name(D2C_GROUP) == "D2C Product Vertical"
    assert group_display_name(SERVICE_GROUP) == "Overall Service Campaigns"
    assert by_name["AMC DayWise"].default_group == "Group7"
    assert by_name["D2C Vertical "].default_group == "Group5"
    assert by_name["Service Campaigns"].default_group == "Group2"
    amc_xml = worksheet_xml(by_name["AMC DayWise"])
    d2c_xml = worksheet_xml(by_name["D2C Vertical "])
    service_xml = worksheet_xml(by_name["Service Campaigns"])
    assert "D2C AMC Vertical" in amc_xml
    assert "D2C Product Vertical" in d2c_xml
    assert "Overall Service Campaigns" in service_xml
    assert "Service | FMS &amp; LMS | Campaigns" in service_xml
    assert "groupResolved" in report_formula(by_name["AMC DayWise"])
    assert "filter_logic_1_group" in report_formula(by_name["AMC DayWise"])
    rows = [
        _kpi_row(filter_logic_1_group="Group7"),
        _kpi_row(filter_logic_1_group="Group5", Delivered=100, **{"Unique Clicks": 40}),
    ]
    amc = aggregate_published_rows(rows, by_name["AMC DayWise"])
    d2c = aggregate_published_rows(rows, by_name["D2C Vertical "])
    assert amc["actual_sent"] == Decimal("10")
    assert d2c["actual_sent"] == Decimal("10")
    assert amc["delivery_rate"] == Decimal("1.000000")
    assert d2c["delivery_rate"] == Decimal("10.000000")
    assert amc["ctr_del_to_clicks"] == Decimal("0.400000")
    assert d2c["ctr_del_to_clicks"] == Decimal("0.400000")


def test_caption_cells_live_on_matching_rows() -> None:
    from xml.etree import ElementTree

    from dfip_web.daily_report import CAPTION_KEY_COL, NS_MAIN

    xml = worksheet_xml(next(item for item in REPORT_CONTRACTS if item.name == VERTICAL))
    root = ElementTree.fromstring(xml)
    ns = {"m": NS_MAIN}
    seen: list[str] = []
    for row in root.findall("m:sheetData/m:row", ns):
        number = row.get("r")
        for cell in row.findall("m:c", ns):
            ref = cell.get("r") or ""
            match = re.search(r"(\d+)$", ref)
            assert match is not None
            assert match.group(1) == number, ref
            if ref.startswith("AZ") or ref.startswith("BA"):
                seen.append(ref)
    assert "AZ1" in seen
    assert "BA7" in seen
    assert f'min="{CAPTION_KEY_COL}"' in xml
    assert 'hidden="1"' in xml
    assert "D2C AMC Vertical" in xml
    assert "Group7" in xml


def test_vertical_fl4_fl5_filters_are_independent() -> None:
    vertical = next(item for item in REPORT_CONTRACTS if item.name == VERTICAL)
    overall = next(item for item in REPORT_CONTRACTS if item.name == "Overall Daywise Report ")
    service = next(item for item in REPORT_CONTRACTS if item.name == "Service Campaigns")
    assert vertical.device_product_filters is True
    assert overall.device_product_filters is False
    assert service.device_product_filters is False
    formula = report_formula(vertical)
    additive = additive_formula(vertical, "Delivered")
    assert "$F$4" in formula and "$G$4" in formula
    assert FL4_HEADER in formula and FL5_HEADER in formula
    assert additive.count("SUMIFS(") == 4
    assert "TRIM($F$4)" in additive
    assert "TRIM($G$4)" in additive
    assert 'IF($F$4="",' not in additive
    assert "$F$4" not in report_formula(overall)
    assert "$G$4" not in report_formula(service)
    assert FL4_HEADER not in worksheet_xml(overall)
    assert FL4_HEADER in worksheet_xml(vertical)
    rows = [
        _kpi_row(
            filter_logic_1_group="Group2",
            **{FL4_HEADER: "DeviceA", FL5_HEADER: "ProdA"},
        ),
        _kpi_row(
            filter_logic_1_group="Group2",
            Delivered=20,
            **{"Unique Clicks": 8, FL4_HEADER: "DeviceB", FL5_HEADER: "ProdA"},
        ),
        _kpi_row(
            filter_logic_1_group="Group2",
            Delivered=30,
            **{"Unique Clicks": 12, FL4_HEADER: "DeviceA", FL5_HEADER: "ProdB"},
        ),
        _kpi_row(
            filter_logic_1_group="Group2",
            Delivered=40,
            **{"Unique Clicks": 16, FL4_HEADER: "", FL5_HEADER: ""},
        ),
    ]
    all_values = aggregate_published_rows(rows, vertical)
    fl4 = aggregate_published_rows(rows, vertical, device="DeviceA")
    fl5 = aggregate_published_rows(rows, vertical, product="ProdA")
    both = aggregate_published_rows(rows, vertical, device="DeviceA", product="ProdA")
    invalid = aggregate_published_rows(rows, vertical, device="NotADevice")
    cleared = aggregate_published_rows(rows, vertical, device="", product="")
    assert all_values["ctr_del_to_clicks"] == Decimal("0.400000")
    assert all_values["actual_sent"] == Decimal("40")
    assert all_values["delivery_rate"] == Decimal("2.500000")
    assert fl4["actual_sent"] == Decimal("20")
    assert fl4["delivery_rate"] == Decimal("2.000000")
    assert fl5["actual_sent"] == Decimal("20")
    assert fl5["delivery_rate"] == Decimal("1.500000")
    assert both["actual_sent"] == Decimal("10")
    assert both["delivery_rate"] == Decimal("1.000000")
    assert both["ctr_del_to_clicks"] == Decimal("0.400000")
    assert invalid["ctr_del_to_clicks"] is None
    assert cleared == all_values
    other = aggregate_published_rows(rows, overall, device="DeviceA", product="ProdB")
    assert other["actual_sent"] == Decimal("40")
    assert other["delivery_rate"] == Decimal("2.500000")


def test_service_default_is_fms_and_group2_without_legacy_fl2_whitelist() -> None:
    service = next(item for item in REPORT_CONTRACTS if item.name == "Service Campaigns")
    source = (
        Path(__file__).resolve().parents[1] / "packages" / "web" / "dfip_web" / "daily_report.py"
    )
    text = source.read_text(encoding="utf-8")
    assert SERVICE_FILTER_LOGIC_1 in text
    assert "FMS BS 1339" not in text
    assert service.default_filter_logic_1 == SERVICE_FILTER_LOGIC_1
    assert service.default_group == SERVICE_GROUP
    assert service.default_filter_logic_1 != "FSC | Service | Campaigns"
    assert "Technician" not in service.default_filter_logic_1
    rows = [
        _kpi_row(
            filter_logic_1_group="Group2",
            **{
                "Filter Logic 1": SERVICE_FILTER_LOGIC_1,
                "Filter Logic 2": "FMS new published value",
            },
        ),
        _kpi_row(
            filter_logic_1_group="Group2",
            Delivered=20,
            **{
                "Unique Clicks": 8,
                "Filter Logic 1": SERVICE_FILTER_LOGIC_1,
                "Filter Logic 2": "FMS BS 1339",
            },
        ),
        _kpi_row(
            filter_logic_1_group="Group2",
            Delivered=100,
            **{"Unique Clicks": 40, "Filter Logic 1": "FSC | Service | Campaigns"},
        ),
        _kpi_row(
            filter_logic_1_group="Group2",
            Delivered=50,
            **{"Unique Clicks": 20, "Filter Logic 1": "Service | Technician | Manual Camp"},
        ),
        _kpi_row(
            filter_logic_1_group="Group7",
            Delivered=80,
            **{"Unique Clicks": 32, "Filter Logic 1": SERVICE_FILTER_LOGIC_1},
        ),
    ]
    result = aggregate_published_rows(rows, service)
    assert result["actual_sent"] == Decimal("20")
    assert result["delivery_rate"] == Decimal("1.500000")
    assert result["ctr_del_to_clicks"] == Decimal("0.400000")


def test_nine_sheet_contracts_and_formulas_remain_valid() -> None:
    grains = {
        "Overall Daywise Report ": ("Month", "Day"),
        "Vertical Level": ("Month", "Filter Logic 1_2"),
        "Channel Wise": ("Month", "Filter Logic 1_2", "Filter Logic 1"),
        "Sub-Split": ("Month", "Filter Logic 1_2", "Filter Logic 1", "Campaign Name"),
        "AMC DayWise": ("Month", "Day"),
        "AMC Split": ("Filter Logic 1", "Filter Logic 2", "Month"),
        "AMC Vertical Monthly Split": ("Filter Logic 1", "Month", "Filter Logic 2"),
        "D2C Vertical ": ("Filter Logic 1", "Month", "Filter Logic 2"),
        "Service Campaigns": ("Filter Logic 1_2", "Month", "Day"),
    }
    assert tuple(item.name for item in REPORT_CONTRACTS) == REPORT_SHEET_NAMES
    mashup = mashup_text()
    assert "D2C AMC Vertical" not in mashup
    for contract in REPORT_CONTRACTS:
        assert contract.grain == grains[contract.name]
        formula = report_formula(contract)
        additive = additive_formula(contract, "Unique Clicks")
        assert "GROUPBY" not in formula
        assert "LAMBDA" not in formula
        assert "HSTACK" not in formula
        assert "DROP(" not in formula
        assert "PublishedFacts" in formula
        assert "filter_logic_1_group" in formula
        assert "groupResolved" in formula
        assert len(formula) < 8192
        assert len(additive) < 8192
        xml = worksheet_xml(contract)
        assert "#NAME?" not in xml
        assert "GROUPBY" not in xml


def test_worksheet_xml_is_well_formed_and_spans_cover_headers() -> None:
    from xml.etree import ElementTree

    from dfip_web.daily_report import CONTRACTS_BY_NAME, SUB_SPLIT

    xml = worksheet_xml(CONTRACTS_BY_NAME[SUB_SPLIT])
    assert 'spans="1:53"' in xml
    assert 'dimension ref="A1:BA10"' in xml
    assert "→" not in xml
    assert "LAMBDA" not in xml
    assert "GROUPBY" not in xml
    assert "CHOOSECOLS" not in xml
    assert "UNIQUE(" in xml
    assert "_xlfn.LET" in xml
    assert "_xlpm.logic1Filter" in xml
    assert "_xlfn._xlws.FILTER" in xml
    assert "_xlfn.ANCHORARRAY($B10)" in xml
    assert "$A:$XFD" in xml
    assert "B10#" not in xml
    assert 't="array" ref="B10:E10"' in xml
    assert 't="array" ref="F10"' in xml
    assert 'cm="1"' in xml
    assert "LOWER(" in xml
    assert "_xlpm.filteredKeys" in xml
    assert "_xlpm.idKeys" in xml
    assert "_xlpm.concatUniq" in xml
    ElementTree.fromstring(xml)


def test_measure_headers_match_source_pivot_data_fields() -> None:
    assert MEASURE_HEADERS[0] == "Total Cost"
    assert MEASURE_HEADERS[3] == "Failed Rate SM"
    assert MEASURE_HEADERS[6] == "Delivery Rate"
    assert MEASURE_HEADERS[-1] == "Overall ROAS"
    assert "Click Through Cost/Conv" not in MEASURE_HEADERS
    assert NATIVE_SHEET_NAMES == ("PublishedFacts", "Facts")


def test_documentation_names_the_daily_report() -> None:
    docs = (ROOT / "documentation" / "DAILY_REPORT.md").read_text(encoding="utf-8")
    assert "Overall Daywise Report" in docs
    assert "PublishedFacts" in docs
    assert "Refresh All" in docs
    assert "UNIQUE" in docs
    assert "PivotTable" in docs
    assert "SUMIFS" in docs
    assert "LOWER" in docs
    assert "D2C AMC Vertical" in docs
    assert "Filter Logic 4" in docs
    assert "not PostgreSQL RLS" in docs or "application-level" in docs.lower()
    assert "0 is a real total of zero" in docs or "mathematical zero" in docs
    assert "blank" in docs and "denominator" in docs


def test_presentation_hides_unused_columns_and_keeps_formulas() -> None:
    from dfip_web.daily_report import (
        MEASURE_FORMATS,
        SUB_SPLIT,
        ZERO_VS_BLANK_HINT,
        last_visible_col,
    )

    assert len(MEASURE_FORMATS) == len(MEASURE_OPS)
    overall = next(item for item in REPORT_CONTRACTS if item.name == "Overall Daywise Report ")
    xml = worksheet_xml(overall)
    last = last_visible_col(overall)
    assert last == 25
    assert f'min="{last + 1}" max="51"' in xml
    assert 'hidden="1"' in xml
    assert 'width="12"' in xml
    assert 'zoomScale="90"' in xml
    assert 'xSplit="3"' in xml
    assert ZERO_VS_BLANK_HINT.split(".")[0] in xml
    assert "UNIQUE(" in xml
    assert "SUMIFS(" in xml
    assert "LAMBDA" not in xml
    assert "GROUPBY" not in xml
    assert "&lt;&gt;" in xml
    assert 'r="B4"' in xml
    sub = next(item for item in REPORT_CONTRACTS if item.name == SUB_SPLIT)
    sub_xml = worksheet_xml(sub)
    assert f'min="{last_visible_col(sub) + 1}" max="51"' in sub_xml


def test_attach_adds_report_number_formats(tmp_path: Path) -> None:
    dest = tmp_path / "styled.xlsx"
    attach_daily_report_sheets(XLSX_PATH, dest)
    with ZipFile(dest) as archive:
        styles = archive.read("xl/styles.xml").decode("utf-8")
        overall = None
        workbook_xml = archive.read("xl/workbook.xml").decode("utf-8")
        rels = archive.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    assert 'numFmtId="164"' in styles
    assert 'formatCode="0.00%"' in styles
    assert 'cellXfs count="21"' in styles
    from dfip_web.daily_report import OVERALL, _sheet_part_map

    parts = _sheet_part_map(workbook_xml, rels)
    with ZipFile(dest) as archive:
        overall = archive.read(parts[OVERALL]).decode("utf-8")
    assert 'hidden="1"' in overall
    assert "UNIQUE(" not in overall
    assert "PublishedFacts" in overall
    assert "<drawing " in overall
