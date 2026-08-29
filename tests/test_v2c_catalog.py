"""V2-C Logic/Labels catalog upload, validation, versioning, and isolation.

In-memory stores only. Does not open PostgreSQL or print credentials.
"""

from __future__ import annotations

from pathlib import Path

from dfip_api.app import create_app
from dfip_config.catalog import UPLOAD_NOTES, InMemoryCatalogStore
from dfip_config.workbook import CatalogWorkbookError, parse_catalog_workbook
from dfip_core.ingest.store import InMemoryIngestStore
from dfip_core.transform.engine import run_transformation
from dfip_core.transform.store import InMemoryFactStore
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from catalog_support import (
    V2C_CAMPAIGN,
    V2C_FL1,
    V2C_FL2,
    V2C_FL3,
    V2C_FL4,
    V2C_FL5,
    V2C_GROUP,
    V2C_MANUAL,
    catalog_xlsx,
    labels_xlsx,
    logic_xlsx,
)
from http_ingest_support import source_row, workbook_bytes
from test_p4_transform import GROUP7_CAMPAIGN, GROUP7_FILTER_LOGIC_1
from test_p5_api import JWT_SECRET, _encode_jwt, make_settings
from test_v2_http_ingest import _error, _upload

CLIENT_A = "a0000000-0000-4000-8000-000000000001"
CLIENT_B = "a0000000-0000-4000-8000-000000000002"
CAMPAIGN_V2 = "a0000000-0000-4000-8000-000000000022"
LABEL_GROUP = "a0000000-0000-4000-8000-000000000041"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _headers(role: str, client_id: str | None) -> dict[str, str]:
    return {"Authorization": f"Bearer {_encode_jwt(role=role, client_id=client_id)}"}


def _app(catalog: InMemoryCatalogStore | None = None) -> tuple[TestClient, InMemoryCatalogStore]:
    store = catalog if catalog is not None else InMemoryCatalogStore()
    client = TestClient(
        create_app(
            settings=make_settings(
                dfip_auth_mode="jwt",
                dfip_auth_secret=JWT_SECRET,
                database_url="",
            ),
            ingest_store=InMemoryIngestStore(),
            fact_store=InMemoryFactStore(),
            catalog_store=store,
        )
    )
    return client, store


def _post_catalog(
    http: TestClient,
    kind: str,
    content: bytes,
    filename: str,
    headers: dict[str, str],
    client_id: str | None = None,
):
    data = {"client_id": client_id} if client_id else None
    return http.post(
        f"/api/v1/catalogs/{kind}",
        headers=headers,
        files={"file": (filename, content, XLSX_TYPE)},
        data=data,
    )


def _unscoped_headers(role: str = "publisher") -> dict[str, str]:
    return {"Authorization": f"Bearer {_encode_jwt(role=role)}"}


def test_parse_valid_logic_and_labels() -> None:
    logic = parse_catalog_workbook("logic", logic_xlsx())
    assert logic.kind == "logic"
    assert logic.rows[0]["campaign_name"] == V2C_CAMPAIGN
    assert logic.rows[0]["filter_logic_1"] == V2C_FL1
    assert logic.rows[0]["filter_logic_2"] == V2C_FL2
    assert logic.rows[0]["amc_status_filter_logic_3"] == V2C_FL3
    assert logic.rows[0]["amc_device_category_filter_logic_4"] == V2C_FL4
    assert logic.rows[0]["amc_product_cat_filter_logic_5"] == V2C_FL5
    assert logic.rows[0]["manual_or_automated"] == V2C_MANUAL
    labels = parse_catalog_workbook("labels", labels_xlsx())
    assert labels.rows[0]["group_name"] == V2C_GROUP
    assert labels.rows[0]["filter_logic_1_value"] == V2C_FL1


def test_parse_rejects_invalid_logic_schema() -> None:
    payload = catalog_xlsx(("Wrong",), [{"Wrong": "x"}])
    try:
        parse_catalog_workbook("logic", payload)
    except CatalogWorkbookError as exc:
        assert any(item.code == "INVALID_SCHEMA" for item in exc.issues)
    else:
        raise AssertionError("expected CatalogWorkbookError")


def test_parse_rejects_invalid_logic_content() -> None:
    payload = logic_xlsx([{"Campaign Name": None, "Filter Logic 1": "x"}])
    try:
        parse_catalog_workbook("logic", payload)
    except CatalogWorkbookError as exc:
        assert any(item.code == "INVALID_CONTENT" for item in exc.issues)
    else:
        raise AssertionError("expected CatalogWorkbookError")


def test_parse_rejects_duplicate_label_values() -> None:
    payload = labels_xlsx(
        [
            {"Group Name": "G1", "Filter Logic 1": V2C_FL1},
            {"Group Name": "G2", "Filter Logic 1": V2C_FL1},
        ]
    )
    try:
        parse_catalog_workbook("labels", payload)
    except CatalogWorkbookError as exc:
        assert any(item.code == "INVALID_CONTENT" for item in exc.issues)
    else:
        raise AssertionError("expected CatalogWorkbookError")


def test_valid_logic_upload_creates_draft() -> None:
    http, store = _app()
    response = _post_catalog(
        http, "logic", logic_xlsx(), "logic.xlsx", _headers("publisher", CLIENT_A)
    )
    assert response.status_code == 201
    body = response.json()
    assert body["validation_status"] == "valid"
    assert body["version"]["status"] == "draft"
    assert body["version"]["origin"] == "uploaded"
    assert body["version"]["row_count"] == 1
    listed = http.get("/api/v1/catalogs/logic", headers=_headers("publisher", CLIENT_A)).json()
    assert listed["processing_active"]["origin"] == "packaged"
    assert listed["processing_active"]["version_id"] == CAMPAIGN_V2
    assert listed["items"][0]["version_id"] == body["version"]["version_id"]
    assert store.get(body["version"]["version_id"]).notes == UPLOAD_NOTES


def test_invalid_logic_file_type() -> None:
    http, store = _app()
    response = http.post(
        "/api/v1/catalogs/logic",
        headers=_headers("publisher", CLIENT_A),
        files={"file": ("logic.csv", b"Campaign Name\nA", "text/csv")},
    )
    assert response.status_code == 422
    assert _error(response)["code"] == "VALIDATION_ERROR"
    assert store.list_for_client(CLIENT_A, "logic") == ()


def test_invalid_logic_schema_http() -> None:
    http, store = _app()
    response = _post_catalog(
        http,
        "logic",
        catalog_xlsx(("Nope",), [{"Nope": "x"}]),
        "logic.xlsx",
        _headers("publisher", CLIENT_A),
    )
    assert response.status_code == 422
    details = _error(response)["details"]
    assert details and details[0]["code"] == "INVALID_SCHEMA"
    assert store.list_for_client(CLIENT_A, "logic") == ()


def test_invalid_logic_content_http() -> None:
    http, _store = _app()
    response = _post_catalog(
        http,
        "logic",
        logic_xlsx([{"Campaign Name": None, "Filter Logic 1": "x"}]),
        "logic.xlsx",
        _headers("publisher", CLIENT_A),
    )
    assert response.status_code == 422
    assert any(item["code"] == "INVALID_CONTENT" for item in _error(response)["details"])


def test_valid_labels_upload_creates_draft() -> None:
    http, _store = _app()
    response = _post_catalog(
        http, "labels", labels_xlsx(), "labels.xlsx", _headers("publisher", CLIENT_A)
    )
    assert response.status_code == 201
    assert response.json()["version"]["status"] == "draft"
    assert response.json()["version"]["kind"] == "labels"


def test_invalid_labels_file_type() -> None:
    http, store = _app()
    response = http.post(
        "/api/v1/catalogs/labels",
        headers=_headers("publisher", CLIENT_A),
        files={"file": ("labels.txt", b"not a workbook", "text/plain")},
    )
    assert response.status_code == 422
    assert store.list_for_client(CLIENT_A, "labels") == ()


def test_invalid_labels_schema_http() -> None:
    http, store = _app()
    response = _post_catalog(
        http,
        "labels",
        catalog_xlsx(("Group Name",), [{"Group Name": "G"}]),
        "labels.xlsx",
        _headers("publisher", CLIENT_A),
    )
    assert response.status_code == 422
    assert any(item["code"] == "INVALID_SCHEMA" for item in _error(response)["details"])
    assert store.list_for_client(CLIENT_A, "labels") == ()


def test_invalid_labels_content_http() -> None:
    http, store = _app()
    response = _post_catalog(
        http,
        "labels",
        labels_xlsx(
            [
                {"Group Name": "G1", "Filter Logic 1": V2C_FL1},
                {"Group Name": "G2", "Filter Logic 1": V2C_FL1},
            ]
        ),
        "labels.xlsx",
        _headers("publisher", CLIENT_A),
    )
    assert response.status_code == 422
    assert store.list_for_client(CLIENT_A, "labels") == ()


def test_version_creation_preserves_old_versions() -> None:
    http, _store = _app()
    first = _post_catalog(http, "logic", logic_xlsx(), "a.xlsx", _headers("publisher", CLIENT_A))
    second_rows = logic_xlsx([{"Campaign Name": "Second", "Filter Logic 1": "FL"}])
    second = _post_catalog(http, "logic", second_rows, "b.xlsx", _headers("publisher", CLIENT_A))
    assert first.status_code == 201
    assert second.status_code == 201
    listed = http.get("/api/v1/catalogs/logic", headers=_headers("publisher", CLIENT_A)).json()
    ids = {item["version_id"] for item in listed["items"]}
    assert first.json()["version"]["version_id"] in ids
    assert second.json()["version"]["version_id"] in ids
    assert listed["items"][0]["version_id"] == second.json()["version"]["version_id"]


def test_invalid_and_non_draft_cannot_activate() -> None:
    http, store = _app()
    created = _post_catalog(http, "logic", logic_xlsx(), "a.xlsx", _headers("publisher", CLIENT_A))
    version_id = created.json()["version"]["version_id"]
    activate = http.post(
        f"/api/v1/catalogs/logic/{version_id}/activate",
        headers=_headers("publisher", CLIENT_A),
    )
    assert activate.status_code == 200
    again = http.post(
        f"/api/v1/catalogs/logic/{version_id}/activate",
        headers=_headers("publisher", CLIENT_A),
    )
    assert again.status_code == 422
    missing = http.post(
        "/api/v1/catalogs/logic/aaaaaaaa-aaaa-4aaa-8aaa-aaaaaaaaaaaa/activate",
        headers=_headers("publisher", CLIENT_A),
    )
    assert missing.status_code == 404
    assert store.active_for(CLIENT_A, "logic").id == version_id


def test_valid_version_activates_and_supersedes() -> None:
    http, _store = _app()
    headers = _headers("publisher", CLIENT_A)
    first_id = _post_catalog(http, "logic", logic_xlsx(), "a.xlsx", headers).json()["version"][
        "version_id"
    ]
    http.post(f"/api/v1/catalogs/logic/{first_id}/activate", headers=headers)
    second_id = _post_catalog(
        http,
        "logic",
        logic_xlsx([{"Campaign Name": "Next", "Filter Logic 1": "N"}]),
        "b.xlsx",
        headers,
    ).json()["version"]["version_id"]
    activated = http.post(f"/api/v1/catalogs/logic/{second_id}/activate", headers=headers)
    assert activated.status_code == 200
    assert activated.json()["status"] == "active"
    listed = http.get("/api/v1/catalogs/logic", headers=headers).json()
    by_id = {item["version_id"]: item for item in listed["items"]}
    assert by_id[first_id]["status"] == "superseded"
    assert by_id[second_id]["status"] == "active"
    assert listed["processing_active"]["version_id"] == second_id


def test_failed_upload_does_not_corrupt_active_version() -> None:
    http, store = _app()
    headers = _headers("publisher", CLIENT_A)
    version_id = _post_catalog(http, "logic", logic_xlsx(), "a.xlsx", headers).json()["version"][
        "version_id"
    ]
    http.post(f"/api/v1/catalogs/logic/{version_id}/activate", headers=headers)
    failed = _post_catalog(
        http, "logic", catalog_xlsx(("Nope",), [{"Nope": "x"}]), "bad.xlsx", headers
    )
    assert failed.status_code == 422
    assert store.active_for(CLIENT_A, "logic").id == version_id
    listed = http.get("/api/v1/catalogs/logic", headers=headers).json()
    assert listed["processing_active"]["version_id"] == version_id


def test_client_isolation_and_unauthorized_access() -> None:
    http, _store = _app()
    publisher_a = _headers("publisher", CLIENT_A)
    publisher_b = _headers("publisher", CLIENT_B)
    reader = _headers("reader", CLIENT_A)
    version_id = _post_catalog(http, "logic", logic_xlsx(), "a.xlsx", publisher_a).json()[
        "version"
    ]["version_id"]
    assert http.get("/api/v1/catalogs/logic", headers=reader).status_code == 403
    assert (
        http.post(
            "/api/v1/catalogs/logic",
            headers=reader,
            files={"file": ("a.xlsx", logic_xlsx(), XLSX_TYPE)},
        ).status_code
        == 403
    )
    leak_list = http.get("/api/v1/catalogs/logic", headers=publisher_b).json()
    assert all(item["version_id"] != version_id for item in leak_list["items"])
    assert http.get(f"/api/v1/catalogs/logic/{version_id}", headers=publisher_b).status_code == 404
    assert (
        http.post(f"/api/v1/catalogs/logic/{version_id}/activate", headers=publisher_b).status_code
        == 404
    )
    assert (
        http.post(
            f"/api/v1/catalogs/logic/{version_id}/deactivate", headers=publisher_b
        ).status_code
        == 404
    )
    assert (
        http.get(f"/api/v1/catalogs/logic/{version_id}/download", headers=publisher_b).status_code
        == 404
    )
    unauth = TestClient(
        create_app(
            settings=make_settings(
                dfip_auth_mode="jwt", dfip_auth_secret=JWT_SECRET, database_url=""
            ),
            ingest_store=InMemoryIngestStore(),
            fact_store=InMemoryFactStore(),
        )
    )
    assert unauth.get("/api/v1/catalogs/logic").status_code == 401
    assert (
        unauth.post(
            "/api/v1/catalogs/logic",
            files={"file": ("a.xlsx", logic_xlsx(), XLSX_TYPE)},
        ).status_code
        == 401
    )
    assert unauth.post(f"/api/v1/catalogs/logic/{version_id}/deactivate").status_code == 401


def test_download_permissions_and_packaged() -> None:
    http, _store = _app()
    headers = _headers("publisher", CLIENT_A)
    version_id = _post_catalog(http, "logic", logic_xlsx(), "a.xlsx", headers).json()["version"][
        "version_id"
    ]
    uploaded = http.get(f"/api/v1/catalogs/logic/{version_id}/download", headers=headers)
    assert uploaded.status_code == 200
    assert "attachment" in uploaded.headers["content-disposition"]
    packaged = http.get("/api/v1/catalogs/logic/packaged/download", headers=headers)
    assert packaged.status_code == 200
    workbook = load_workbook(filename := __import__("io").BytesIO(uploaded.content))
    del filename
    sheet = workbook.active
    assert sheet.cell(1, 1).value == "Campaign Name"
    assert sheet.cell(2, 1).value == V2C_CAMPAIGN
    workbook.close()
    assert (
        http.get(
            f"/api/v1/catalogs/logic/{version_id}/download",
            headers=_headers("reader", CLIENT_A),
        ).status_code
        == 403
    )


def test_processing_uses_active_logic_and_labels(tmp_path) -> None:
    http, store = _app()
    headers = _headers("publisher", CLIENT_A)
    logic_id = _post_catalog(http, "logic", logic_xlsx(), "logic.xlsx", headers).json()["version"][
        "version_id"
    ]
    labels_id = _post_catalog(http, "labels", labels_xlsx(), "labels.xlsx", headers).json()[
        "version"
    ]["version_id"]
    assert (
        http.post(f"/api/v1/catalogs/logic/{logic_id}/activate", headers=headers).status_code == 200
    )
    assert (
        http.post(f"/api/v1/catalogs/labels/{labels_id}/activate", headers=headers).status_code
        == 200
    )
    content = workbook_bytes(
        tmp_path / "v2c.xlsx",
        [source_row(**{"Campaign ID": "v2c-1", "Campaign Name": V2C_CAMPAIGN})],
    )
    uploaded = _upload(http, content, "v2c.xlsx", headers=headers)
    assert uploaded.status_code == 201
    run = uploaded.json()["processing_run"]
    assert run["campaign_label_version_id"] == logic_id
    assert run["label_group_version_id"] == labels_id
    facts = http.get("/api/v1/facts?limit=200", headers=headers).json()
    fact = next(item for item in facts["items"] if item["campaign_id"] == "v2c-1")
    assert fact["filter_logic_1"] == V2C_FL1
    assert fact["filter_logic_2"] == V2C_FL2
    assert fact["amc_status_filter_logic_3"] == V2C_FL3
    assert fact["amc_device_category_filter_logic_4"] == V2C_FL4
    assert fact["amc_product_cat_filter_logic_5"] == V2C_FL5
    assert fact["manual_or_automated"] == V2C_MANUAL
    assert fact["filter_logic_1_group"] == V2C_GROUP
    assert fact["label_match_status"] == "matched"
    assert fact["campaign_label_version_id"] == logic_id
    assert fact["label_group_version_id"] == labels_id


def test_draft_catalog_does_not_bind_processing(tmp_path) -> None:
    http, store = _app()
    headers = _headers("publisher", CLIENT_A)
    created = _post_catalog(http, "logic", logic_xlsx(), "logic.xlsx", headers)
    assert created.status_code == 201
    assert created.json()["version"]["status"] == "draft"
    assert store.active_for(CLIENT_A, "logic") is None
    content = workbook_bytes(
        tmp_path / "packaged.xlsx",
        [source_row(**{"Campaign ID": "packaged-1", "Campaign Name": GROUP7_CAMPAIGN})],
    )
    uploaded = _upload(http, content, "packaged.xlsx", headers=headers)
    assert uploaded.status_code == 201
    run = uploaded.json()["processing_run"]
    assert run["campaign_label_version_id"] == CAMPAIGN_V2
    facts = http.get("/api/v1/facts?limit=200", headers=headers).json()
    fact = next(item for item in facts["items"] if item["campaign_id"] == "packaged-1")
    assert fact["filter_logic_1"] == GROUP7_FILTER_LOGIC_1
    assert fact["filter_logic_1_group"] == "Group7"
    assert fact["campaign_label_version_id"] == CAMPAIGN_V2


def test_deactivate_falls_back_to_packaged_processing(tmp_path) -> None:
    http, store = _app()
    headers = _headers("publisher", CLIENT_A)
    logic_id = _post_catalog(http, "logic", logic_xlsx(), "logic.xlsx", headers).json()["version"][
        "version_id"
    ]
    labels_id = _post_catalog(http, "labels", labels_xlsx(), "labels.xlsx", headers).json()[
        "version"
    ]["version_id"]
    http.post(f"/api/v1/catalogs/logic/{logic_id}/activate", headers=headers)
    http.post(f"/api/v1/catalogs/labels/{labels_id}/activate", headers=headers)
    deactivated = http.post(f"/api/v1/catalogs/logic/{logic_id}/deactivate", headers=headers)
    assert deactivated.status_code == 200
    assert deactivated.json()["status"] == "superseded"
    http.post(f"/api/v1/catalogs/labels/{labels_id}/deactivate", headers=headers)
    assert store.active_for(CLIENT_A, "logic") is None
    listed = http.get("/api/v1/catalogs/logic", headers=headers).json()
    assert listed["processing_active"]["origin"] == "packaged"
    content = workbook_bytes(
        tmp_path / "after-deactivate.xlsx",
        [source_row(**{"Campaign ID": "after-1", "Campaign Name": GROUP7_CAMPAIGN})],
    )
    uploaded = _upload(http, content, "after-deactivate.xlsx", headers=headers)
    fact = next(
        item
        for item in http.get("/api/v1/facts?limit=200", headers=headers).json()["items"]
        if item["campaign_id"] == "after-1"
    )
    assert uploaded.json()["processing_run"]["campaign_label_version_id"] == CAMPAIGN_V2
    assert fact["filter_logic_1"] == GROUP7_FILTER_LOGIC_1
    assert fact["filter_logic_1_group"] == "Group7"


def test_historical_run_keeps_original_versions(tmp_path) -> None:
    http, store = _app()
    headers = _headers("publisher", CLIENT_A)
    first_logic = _post_catalog(http, "logic", logic_xlsx(), "a.xlsx", headers).json()["version"][
        "version_id"
    ]
    first_labels = _post_catalog(http, "labels", labels_xlsx(), "a.xlsx", headers).json()[
        "version"
    ]["version_id"]
    http.post(f"/api/v1/catalogs/logic/{first_logic}/activate", headers=headers)
    http.post(f"/api/v1/catalogs/labels/{first_labels}/activate", headers=headers)
    content = workbook_bytes(
        tmp_path / "hist.xlsx",
        [source_row(**{"Campaign ID": "hist-1", "Campaign Name": V2C_CAMPAIGN})],
    )
    uploaded = _upload(http, content, "hist.xlsx", headers=headers)
    run_id = uploaded.json()["processing_run"]["processing_run_id"]
    batch_id = uploaded.json()["batch"]["batch_id"]
    second_logic = _post_catalog(
        http,
        "logic",
        logic_xlsx([{"Campaign Name": V2C_CAMPAIGN, "Filter Logic 1": "CHANGED"}]),
        "b.xlsx",
        headers,
    ).json()["version"]["version_id"]
    http.post(f"/api/v1/catalogs/logic/{second_logic}/activate", headers=headers)
    ingest = http.app.state.ingest_store
    facts = http.app.state.fact_store
    result = run_transformation(ingest, facts, batch_id, processing_run_id=run_id, catalog=store)
    assert result.processing_run.id == run_id
    fact = next(item for item in facts.facts.values() if item.campaign_id == "hist-1")
    assert fact.campaign_label_version_id == first_logic
    assert fact.filter_logic_1 == V2C_FL1
    assert store.active_for(CLIENT_A, "logic").id == second_logic


def test_failed_processing_does_not_corrupt_catalog(tmp_path) -> None:
    http, store = _app()
    headers = _headers("publisher", CLIENT_A)
    logic_id = _post_catalog(http, "logic", logic_xlsx(), "logic.xlsx", headers).json()["version"][
        "version_id"
    ]
    http.post(f"/api/v1/catalogs/logic/{logic_id}/activate", headers=headers)
    bad = workbook_bytes(
        tmp_path / "bad.xlsx",
        [source_row()],
        header_override={12: "Campaign"},
    )
    uploaded = _upload(http, bad, "bad.xlsx", headers=headers)
    assert (
        uploaded.json()["processing_run"] is None or uploaded.json()["batch"]["status"] == "failed"
    )
    assert store.active_for(CLIENT_A, "logic").id == logic_id
    assert store.active_for(CLIENT_A, "logic").status == "active"


def test_unscoped_publisher_activates_logic_with_query_client_id() -> None:
    http, store = _app()
    headers = _unscoped_headers()
    created = _post_catalog(
        http, "logic", logic_xlsx(), "logic.xlsx", headers, client_id=CLIENT_A
    )
    assert created.status_code == 201
    version_id = created.json()["version"]["version_id"]
    assert store.get(version_id).status == "draft"
    listed = http.get(
        "/api/v1/catalogs/logic", headers=headers, params={"client_id": CLIENT_A}
    ).json()
    assert listed["processing_active"]["origin"] == "packaged"
    assert listed["items"][0]["status"] == "draft"
    activated = http.post(
        f"/api/v1/catalogs/logic/{version_id}/activate",
        headers=headers,
        params={"client_id": CLIENT_A},
    )
    assert activated.status_code == 200
    assert activated.json()["status"] == "active"
    assert activated.json()["client_id"] == CLIENT_A
    assert store.get(version_id).status == "active"
    after = http.get(
        "/api/v1/catalogs/logic", headers=headers, params={"client_id": CLIENT_A}
    ).json()
    assert after["processing_active"]["version_id"] == version_id
    assert after["processing_active"]["status"] == "active"


def test_unscoped_publisher_activates_labels_with_query_client_id() -> None:
    http, store = _app()
    headers = _unscoped_headers()
    created = _post_catalog(
        http, "labels", labels_xlsx(), "labels.xlsx", headers, client_id=CLIENT_A
    )
    assert created.status_code == 201
    version_id = created.json()["version"]["version_id"]
    assert store.get(version_id).status == "draft"
    activated = http.post(
        f"/api/v1/catalogs/labels/{version_id}/activate",
        headers=headers,
        params={"client_id": CLIENT_A},
    )
    assert activated.status_code == 200
    assert activated.json()["status"] == "active"
    assert activated.json()["kind"] == "labels"
    after = http.get(
        "/api/v1/catalogs/labels", headers=headers, params={"client_id": CLIENT_A}
    ).json()
    assert after["processing_active"]["version_id"] == version_id


def test_unscoped_activate_missing_client_id_returns_422() -> None:
    http, store = _app()
    headers = _unscoped_headers()
    created = _post_catalog(
        http, "logic", logic_xlsx(), "logic.xlsx", headers, client_id=CLIENT_A
    )
    version_id = created.json()["version"]["version_id"]
    missing = http.post(f"/api/v1/catalogs/logic/{version_id}/activate", headers=headers)
    assert missing.status_code == 422
    body = _error(missing)
    assert body["code"] == "VALIDATION_ERROR"
    assert "client_id is required" in body["message"]
    assert store.get(version_id).status == "draft"
    labels = _post_catalog(
        http, "labels", labels_xlsx(), "labels.xlsx", headers, client_id=CLIENT_A
    )
    labels_id = labels.json()["version"]["version_id"]
    labels_missing = http.post(
        f"/api/v1/catalogs/labels/{labels_id}/activate", headers=headers
    )
    assert labels_missing.status_code == 422
    assert "client_id is required" in _error(labels_missing)["message"]
    assert store.get(labels_id).status == "draft"


def test_activate_wrong_client_keeps_existing_authorization() -> None:
    http, store = _app()
    bound_a = _headers("publisher", CLIENT_A)
    bound_b = _headers("publisher", CLIENT_B)
    unscoped = _unscoped_headers()
    version_id = _post_catalog(http, "logic", logic_xlsx(), "a.xlsx", bound_a).json()[
        "version"
    ]["version_id"]
    mismatch = http.post(
        f"/api/v1/catalogs/logic/{version_id}/activate",
        headers=bound_a,
        params={"client_id": CLIENT_B},
    )
    assert mismatch.status_code == 403
    assert _error(mismatch)["code"] == "AUTHORIZATION_FAILED"
    assert store.get(version_id).status == "draft"
    other_client = http.post(
        f"/api/v1/catalogs/logic/{version_id}/activate",
        headers=unscoped,
        params={"client_id": CLIENT_B},
    )
    assert other_client.status_code == 404
    assert store.get(version_id).status == "draft"
    assert (
        http.post(
            f"/api/v1/catalogs/logic/{version_id}/activate", headers=bound_b
        ).status_code
        == 404
    )
    assert store.get(version_id).status == "draft"


def test_activate_one_kind_does_not_activate_the_other() -> None:
    http, store = _app()
    headers = _unscoped_headers()
    logic_id = _post_catalog(
        http, "logic", logic_xlsx(), "logic.xlsx", headers, client_id=CLIENT_A
    ).json()["version"]["version_id"]
    labels_id = _post_catalog(
        http, "labels", labels_xlsx(), "labels.xlsx", headers, client_id=CLIENT_A
    ).json()["version"]["version_id"]
    assert store.get(logic_id).status == "draft"
    assert store.get(labels_id).status == "draft"
    logic = http.post(
        f"/api/v1/catalogs/logic/{logic_id}/activate",
        headers=headers,
        params={"client_id": CLIENT_A},
    )
    assert logic.status_code == 200
    assert store.get(logic_id).status == "active"
    assert store.get(labels_id).status == "draft"
    assert store.active_for(CLIENT_A, "labels") is None
    listed_labels = http.get(
        "/api/v1/catalogs/labels", headers=headers, params={"client_id": CLIENT_A}
    ).json()
    assert listed_labels["processing_active"]["origin"] == "packaged"
    labels = http.post(
        f"/api/v1/catalogs/labels/{labels_id}/activate",
        headers=headers,
        params={"client_id": CLIENT_A},
    )
    assert labels.status_code == 200
    assert store.get(labels_id).status == "active"
    assert store.get(logic_id).status == "active"
    assert store.active_for(CLIENT_A, "logic").id == logic_id
    assert store.active_for(CLIENT_A, "labels").id == labels_id
    listed_logic = http.get(
        "/api/v1/catalogs/logic", headers=headers, params={"client_id": CLIENT_A}
    ).json()
    assert listed_logic["processing_active"]["version_id"] == logic_id
    listed_labels = http.get(
        "/api/v1/catalogs/labels", headers=headers, params={"client_id": CLIENT_A}
    ).json()
    assert listed_labels["processing_active"]["version_id"] == labels_id


def test_spa_catalog_activate_propagates_loaded_query_client_id() -> None:
    root = Path(__file__).resolve().parents[1] / "apps" / "web" / "static" / "js"
    views = (root / "views.js").read_text(encoding="utf-8")
    app_js = (root / "app.js").read_text(encoding="utf-8")
    catalog = views[views.index("export function catalogView") :]
    default_client = catalog[catalog.index("const defaultClient") : catalog.index("const items")]
    assert 'query.get("client_id")' in default_client
    assert "data-catalog-activate" in catalog
    assert "data-catalog-client=" in catalog
    action = app_js[
        app_js.index("async function runCatalogAction") : app_js.index(
            "async function reloadPublications"
        )
    ]
    assert "data-catalog-client" in action
    assert 'query.get("client_id")' in action
    load = app_js[
        app_js.index("async function loadCatalogView") : app_js.index(
            "async function runCatalogAction"
        )
    ]
    assert 'query.get("client_id")' in load


def test_spa_catalog_wiring_does_not_claim_excel_export() -> None:
    root = Path(__file__).resolve().parents[1] / "apps" / "web" / "static" / "js"
    client_js = (root / "api-client.js").read_text(encoding="utf-8")
    app_js = (root / "app.js").read_text(encoding="utf-8")
    views = (root / "views.js").read_text(encoding="utf-8")
    assert "uploadCatalog" in client_js
    assert "/catalogs/" in client_js
    assert "uploadCatalog" in app_js
    assert "admin-catalogs" in app_js
    assert "data-catalog-upload-form" in views
    assert "excel export" not in client_js.lower()
    assert "excel export" not in app_js.lower()
    assert "excel export" not in views.lower()


def test_postgres_label_member_insert_matches_five_columns() -> None:
    source = (
        Path(__file__).resolve().parents[1] / "packages" / "db" / "dfip_db" / "catalog_store.py"
    ).read_text(encoding="utf-8")
    start = source.index("INSERT INTO label_group_member")
    chunk = source[start : start + 500]
    values = chunk[chunk.index("VALUES") : chunk.index(")", chunk.index("VALUES")) + 1]
    assert values.count("%s") == 5
    columns = chunk[chunk.index("(") + 1 : chunk.index(")")]
    names = [part.strip() for part in columns.split(",") if part.strip()]
    assert names == [
        "id",
        "version_id",
        "group_name",
        "filter_logic_1_value",
        "row_order",
    ]
