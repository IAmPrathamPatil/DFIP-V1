"""RUN 008: hide PublishedFacts and Facts on generated client workbooks."""

from __future__ import annotations

import io
import re
import time
from zipfile import ZipFile

import pytest
from dfip_web.client_report_download import (
    ARTIFACT_REFRESHABLE,
    ARTIFACT_STATIC,
    FACT_VALUE_FIELDS,
    QUERY_TABLE_PART,
    TECHNICAL_SHEET_NAMES,
    assert_technical_sheets_hidden,
    render_client_report_xlsx,
    sheet_visibility_states,
)
from dfip_web.client_workbook import FACT_HEADERS, XLSX_PATH
from dfip_web.daily_report import CLIENT_WORKBOOK_SHEET_NAMES, REPORT_SHEET_NAMES
from dfip_web.pivot_report import PIVOT_CACHE_PART, PIVOT_TABLE_NAMES, assert_native_pivot_package
from dfip_web.slicer_defaults import assert_slicer_package
from openpyxl import load_workbook

from test_p8_client_report import _campaign_ids, _settings_value
from test_p11_pivot_report import _published_fact


def _rows(*, campaign: str, channel: str = "WhatsApp"):
    return [_published_fact(campaign_id=campaign, channel=channel, month_label="Aug-25")]


def _workbook(
    rows,
    *,
    artifact: str = ARTIFACT_STATIC,
    client_id: str,
    code: str,
):
    return render_client_report_xlsx(
        rows,
        published_at=None,
        client_id=client_id,
        client_name=code,
        client_code=code,
        publication_id=f"pub-{code}",
        artifact=artifact,
        api_base_url="http://127.0.0.1:8000" if artifact == ARTIFACT_REFRESHABLE else None,
    )


def _states(body: bytes) -> dict[str, str]:
    with ZipFile(io.BytesIO(body)) as archive:
        return sheet_visibility_states(archive.read("xl/workbook.xml").decode("utf-8"))


def test_generated_workbooks_hide_technical_sheets_not_reports() -> None:
    static = _workbook(
        _rows(campaign="camp-a"),
        client_id="a0000000-0000-4000-8000-000000000001",
        code="C1",
    )
    refreshable = _workbook(
        _rows(campaign="camp-a"),
        artifact=ARTIFACT_REFRESHABLE,
        client_id="a0000000-0000-4000-8000-000000000001",
        code="C1",
    )
    for body, refresh in ((static, False), (refreshable, True)):
        assert_native_pivot_package(body)
        assert_slicer_package(body)
        assert_technical_sheets_hidden(body, facts_visible=False)
        states = _states(body)
        assert states["PublishedFacts"] == "hidden"
        assert states["Facts"] == "hidden"
        for name in REPORT_SHEET_NAMES:
            assert states[name] == "visible"
        with ZipFile(io.BytesIO(body)) as archive:
            workbook_xml = archive.read("xl/workbook.xml").decode("utf-8")
            cache = archive.read(PIVOT_CACHE_PART).decode("utf-8")
            names = set(archive.namelist())
        assert 'sheet="PublishedFacts"' in cache
        assert list(FACT_HEADERS)[:3] == ["Filter Logic 1", "Filter Logic 2", "Template Status"]
        assert len(FACT_HEADERS) == 45
        assert len(FACT_VALUE_FIELDS) == 45
        assert 'activeTab="2"' in workbook_xml
        assert 'firstSheet="2"' in workbook_xml
        assert "veryHidden" not in workbook_xml
        facts_xml = ZipFile(io.BytesIO(body)).read("xl/worksheets/sheet2.xml").decode("utf-8")
        overall_xml = ZipFile(io.BytesIO(body)).read("xl/worksheets/sheet3.xml").decode("utf-8")
        assert "tabSelected" not in facts_xml
        assert 'tabSelected="1"' in overall_xml
        assert (QUERY_TABLE_PART in names) is refresh
        if refresh:
            connections = ZipFile(io.BytesIO(body)).read("xl/connections.xml").decode("utf-8")
            assert "Location=PublishedFacts" in connections
            assert "BearerToken" not in connections
            assert "A1:AS1048576" not in cache
            assert 'name="ExternalData_1"' in cache
            assert "worksheetSource ref=" not in cache
            assert "REFRESH REQUIRED" in facts_xml
        workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=False)
        assert list(workbook.sheetnames) == list(CLIENT_WORKBOOK_SHEET_NAMES)
        assert workbook["PublishedFacts"].sheet_state == "hidden"
        assert workbook["Facts"].sheet_state == "hidden"
        assert workbook[REPORT_SHEET_NAMES[0]].sheet_state == "visible"
        workbook.close()
        assert _campaign_ids(body) == ["camp-a"]
        assert _settings_value(body, "BearerToken") in {None, ""}
        assert _settings_value(body, "ClientId") in {None, ""}


def test_company_workbooks_hide_technical_sheets_without_mixing_data() -> None:
    company_a = _workbook(
        _rows(campaign="camp-a", channel="WhatsApp"),
        client_id="a0000000-0000-4000-8000-000000000001",
        code="C1",
    )
    company_b = _workbook(
        _rows(campaign="camp-b", channel="RCS"),
        client_id="b0000000-0000-4000-8000-000000000002",
        code="C2",
    )
    new_company = _workbook(
        _rows(campaign="camp-n", channel="Email"),
        client_id="c0000000-0000-4000-8000-000000000003",
        code="NEWCO",
    )
    assert _campaign_ids(company_a) == ["camp-a"]
    assert _campaign_ids(company_b) == ["camp-b"]
    assert _campaign_ids(new_company) == ["camp-n"]
    for body, client_id in (
        (company_a, "a0000000-0000-4000-8000-000000000001"),
        (company_b, "b0000000-0000-4000-8000-000000000002"),
        (new_company, "c0000000-0000-4000-8000-000000000003"),
    ):
        assert_technical_sheets_hidden(body)
        with ZipFile(io.BytesIO(body)) as archive:
            custom = archive.read("docProps/custom.xml").decode("utf-8")
            cache = archive.read(PIVOT_CACHE_PART).decode("utf-8")
            slicers = [
                n
                for n in archive.namelist()
                if n.startswith("xl/slicerCaches/") and n.endswith(".xml")
            ]
        assert client_id in custom
        assert 'sheet="PublishedFacts"' in cache
        assert len(slicers) == 35
        names = re.findall(r'<cacheField name="([^"]+)"', cache)
        assert names[:45] == list(FACT_HEADERS)


def test_tracked_template_keeps_technical_sheets_visible() -> None:
    with ZipFile(XLSX_PATH) as archive:
        states = sheet_visibility_states(archive.read("xl/workbook.xml").decode("utf-8"))
    for name in TECHNICAL_SHEET_NAMES:
        assert states[name] == "visible"
    for name in REPORT_SHEET_NAMES:
        assert states[name] == "visible"


def _dispatch_excel():
    win32com_client = pytest.importorskip("win32com.client")
    try:
        excel = win32com_client.DispatchEx("Excel.Application")
    except Exception as exc:
        pytest.skip(f"Excel COM not available: {exc}")
    excel.Visible = False
    excel.DisplayAlerts = False
    excel.AskToUpdateLinks = False
    excel.EnableEvents = False
    return excel


def test_desktop_excel_hidden_sheets_save_reopen(tmp_path) -> None:
    """Hidden technical sheets, pivots/slicers, PivotCache.Refresh. Not live Refresh All."""
    xl_sheet_hidden = 0
    xl_sheet_visible = -1
    static_path = tmp_path / "run008_static.xlsx"
    refresh_path = tmp_path / "run008_refreshable.xlsm"
    rows = _rows(campaign="camp-a")
    static_path.write_bytes(
        _workbook(rows, client_id="a0000000-0000-4000-8000-000000000001", code="C1")
    )
    refresh_path.write_bytes(
        _workbook(
            rows,
            artifact=ARTIFACT_REFRESHABLE,
            client_id="a0000000-0000-4000-8000-000000000001",
            code="C1",
        )
    )
    excel = _dispatch_excel()
    try:
        for path in (static_path, refresh_path):
            workbook = excel.Workbooks.Open(str(path.resolve()), UpdateLinks=0, ReadOnly=False)
            excel.ActiveWindow.SelectedSheets.Item(1).Select()
            assert workbook.Worksheets.Count == 11
            assert workbook.Worksheets("PublishedFacts").Visible == xl_sheet_hidden
            assert workbook.Worksheets("Facts").Visible == xl_sheet_hidden
            visible = [
                str(ws.Name) for ws in workbook.Worksheets if int(ws.Visible) == xl_sheet_visible
            ]
            assert visible == list(REPORT_SHEET_NAMES)
            assert workbook.SlicerCaches.Count == 35
            assert sum(int(ws.PivotTables().Count) for ws in workbook.Worksheets) == 9
            overall = workbook.Worksheets(REPORT_SHEET_NAMES[0])
            overall.PivotTables(1).PivotCache().Refresh()
            assert workbook.Worksheets("PublishedFacts").Visible == xl_sheet_hidden
            facts = workbook.Worksheets("Facts")
            assert facts.Visible == xl_sheet_hidden
            assert facts.ListObjects.Count >= 1
            assert workbook.SlicerCaches.Count == 35
            assert str(overall.PivotTables(1).Name) in PIVOT_TABLE_NAMES.values()
            saved = path.with_name(path.stem + "_reopen.xlsx")
            workbook.SaveAs(str(saved.resolve()))
            workbook.Close(False)
            reopened = excel.Workbooks.Open(str(saved.resolve()), UpdateLinks=0, ReadOnly=True)
            assert reopened.Worksheets("PublishedFacts").Visible == xl_sheet_hidden
            assert reopened.Worksheets("Facts").Visible == xl_sheet_hidden
            assert reopened.SlicerCaches.Count == 35
            assert sum(int(ws.PivotTables().Count) for ws in reopened.Worksheets) == 9
            reopened.Close(False)
    finally:
        try:
            excel.Quit()
        except Exception:
            pass
        time.sleep(0.4)
