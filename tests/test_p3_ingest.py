"""P3 ingestion tests. In-memory store only. No live PostgreSQL. No P4 engines."""

from __future__ import annotations

import inspect
import shutil
from datetime import date, datetime
from pathlib import Path

import pytest
from dfip_core import ingest as ingest_pkg
from dfip_core.ingest import (
    InMemoryIngestStore,
    bind_versions_for_day,
    expected_source_headers,
    ingest_workbook,
    inspect_workbook,
    sha256_file,
)
from dfip_core.ingest.headers import HeaderContractError, locate_source_headers
from dfip_core.ingest.pipeline import ingest_workbook as pipeline_fn
from dfip_core.ingest.reader import payload_from_values, serialize_cell
from dfip_db.catalog import SOURCE_COLUMNS, WEB_ENGAGE_SOURCE_HEADERS
from dfip_db.paths import repo_root
from openpyxl import Workbook

DERIVED_HEADERS = tuple(
    col.excel_header for col in SOURCE_COLUMNS if col.source_role == "derived_excel"
)
PROD_SAMPLE = repo_root() / "14. WE Report Raw Data Oct-25 VJ Prod.xlsx"
DAILY_REPORT = repo_root() / "Web Engage - Daily Report - FY-2026.xlsx"


def _write_workbook(
    path: Path,
    rows: list[dict[str, object]],
    *,
    layout: str = "legacy",
    header_override: dict[int, str] | None = None,
    sheet: str = "Web-Engage Raw",
) -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    headers = list(WEB_ENGAGE_SOURCE_HEADERS)
    start = 11 if layout == "legacy" else 1
    if layout == "legacy":
        for index, header in enumerate(DERIVED_HEADERS, start=1):
            worksheet.cell(1, index, header)
    for index, header in enumerate(headers, start=start):
        worksheet.cell(1, index, header)
    if header_override:
        for col, value in header_override.items():
            worksheet.cell(1, col, value)
    for offset, payload in enumerate(rows):
        excel_row = 2 + offset
        for index, header in enumerate(headers):
            if header in payload:
                worksheet.cell(excel_row, start + index, payload[header])
    workbook.save(path)
    workbook.close()
    return path


def test_exact_fifty_seven_column_header_contract() -> None:
    headers = expected_source_headers()
    assert len(headers) == 57
    assert headers == WEB_ENGAGE_SOURCE_HEADERS
    assert headers[0] == "Day"
    assert headers[-1] == "Template Name (WhatsApp)"
    assert "Filter Logic 1" not in headers
    assert "Total Cost" not in headers


def test_header_mismatch_is_explicit(tmp_path: Path) -> None:
    path = _write_workbook(
        tmp_path / "bad.xlsx",
        [{"Day": datetime(2025, 4, 1), "Campaign Name": "X"}],
        header_override={12: "Campaign"},
    )
    result = ingest_workbook(path, InMemoryIngestStore())
    assert result.batch.status == "failed"
    assert result.batch.row_count_rejected == 1
    assert result.batch.error_summary is not None
    assert "HEADER_CONTRACT" in result.batch.error_summary


def test_locate_headers_reports_missing_and_unexpected() -> None:
    expected = list(WEB_ENGAGE_SOURCE_HEADERS)
    expected[1] = "Campaign"
    with pytest.raises(HeaderContractError) as caught:
        locate_source_headers(expected)
    assert caught.value.reason_code == "HEADER_CONTRACT"
    assert "Campaign Name" in caught.value.missing
    assert "Campaign" in caught.value.unexpected


def test_duplicate_header_detection() -> None:
    values = list(WEB_ENGAGE_SOURCE_HEADERS)
    values[1] = "Day"
    with pytest.raises(HeaderContractError) as caught:
        locate_source_headers(values)
    assert caught.value.reason_code == "DUPLICATE_HEADERS"
    assert "Day" in caught.value.duplicates


def test_sha256_is_deterministic(tmp_path: Path) -> None:
    rows = [{"Day": datetime(2025, 4, 1), "Campaign Name": "Alpha", "Campaign ID": "c1"}]
    first = _write_workbook(tmp_path / "a.xlsx", rows)
    second = tmp_path / "b.xlsx"
    shutil.copyfile(first, second)
    assert sha256_file(first) == sha256_file(second)
    assert len(sha256_file(first)) == 64


def test_source_file_identity_and_batch_idempotency(tmp_path: Path) -> None:
    path = _write_workbook(
        tmp_path / "same.xlsx",
        [
            {
                "Day": datetime(2025, 6, 15),
                "Campaign Name": "Alpha",
                "Campaign ID": "c1",
                "Variation ID": "v1",
            }
        ],
    )
    store = InMemoryIngestStore()
    first = ingest_workbook(path, store)
    second = ingest_workbook(path, store)
    assert first.replayed is False
    assert second.replayed is True
    assert first.source_file.id == second.source_file.id
    assert first.batch.id == second.batch.id
    assert first.source_file.sha256 == sha256_file(path)
    assert len(store.batches) == 1
    assert len(store.staged_rows) == 1
    assert first.batch.status == "staged"


def test_batch_creation_metadata(tmp_path: Path) -> None:
    path = _write_workbook(
        tmp_path / "meta.xlsx",
        [{"Day": datetime(2025, 8, 2), "Campaign Name": "A", "Campaign ID": "1"}],
    )
    result = ingest_workbook(path, InMemoryIngestStore())
    assert result.batch.worksheet_name == "Web-Engage Raw"
    assert result.batch.header_row == 1
    assert result.batch.source_start_column == "K"
    assert result.source_file.source_kind == "legacy_workbook"
    assert result.processing_run is not None
    assert result.processing_run.status == "pending"


def test_source_row_order_and_empty_row_skip(tmp_path: Path) -> None:
    path = _write_workbook(
        tmp_path / "order.xlsx",
        [
            {"Day": datetime(2025, 4, 1), "Campaign Name": "first", "Campaign ID": "1"},
            {},
            {"Day": datetime(2025, 4, 2), "Campaign Name": "third", "Campaign ID": "3"},
        ],
    )
    store = InMemoryIngestStore()
    result = ingest_workbook(path, store)
    rows = store.staged_for_batch(result.batch.id)
    assert [row.source_row_number for row in rows] == [2, 4]
    assert [row.raw["Campaign Name"] for row in rows] == ["first", "third"]
    assert result.empty_row_count == 1


def test_exact_value_preservation_no_trim_blank_null(tmp_path: Path) -> None:
    path = _write_workbook(
        tmp_path / "preserve.xlsx",
        [
            {
                "Day": datetime(2025, 4, 1),
                "Campaign Name": " Email Additional: Drop_off 14th Oct",
                "Campaign ID": " id ",
                "Channel": None,
                "Sent": 0,
            }
        ],
    )
    store = InMemoryIngestStore()
    result = ingest_workbook(path, store)
    raw = store.staged_for_batch(result.batch.id)[0].raw
    assert raw["Campaign Name"] == " Email Additional: Drop_off 14th Oct"
    assert raw["Campaign ID"] == " id "
    assert raw["Variation ID"] is None
    assert raw["Channel"] is None
    assert raw["Sent"] == 0
    assert raw["Day"] == "2025-04-01"
    assert "Filter Logic 1" not in raw
    assert len(raw) == 57


def test_serialize_does_not_trim() -> None:
    assert serialize_cell("  x  ") == "  x  "
    assert serialize_cell("") == ""
    assert serialize_cell(None) is None


def test_native_export_headers_start_at_a(tmp_path: Path) -> None:
    path = _write_workbook(
        tmp_path / "native.xlsx",
        [{"Day": datetime(2025, 4, 1), "Campaign Name": "n", "Campaign ID": "1"}],
        layout="native",
        sheet="Export",
    )
    result = ingest_workbook(path, InMemoryIngestStore())
    assert result.batch.status == "staged"
    assert result.batch.source_start_column == "A"
    assert result.source_file.source_kind == "native_export"


def test_missing_sheet_is_rejected(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active.title = "Overall Daywise Report "
    workbook.active["A1"] = "Something else"
    miss = tmp_path / "daily-like.xlsx"
    workbook.save(miss)
    workbook.close()
    result = ingest_workbook(miss, InMemoryIngestStore())
    assert result.batch.status == "failed"
    assert result.batch.row_count_rejected == 1
    assert result.processing_run is None


def test_processing_run_binds_p2_versions_without_resolving(tmp_path: Path) -> None:
    path = _write_workbook(
        tmp_path / "june.xlsx",
        [{"Day": datetime(2025, 6, 15), "Campaign Name": "A", "Campaign ID": "1"}],
    )
    result = ingest_workbook(path, InMemoryIngestStore())
    binding = result.version_binding
    assert binding is not None
    assert binding.campaign_version_label == "campaign-v1"
    assert binding.template_version_label == "template-v2"
    assert binding.rate_card_version_label == "rate-v1"
    assert binding.label_group_version_label == "fl1-group-v1"
    assert binding.campaign_label_version_id == "a0000000-0000-4000-8000-000000000021"
    assert binding.rate_card_version_id == "a0000000-0000-4000-8000-000000000011"
    august = bind_versions_for_day(date(2025, 8, 1))
    assert august.campaign_version_label == "campaign-v2"
    assert august.template_version_label == "template-v4"
    assert august.rate_card_version_label == "rate-v2"


def test_ingest_does_not_implement_p4_logic() -> None:
    source = inspect.getsource(ingest_pkg) + inspect.getsource(pipeline_fn)
    lowered = source.lower()
    assert "resolve_campaign_label" not in source
    assert "resolve_template_status" not in source
    assert "resolve_rate_card_rule" not in source
    assert "delivered *" not in lowered
    assert "total_cost" not in lowered
    assert "fastapi" not in lowered
    assert "row level security" not in lowered


def test_payload_mapping_uses_exact_headers() -> None:
    values = [None] * 67
    values[10] = datetime(2025, 4, 1)
    values[11] = " keep "
    values[14] = ""
    raw = payload_from_values(values, 10)
    assert list(raw) == list(WEB_ENGAGE_SOURCE_HEADERS)
    assert raw["Campaign Name"] == " keep "
    assert raw["Variation ID"] == ""


@pytest.mark.skipif(not PROD_SAMPLE.exists(), reason="supplied Prod workbook not present")
def test_supplied_prod_workbook_header_contract() -> None:
    layout = inspect_workbook(PROD_SAMPLE)
    assert layout.worksheet_name == "Web-Engage Raw"
    assert layout.match.start_column_letter == "K"
    assert layout.match.header_row == 1
    assert layout.match.headers == WEB_ENGAGE_SOURCE_HEADERS
    assert layout.match.source_kind == "legacy_workbook"


@pytest.mark.skipif(not DAILY_REPORT.exists(), reason="Daily Report workbook not present")
def test_daily_report_is_not_a_source_workbook() -> None:
    with pytest.raises(HeaderContractError) as caught:
        inspect_workbook(DAILY_REPORT)
    assert caught.value.reason_code == "MISSING_SHEET"


def test_p3_migration_adds_batch_metadata_without_rls() -> None:
    sql = (
        repo_root() / "supabase" / "migrations" / "20260823000010_p3_batch_ingest_metadata.sql"
    ).read_text(encoding="utf-8")
    assert "worksheet_name" in sql
    assert "header_row" in sql
    assert "source_start_column" in sql
    assert "enable row level security" not in sql.lower()
    assert "create policy" not in sql.lower()
