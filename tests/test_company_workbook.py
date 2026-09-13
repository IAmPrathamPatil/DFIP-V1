"""RUN 005B: company-bound static workbook download, provenance, and recovery.

Does not implement live Refresh All. Does not open hosted databases.
PostgreSQL cases require DFIP_TEST_DATABASE_URL.
"""

from __future__ import annotations

import io
import re
from datetime import UTC, datetime
from pathlib import Path
from zipfile import ZipFile

from dfip_api.app import create_app
from dfip_api.publication_service import (
    NO_PUBLICATION_WORKBOOK_MESSAGE,
    ROW_CAP_WORKBOOK_MESSAGE,
)
from dfip_db.local_demo_guard import COMPANY_2_CLIENT_ID, DEFAULT_CLIENT_ID
from dfip_web.client_report_download import (
    CLIENT_REPORT_DOWNLOAD_NAME,
    CLIENT_REPORT_DOWNLOAD_SUFFIX,
    CUSTOM_PROPS_PART,
    client_report_download_filename,
    render_client_report_xlsx,
)
from dfip_web.client_workbook import FACT_HEADERS
from dfip_web.daily_report import CLIENT_WORKBOOK_SHEET_NAMES, REPORT_SHEET_NAMES
from dfip_web.pivot_report import PIVOT_CACHE_PART, assert_native_pivot_package
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from postgres_support import postgres_only, requires_postgres
from test_company_registry import (
    CLIENT2_PASSWORD,
    CLIENT_PASSWORD,
    DEMO_CLIENT_2_SUBJECT,
    DEMO_CLIENT_SUBJECT,
    DEMO_PUBLISHER_2_SUBJECT,
    DEMO_PUBLISHER_SUBJECT,
    PUBLISHER2_PASSWORD,
    PUBLISHER_PASSWORD,
    WEB_STATIC,
    _bearer,
    _memory_app,
    _publish_synthetic,
    _select,
    _token,
)
from test_p5_api import AUTH, CLIENT_ID, JWT_SECRET, make_settings
from test_p7_publication import _error, _publish, publisher_app
from test_p8_client_report import SECRET_BYTES, _campaign_ids, _scan_workbook_text, _settings_value

CAMP_A = "camp-005b-a"
CAMP_B = "camp-005b-b"
CAMP_A2 = "camp-005b-a2"
FILENAME_RE = re.compile(
    r'^attachment; filename="DFIP_[A-Za-z0-9._-]+_\d{4}-\d{2}-\d{2}_Client_Report\.xlsx"$'
)


def _custom_props(body: bytes) -> dict[str, str]:
    with ZipFile(io.BytesIO(body)) as archive:
        xml = archive.read(CUSTOM_PROPS_PART).decode("utf-8")
    found: dict[str, str] = {}
    for name, value in re.findall(
        r'name="([^"]+)"[^>]*>\s*<vt:lpwstr>([^<]*)</vt:lpwstr>',
        xml,
    ):
        found[name] = (
            value.replace("&lt;", "<")
            .replace("&gt;", ">")
            .replace("&amp;", "&")
            .replace("&quot;", '"')
        )
    return found


def _facts_row5(body: bytes) -> dict[str, str]:
    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=False)
    sheet = workbook["Facts"]
    values = [cell.value for cell in next(sheet.iter_rows(min_row=5, max_row=5, max_col=6))]
    workbook.close()
    pairs = dict(zip(values[0::2], values[1::2], strict=False))
    return {str(key): "" if value is None else str(value) for key, value in pairs.items() if key}


def _published_headers(body: bytes) -> tuple[str, ...]:
    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=False)
    sheet = workbook["PublishedFacts"]
    headers = tuple(
        "" if cell.value is None else str(cell.value)
        for cell in next(sheet.iter_rows(min_row=1, max_row=1))
    )
    workbook.close()
    return headers


def _cache_source_ref(body: bytes) -> str:
    with ZipFile(io.BytesIO(body)) as archive:
        cache = archive.read(PIVOT_CACHE_PART).decode("utf-8")
    match = re.search(r'<worksheetSource ref="([^"]+)" sheet="PublishedFacts"/>', cache)
    assert match is not None
    return match.group(1)


def _cache_field_names(body: bytes) -> list[str]:
    with ZipFile(io.BytesIO(body)) as archive:
        cache = archive.read(PIVOT_CACHE_PART).decode("utf-8")
    return re.findall(r'<cacheField name="([^"]+)"', cache)


def _part_count(body: bytes, folder: str) -> int:
    with ZipFile(io.BytesIO(body)) as archive:
        return sum(
            1
            for name in archive.namelist()
            if name.startswith(folder) and name.endswith(".xml") and "/_rels/" not in name
        )


def _assert_static_and_native(body: bytes) -> None:
    assert body.startswith(b"PK")
    assert_native_pivot_package(body)
    assert _part_count(body, "xl/pivotTables/") == 9
    assert _part_count(body, "xl/slicerCaches/") == 35
    workbook_xml = ZipFile(io.BytesIO(body)).read("xl/workbook.xml").decode("utf-8")
    for name in CLIENT_WORKBOOK_SHEET_NAMES:
        assert f'name="{name}"' in workbook_xml
    assert list(REPORT_SHEET_NAMES) == list(CLIENT_WORKBOOK_SHEET_NAMES)[2:]
    assert _published_headers(body) == FACT_HEADERS
    last_col = get_column_letter(len(FACT_HEADERS))
    assert _cache_source_ref(body).startswith(f"A1:{last_col}")
    assert _cache_field_names(body)[: len(FACT_HEADERS)] == list(FACT_HEADERS)
    assert _settings_value(body, "BearerToken") in {None, ""}
    assert _settings_value(body, "ClientId") in {None, ""}
    assert _settings_value(body, "ApiBaseUrl") in {None, ""}
    text = _scan_workbook_text(body)
    for token in SECRET_BYTES:
        assert token not in text
    with ZipFile(io.BytesIO(body)) as archive:
        connections = archive.read("xl/connections.xml").decode("utf-8")
        names = set(archive.namelist())
    assert 'refreshOnLoad="0"' in connections
    assert not any(name.startswith("xl/queryTables/") for name in names)


def _assert_filename(disposition: str, client_code: str, published_at: str | None) -> None:
    assert FILENAME_RE.match(disposition)
    assert CLIENT_REPORT_DOWNLOAD_NAME in disposition
    assert disposition.endswith(f'_{CLIENT_REPORT_DOWNLOAD_SUFFIX}"')
    stamp = None
    if published_at:
        parsed = datetime.fromisoformat(published_at.replace("Z", "+00:00"))
        stamp = client_report_download_filename(client_code, parsed)
    assert f"DFIP_{client_code}_" in disposition
    if stamp:
        assert stamp in disposition


def test_filename_helper_uses_code_not_display_name() -> None:
    published = datetime(2026, 9, 1, 12, 0, tzinfo=UTC)
    name = client_report_download_filename("default", published)
    assert name == "DFIP_default_2026-09-01_Client_Report.xlsx"
    assert "Alpha" not in name
    assert CLIENT_REPORT_DOWNLOAD_NAME == "Client_Report.xlsx"


def test_no_publication_returns_explicit_error() -> None:
    app, *_rest = publisher_app()
    http = TestClient(app)
    response = http.get(
        "/api/v1/publications/current/client-report.xlsx",
        headers=AUTH,
        params={"client_id": CLIENT_ID},
    )
    assert response.status_code == 404
    assert _error(response)["code"] == "NOT_FOUND"
    assert _error(response)["message"] == NO_PUBLICATION_WORKBOOK_MESSAGE
    assert not response.content.startswith(b"PK")


def test_row_cap_returns_explicit_error() -> None:
    app, *_rest = publisher_app(dfip_download_max_rows=1)
    http = TestClient(app)
    created = _publish(http)
    assert created.status_code == 201
    response = http.get(
        "/api/v1/publications/current/client-report.xlsx",
        headers=AUTH,
        params={"client_id": CLIENT_ID},
    )
    assert response.status_code == 422
    assert _error(response)["code"] == "VALIDATION_ERROR"
    assert _error(response)["message"] == ROW_CAP_WORKBOOK_MESSAGE
    assert not response.content.startswith(b"PK")


def test_company_a_and_b_workbooks_are_isolated(tmp_path: Path) -> None:
    http = TestClient(_memory_app())
    pub_a = _bearer(_token(http, DEMO_PUBLISHER_SUBJECT, PUBLISHER_PASSWORD, DEFAULT_CLIENT_ID))
    pub_b = _bearer(_token(http, DEMO_PUBLISHER_2_SUBJECT, PUBLISHER2_PASSWORD))
    first = _publish_synthetic(
        http, tmp_path, publisher=pub_a, client_id=DEFAULT_CLIENT_ID, campaign_id=CAMP_A
    )
    second = _publish_synthetic(
        http, tmp_path, publisher=pub_b, client_id=COMPANY_2_CLIENT_ID, campaign_id=CAMP_B
    )
    report_a = http.get("/api/v1/publications/current/client-report.xlsx", headers=pub_a)
    report_b = http.get("/api/v1/publications/current/client-report.xlsx", headers=pub_b)
    assert report_a.status_code == 200, report_a.text
    assert report_b.status_code == 200, report_b.text
    steal = http.get(
        "/api/v1/publications/current/client-report.xlsx",
        headers=pub_a,
        params={"client_id": COMPANY_2_CLIENT_ID},
    )
    assert steal.status_code == 403
    other = http.get(
        f"/api/v1/publications/{second['publication']['publication_id']}/client-report.xlsx",
        headers=pub_a,
    )
    assert other.status_code == 404

    _assert_filename(
        report_a.headers["content-disposition"],
        "default",
        first["publication"]["published_at"],
    )
    _assert_filename(
        report_b.headers["content-disposition"],
        "company-2",
        second["publication"]["published_at"],
    )
    assert "Alpha" not in report_a.headers["content-disposition"]
    assert "Beta" not in report_b.headers["content-disposition"]

    body_a = report_a.content
    body_b = report_b.content
    _assert_static_and_native(body_a)
    _assert_static_and_native(body_b)
    assert _campaign_ids(body_a) == [CAMP_A]
    assert _campaign_ids(body_b) == [CAMP_B]
    text_a = _scan_workbook_text(body_a)
    text_b = _scan_workbook_text(body_b)
    assert CAMP_B not in text_a
    assert CAMP_A not in text_b
    assert COMPANY_2_CLIENT_ID not in text_a
    assert DEFAULT_CLIENT_ID in text_a
    assert DEFAULT_CLIENT_ID not in text_b
    assert COMPANY_2_CLIENT_ID in text_b

    props_a = _custom_props(body_a)
    props_b = _custom_props(body_b)
    assert props_a["client_id"] == DEFAULT_CLIENT_ID
    assert props_b["client_id"] == COMPANY_2_CLIENT_ID
    assert props_a["publication_id"] == first["publication"]["publication_id"]
    assert props_b["publication_id"] == second["publication"]["publication_id"]
    assert props_a["published_at"]
    assert props_b["published_at"]
    row_a = _facts_row5(body_a)
    row_b = _facts_row5(body_b)
    assert row_a["Company"] == "Alpha Co"
    assert row_a["Code"] == "default"
    assert row_b["Company"] == "Beta Co"
    assert row_b["Code"] == "company-2"
    assert "Beta Co" not in text_a
    assert "Alpha Co" not in text_b
    assert _settings_value(body_a, "ClientId") in {None, ""}
    assert _settings_value(body_b, "ClientId") in {None, ""}


def test_redownload_after_new_publication_is_new_snapshot(tmp_path: Path) -> None:
    http = TestClient(_memory_app())
    publisher = _bearer(_token(http, DEMO_PUBLISHER_SUBJECT, PUBLISHER_PASSWORD, DEFAULT_CLIENT_ID))
    first = _publish_synthetic(
        http, tmp_path, publisher=publisher, client_id=DEFAULT_CLIENT_ID, campaign_id=CAMP_A
    )
    old = http.get("/api/v1/publications/current/client-report.xlsx", headers=publisher)
    assert old.status_code == 200
    old_body = old.content
    pub_a = first["publication"]["publication_id"]
    second = _publish_synthetic(
        http, tmp_path, publisher=publisher, client_id=DEFAULT_CLIENT_ID, campaign_id=CAMP_A2
    )
    recovered = http.get("/api/v1/publications/current/client-report.xlsx", headers=publisher)
    historical = http.get(
        f"/api/v1/publications/{pub_a}/client-report.xlsx",
        headers=publisher,
    )
    assert recovered.status_code == 200
    assert historical.status_code == 200
    assert _campaign_ids(old_body) == [CAMP_A]
    assert _campaign_ids(historical.content) == [CAMP_A]
    assert _campaign_ids(recovered.content) == [CAMP_A2]
    assert CAMP_A2 not in _scan_workbook_text(old_body)
    assert _custom_props(old_body)["publication_id"] == pub_a
    assert (
        _custom_props(recovered.content)["publication_id"]
        == second["publication"]["publication_id"]
    )
    assert _custom_props(historical.content)["publication_id"] == pub_a


def test_bound_publisher_cannot_download_other_company_while_selected(tmp_path: Path) -> None:
    http = TestClient(_memory_app())
    token = _token(http, DEMO_PUBLISHER_SUBJECT, PUBLISHER_PASSWORD)
    selected = _select(http, token, DEFAULT_CLIENT_ID)
    assert selected.status_code == 200, selected.text
    publisher = _bearer(selected.json()["access_token"])
    other = _bearer(_token(http, DEMO_PUBLISHER_2_SUBJECT, PUBLISHER2_PASSWORD))
    _publish_synthetic(
        http, tmp_path, publisher=publisher, client_id=DEFAULT_CLIENT_ID, campaign_id=CAMP_A
    )
    _publish_synthetic(
        http, tmp_path, publisher=other, client_id=COMPANY_2_CLIENT_ID, campaign_id=CAMP_B
    )
    own = http.get("/api/v1/publications/current/client-report.xlsx", headers=publisher)
    denied = http.get(
        "/api/v1/publications/current/client-report.xlsx",
        headers=publisher,
        params={"client_id": COMPANY_2_CLIENT_ID},
    )
    assert own.status_code == 200
    assert denied.status_code == 403
    assert _campaign_ids(own.content) == [CAMP_A]


def test_new_company_workbook_after_publication(tmp_path: Path) -> None:
    http = TestClient(_memory_app())
    publisher = _bearer(_token(http, DEMO_PUBLISHER_SUBJECT, PUBLISHER_PASSWORD))
    created = http.post("/api/v1/clients", headers=publisher, json={"name": "Gamma Co"})
    assert created.status_code == 201, created.text
    new_id = created.json()["client_id"]
    new_code = created.json()["code"]
    selected = http.post(
        "/api/v1/auth/select-client",
        headers=publisher,
        json={"client_id": new_id},
    )
    assert selected.status_code == 200, selected.text
    bound = _bearer(selected.json()["access_token"])
    missing = http.get("/api/v1/publications/current/client-report.xlsx", headers=bound)
    assert missing.status_code == 404
    assert _error(missing)["message"] == NO_PUBLICATION_WORKBOOK_MESSAGE
    published = _publish_synthetic(
        http, tmp_path, publisher=bound, client_id=new_id, campaign_id="camp-005b-new"
    )
    report = http.get("/api/v1/publications/current/client-report.xlsx", headers=bound)
    assert report.status_code == 200, report.text
    _assert_filename(
        report.headers["content-disposition"],
        new_code,
        published["publication"]["published_at"],
    )
    assert _campaign_ids(report.content) == ["camp-005b-new"]
    assert _custom_props(report.content)["client_id"] == new_id
    assert _facts_row5(report.content)["Company"] == "Gamma Co"
    assert CAMP_A not in _scan_workbook_text(report.content)


def test_generator_stamps_inert_provenance_and_keeps_settings_cleared() -> None:
    body = render_client_report_xlsx(
        [],
        published_at=datetime(2026, 9, 1, tzinfo=UTC),
        client_id=DEFAULT_CLIENT_ID,
        client_code="default",
        client_name="Alpha Co",
        publication_id="p0000000-0000-4000-8000-000000000001",
    )
    _assert_static_and_native(body)
    props = _custom_props(body)
    assert props["client_id"] == DEFAULT_CLIENT_ID
    assert props["publication_id"] == "p0000000-0000-4000-8000-000000000001"
    row = _facts_row5(body)
    assert row["Company"] == "Alpha Co"
    assert row["Code"] == "default"
    assert _settings_value(body, "ClientId") in {None, ""}


def test_spa_exposes_download_company_workbook() -> None:
    views = (WEB_STATIC / "js" / "views.js").read_text(encoding="utf-8")
    app_js = (WEB_STATIC / "js" / "app.js").read_text(encoding="utf-8")
    assert "Download Company Workbook" in views
    assert "data-download-company-workbook" in views
    assert "data-download-client-report" in views
    assert "data-active-company-download" in views
    assert "Company workbook downloaded." in app_js
    assert "downloadClientReport" in app_js


@requires_postgres
@postgres_only
def test_postgres_company_workbook_isolation(pg_conn, postgres_url: str, tmp_path: Path) -> None:
    from dfip_api.local_demo_seed import seed_company2_identities, seed_demo_identities

    seed_demo_identities(
        pg_conn,
        publisher_password=PUBLISHER_PASSWORD,
        client_password=CLIENT_PASSWORD,
        iterations=1000,
    )
    seed_company2_identities(
        pg_conn,
        publisher_password=PUBLISHER2_PASSWORD,
        client_password=CLIENT2_PASSWORD,
        iterations=1000,
    )
    pg_conn.commit()
    app = create_app(
        settings=make_settings(
            dfip_auth_mode="jwt",
            dfip_auth_secret=JWT_SECRET,
            database_url=postgres_url,
            dfip_password_pbkdf2_iterations=1000,
            dfip_env="test",
        )
    )
    with TestClient(app) as http:
        pub_a = _bearer(_token(http, DEMO_PUBLISHER_SUBJECT, PUBLISHER_PASSWORD, DEFAULT_CLIENT_ID))
        pub_b = _bearer(_token(http, DEMO_PUBLISHER_2_SUBJECT, PUBLISHER2_PASSWORD))
        _publish_synthetic(
            http, tmp_path, publisher=pub_a, client_id=DEFAULT_CLIENT_ID, campaign_id=CAMP_A
        )
        _publish_synthetic(
            http, tmp_path, publisher=pub_b, client_id=COMPANY_2_CLIENT_ID, campaign_id=CAMP_B
        )
        client_a = _bearer(_token(http, DEMO_CLIENT_SUBJECT, CLIENT_PASSWORD))
        client_b = _bearer(_token(http, DEMO_CLIENT_2_SUBJECT, CLIENT2_PASSWORD))
        report_a = http.get("/api/v1/publications/current/client-report.xlsx", headers=client_a)
        report_b = http.get("/api/v1/publications/current/client-report.xlsx", headers=client_b)
        assert report_a.status_code == 200, report_a.text
        assert report_b.status_code == 200, report_b.text
        assert _custom_props(report_a.content)["client_id"] == DEFAULT_CLIENT_ID
        assert _custom_props(report_b.content)["client_id"] == COMPANY_2_CLIENT_ID
        assert CAMP_B not in _scan_workbook_text(report_a.content)
        assert CAMP_A not in _scan_workbook_text(report_b.content)
        assert CLIENT_REPORT_DOWNLOAD_NAME in report_a.headers.get("content-disposition", "")
